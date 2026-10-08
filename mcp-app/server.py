"""PowerPoint Live - an MCP App that drives desktop PowerPoint (Windows, COM) and shows the deck live.

    uvx --with "mcp<2" --with python-pptx --with pillow --with "pywin32; sys_platform == 'win32'" python mcp-app/server.py

Hosts that support MCP Apps (Claude Desktop, claude.ai) show the live view; elsewhere the tools still work.

Tools for the model (the view uses them too):
  powerpoint_open      open a deck (or attach to an open one) and pick the current slide
  powerpoint_run       run Python against the deck: app, prs, slide, goto(n), win32com in scope
  powerpoint_show      make slide n current
  powerpoint_move      move slides (by SlideID) before a slide or to the end of a section
  powerpoint_sections  add, rename or delete sections
  powerpoint_hide      hide or unhide slides
  powerpoint_fix       apply a lint fix to one shape, live
  powerpoint_history   list saved versions; powerpoint_restore puts one back
  powerpoint_slideshow start, step through, black out and end the real slide show (the Rehearse view)
  powerpoint_apply_direction  restyle the deck in a design direction (theme colours + fonts)
  powerpoint_set_theme change theme colour slots and fonts
  powerpoint_layout_variants  how to make hidden layout variants of a slide; powerpoint_choose_variant /
                       powerpoint_discard_variants keep one or drop them
  powerpoint_resume    re-arm the server after an unexpected error halted it
Tools only the view calls: slide_state, slide_image, deck_outline, slide_thumbs, deck_lint, history_thumb,
design_previews, theme_info, layout_variants.

Every change made through these tools is preceded by a saved version, so it can be undone from the History view.

Errors follow scripts/kShared.py: expected states (no deck open, file not found, not Windows) come back as tool
errors with a plain message; the first UNEXPECTED error is reported to kS.GlobalErrorHandler and halts the server.
While halted every tool answers "PowerPoint Live halted ..." until a person presses Resume (powerpoint_resume).
"""
import asyncio
import base64
import contextlib
import io
import os
import re
import shutil
import sys
import tempfile
import time
import traceback
from concurrent.futures import ThreadPoolExecutor

from mcp.server.fastmcp import FastMCP

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(HERE), "scripts")
if SCRIPTS not in sys.path:
    sys.path.insert(0, SCRIPTS)

from kShared import ToolInputException, ToolReportableException, kErrorReport, kS, kToolException  # noqa: E402

kErrorReport.HoldForView = True  # a server never exits: each error report waits for the person's Yes/No
kS.InstallUnhandledExceptionCapture()  # errors that escape every guarded method are reported too

URI = "ui://powerpoint-live/slide-view"
MIME = "text/html;profile=mcp-app"
UI = {"ui": {"resourceUri": URI}, "ui/resourceUri": URI}
APP_ONLY = {"ui": {"resourceUri": URI, "visibility": ["app"]}, "ui/resourceUri": URI}
RAMP = [0, 0.35, -0.3, 0.6, -0.5, 0.75]
LINE_CHARTS = {4, 63, 64, 65, 66, 67}
FIXABLE = {"unused_placeholder", "body_below_floor", "text_overflow", "picture_stretched", "a11y_missing_alt_text",
           "chart_default_palette", "chart_legend_steals_plot", "chart_accounting_zero_dash"}
MAX_VERSIONS = 25
NO_TEXT_TYPES = {7, 10, 12, 16}   # msoEmbeddedOLEObject, msoLinkedOLEObject, msoOLEControlObject, msoMedia
GOTO_VIEWS = {1, 3, 9}            # ppViewSlide, ppViewNotesPage, ppViewNormal: the views that can GotoSlide
SECTION_ACTIONS = ("add", "rename", "delete")
SHOW_ACTIONS = ("start", "next", "previous", "black", "end", "state")
SHOW_STATES = {1: "running", 2: "paused", 3: "black", 4: "white", 5: "done"}   # PpSlideShowState
WPM = 130                         # speaking rate behind every talk-length estimate (reference/CONTENT.md)
# A timing marker opens the notes (reference/PRESENTING.md): [15 sec], [2 min], [1:30]
MARKER = re.compile(r"^\s*\[\s*(?:(\d+)\s*:\s*(\d{1,2})|(\d+(?:\.\d+)?)\s*(s|secs?|seconds?|m|mins?|minutes?))\s*\]",
                    re.IGNORECASE)
MAX_PREVIEWS = 6
DEFAULT_DIRECTIONS = ["clean-corporate", "editorial-serif", "dark-stage", "nordic-calm", "bold-signal", "sunset-warm"]
DIRECTION_SLOTS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6")
SLOT_INDEX = {"dk1": 1, "lt1": 2, "dk2": 3, "lt2": 4, "accent1": 5, "accent2": 6, "accent3": 7, "accent4": 8,
              "accent5": 9, "accent6": 10, "hlink": 11, "folHlink": 12}   # MsoThemeColorSchemeIndex
HEX = re.compile(r"^#?([0-9A-Fa-f]{6})$")
TEXT_PAIRS = [("dk1", "lt1"), ("dk2", "lt1"), ("dk1", "lt2"), ("dk2", "lt2"), ("lt1", "dk1"), ("lt1", "dk2"),
              ("hlink", "lt1")]
GRAPHIC_PAIRS = [(f"accent{I}", "lt1") for I in range(1, 7)]
VARIANT_TAG = "PPTLIVE-VARIANT-OF"     # Slide.Tags: the SlideID this hidden slide is a layout variant of
VARIANT_NAME_TAG = "PPTLIVE-VARIANT-NAME"
PP_SAVE_AS_PPTX = 24                  # ppSaveAsOpenXMLPresentation: python-pptx reads it even when the deck is .pptm

mcp = FastMCP("powerpoint-live")


# ---------------------------------------------------------------- shape maps and change highlights

class kChangeTracker:
    """What each slide looked like when last seen, and which shapes changed since."""

    def __init__(self, Live):
        try:
            self.Live = Live
            self.Seen = {}      # SlideID -> shape map last seen
            self.Changes = {}   # SlideID -> {"stamp", "boxes"}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kChangeTracker.__init__")

    def Clear(self):
        """Forget everything (another deck, or a restore)."""
        if kS.ErrorMode:
            return None
        try:
            self.Seen.clear()
            self.Changes.clear()
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kChangeTracker.Clear")
            return None

    def ShapeText(self, Shape):
        """The text of a shape, or "" for shapes that cannot hold text."""
        if kS.ErrorMode:
            return ""
        try:
            if Shape.Type in NO_TEXT_TYPES or not Shape.HasTextFrame:
                return ""
            return Shape.TextFrame.TextRange.Text if Shape.TextFrame.HasText else ""
        except Exception as e:
            kS.GlobalErrorHandler(e, "kChangeTracker.ShapeText")
            return ""

    def ShapeMap(self, Slide):
        """Top-level shapes of a slide: id -> (name, box, text hash)."""
        if kS.ErrorMode:
            return {}
        try:
            Out = {}
            for Shape in Slide.Shapes:
                Text = self.ShapeText(Shape)
                Out[Shape.Id] = {"name": Shape.Name,
                                 "box": [round(Shape.Left, 1), round(Shape.Top, 1), round(Shape.Width, 1),
                                         round(Shape.Height, 1)], "text": hash(Text), "snippet": Text[:80]}
            return Out
        except Exception as e:
            kS.GlobalErrorHandler(e, "kChangeTracker.ShapeMap")
            return {}

    def Signature(self, Prs, N):
        """Cheap fingerprint of slide n. Changes when anyone edits the slide."""
        if kS.ErrorMode:
            return "halted"
        try:
            if not N:
                return "empty"
            Slide = Prs.Slides(N)
            Map = self.ShapeMap(Slide)
            return str(hash((Slide.SlideID, bool(Slide.SlideShowTransition.Hidden), self.Live.ThemeEpoch,
                             tuple((K, tuple(V["box"]), V["text"]) for K, V in Map.items()))))
        except Exception as e:
            kS.GlobalErrorHandler(e, "kChangeTracker.Signature")
            return "halted"

    def Diff(self, Before, After):
        """Boxes of the shapes added, changed or removed between two shape maps."""
        if kS.ErrorMode:
            return []
        try:
            Boxes = []
            for K, V in After.items():
                Old = Before.get(K)
                if Old is None:
                    Boxes.append({"id": K, "name": V["name"], "kind": "added", "box": V["box"]})
                elif Old["box"] != V["box"] or Old["text"] != V["text"]:
                    Boxes.append({"id": K, "name": V["name"], "kind": "changed", "box": V["box"]})
            for K, V in Before.items():
                if K not in After:
                    Boxes.append({"id": K, "name": V["name"], "kind": "removed", "box": V["box"]})
            return Boxes
        except Exception as e:
            kS.GlobalErrorHandler(e, "kChangeTracker.Diff")
            return []

    def NoteChanges(self, Slide, Before=None):
        """Compare a slide with what was last seen (or `Before`) and remember what changed."""
        if kS.ErrorMode:
            return {}
        try:
            After = self.ShapeMap(Slide)
            Prev = Before if Before is not None else self.Seen.get(Slide.SlideID)
            if Prev is not None:
                Boxes = self.Diff(Prev, After)
                if Boxes:
                    self.Changes[Slide.SlideID] = {"stamp": time.time(), "boxes": Boxes}
            self.Seen[Slide.SlideID] = After
            return After
        except Exception as e:
            kS.GlobalErrorHandler(e, "kChangeTracker.NoteChanges")
            return {}

    def NoteNewSlide(self, Prs, Slide):
        """A slide that did not exist before a change: highlight the whole slide."""
        if kS.ErrorMode:
            return None
        try:
            self.Changes[Slide.SlideID] = {"stamp": time.time(), "boxes": [{
                "id": 0, "name": "new slide", "kind": "added",
                "box": [0, 0, Prs.PageSetup.SlideWidth, Prs.PageSetup.SlideHeight]}]}
            self.Seen[Slide.SlideID] = self.ShapeMap(Slide)
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kChangeTracker.NoteNewSlide")
            return None


# ---------------------------------------------------------------- versions (undo)

class kVersionHistory:
    """Copies of the deck saved before every change, newest last."""

    def __init__(self, Live):
        try:
            self.Live = Live
            self.Versions = []   # [{"v", "time", "label", "path", "thumb", "slide", "count"}]
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVersionHistory.__init__")

    def Find(self, V):
        """The version entry numbered V, or None."""
        if kS.ErrorMode:
            return None
        try:
            for Entry in self.Versions:
                if Entry["v"] == int(V):
                    return Entry
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVersionHistory.Find")
            return None

    def Snapshot(self, Label):
        """Save a copy of the deck as a version before a change. SaveCopyAs never rebinds the open deck."""
        if kS.ErrorMode:
            return 0
        try:
            Live = self.Live
            Prs = Live.Prs()
            V = (self.Versions[-1]["v"] + 1) if self.Versions else 1
            Ext = os.path.splitext(Prs.FullName)[1] or ".pptx"
            Path = os.path.join(Live.Work, f"v{V:03d}{Ext}")
            Prs.SaveCopyAs(Path)
            N = Live.Clamp(Prs, Live.Current or 1)
            Thumb = Live.Png(Prs.Slides(N), 240) if N else None
            self.Versions.append({"v": V, "time": time.time(), "label": Label, "path": Path, "thumb": Thumb,
                                  "slide": N, "count": Prs.Slides.Count})
            while len(self.Versions) > MAX_VERSIONS:
                Old = self.Versions.pop(0)
                if os.path.exists(Old["path"]):
                    os.remove(Old["path"])
            return V
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVersionHistory.Snapshot")
            return 0

    def List(self):
        """The versions, newest first, without paths or pictures."""
        if kS.ErrorMode:
            return {"versions": []}
        try:
            return {"versions": [{K: H[K] for K in ("v", "time", "label", "slide", "count")}
                                 for H in reversed(self.Versions)]}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVersionHistory.List")
            return {"versions": []}

    def Thumb(self, V):
        """The thumbnail saved with version V."""
        if kS.ErrorMode:
            return {"v": 0, "png": None}
        try:
            Entry = self.Find(V)
            return {"v": int(V), "png": Entry["thumb"] if Entry else None}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVersionHistory.Thumb")
            return {"v": 0, "png": None}

    def Restore(self, V):
        """Replace the deck file with version V and reopen it; the state left behind is saved first."""
        if kS.ErrorMode:
            return None
        try:
            Live = self.Live
            Entry = self.Find(V)
            if Entry is None:
                raise ToolInputException(f"no version {V}")
            Prs = Live.Prs()
            Path = Prs.FullName
            if not os.path.isabs(Path):
                raise ToolReportableException("the deck has never been saved; save it once, then versions can be "
                                              "restored")
            self.Snapshot(f"before restoring v{V}")
            WithWindow = Prs.Windows.Count > 0
            Prs.Close()                  # COM Close discards unsaved changes without a prompt: they are in the version
            shutil.copyfile(Entry["path"], Path)
            New = Live.App().Presentations.Open(Path, 0, 0, -1 if WithWindow else 0)
            Live.PrsName = New.FullName
            Live.Current = Live.Clamp(New, Entry["slide"] or 1)
            Live.Tracker.Seen.clear()
            Live.Version += 1
            Live.SyncWindow(New)
            return {**Live.Summary(New), "restored": int(V)}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVersionHistory.Restore")
            return None


# ---------------------------------------------------------------- changes (handler classes, one per kind)

class kMutation:
    """One change to the deck. kPowerPointLive.Mutate saves a version, calls Apply, then records what changed."""

    def __init__(self, Live, Label):
        try:
            self.Live = Live
            self.Label = Label
        except Exception as e:
            kS.GlobalErrorHandler(e, "kMutation.__init__")

    def Apply(self, Prs):
        """Make the change; return extra keys for the result."""
        if kS.ErrorMode:
            return {}
        try:
            return {}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kMutation.Apply")
            return {}


class kRunMutation(kMutation):
    """powerpoint_run: the user's (Claude's) Python against the deck."""

    def __init__(self, Live, Code, Slide, Label):
        try:
            super().__init__(Live, Label or "powerpoint_run")
            self.Code = Code
            self.Slide = Slide
            self.Prs = None
            self.Output = io.StringIO()
            self.Error = None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kRunMutation.__init__")

    def Goto(self, N):
        """goto(n) in the user's code: change the current slide."""
        if kS.ErrorMode:
            return None
        try:
            self.Live.Current = self.Live.Clamp(self.Prs, N)
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kRunMutation.Goto")
            return None

    def Apply(self, Prs):
        """Run the code with app, prs, slide, goto and win32com in scope; print() output is kept."""
        if kS.ErrorMode:
            return {}
        try:
            Live = self.Live
            self.Prs = Prs
            if self.Slide:
                Live.Current = Live.Clamp(Prs, self.Slide)
            import win32com.client
            Scope = {"app": Live.App(), "prs": Prs, "slide": Prs.Slides(Live.Current) if Live.Current else None,
                     "goto": self.Goto, "win32com": win32com}
            with contextlib.redirect_stdout(self.Output):
                try:
                    exec(compile(self.Code, "<powerpoint_run>", "exec"), Scope)
                except Exception:  # ERROR-SUPPRESSED-JUSTIFIED: the user's code failing is the tool's result, reported back to Claude
                    self.Error = traceback.format_exc(limit=3)
            return {}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kRunMutation.Apply")
            return {}


class kMoveMutation(kMutation):
    """powerpoint_move: slides before a slide, to the end of a section, or to the end of the deck."""

    def __init__(self, Live, SlideIds, BeforeId, Section):
        try:
            super().__init__(Live, "move slides")
            self.SlideIds = [int(I) for I in SlideIds]
            self.BeforeId = int(BeforeId or 0)
            self.Section = int(Section or 0)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kMoveMutation.__init__")

    def Apply(self, Prs):
        """Move each slide in deck order, so their order is kept; the current slide stays current."""
        if kS.ErrorMode:
            return {}
        try:
            Live = self.Live
            Sorter = Live.Sorter
            Cur = Prs.Slides(Live.Current).SlideID if Live.Current else None
            Ordered = sorted((Sorter.SlideById(Prs, I).SlideIndex, I) for I in self.SlideIds)
            Ids = [I for _, I in Ordered if I != self.BeforeId]
            for Sid in Ids:
                Sorter.MoveOne(Prs, Sorter.SlideById(Prs, Sid), self.BeforeId, self.Section)
            if Cur is not None:
                Live.Current = Sorter.SlideById(Prs, Cur).SlideIndex
            return {"moved": Ids}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kMoveMutation.Apply")
            return {}


class kSectionMutation(kMutation):
    """powerpoint_sections: add, rename or delete a section."""

    def __init__(self, Live, Action, Section, Name, BeforeId):
        try:
            super().__init__(Live, f"section {Action}")
            self.Action = Action
            self.Section = int(Section or 0)
            self.Name = Name
            self.BeforeId = int(BeforeId or 0)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSectionMutation.__init__")

    def Apply(self, Prs):
        """Edit the sections; deleting keeps the slides (they join the section before)."""
        if kS.ErrorMode:
            return {}
        try:
            Sp = Prs.SectionProperties
            if self.Action == "add":
                Idx = self.Live.Sorter.SlideById(Prs, self.BeforeId).SlideIndex if self.BeforeId \
                    else Prs.Slides.Count + 1
                if Idx > Prs.Slides.Count:
                    Sp.AddSection(Sp.Count + 1, self.Name or "New section")
                else:
                    Sp.AddBeforeSlide(Idx, self.Name or "New section")
            elif self.Action == "rename":
                Sp.Rename(self.Section, self.Name)
            else:
                Sp.Delete(self.Section, False)
            return {}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSectionMutation.Apply")
            return {}


class kHideMutation(kMutation):
    """powerpoint_hide: hide or unhide slides in the slide show."""

    def __init__(self, Live, SlideIds, Hidden):
        try:
            super().__init__(Live, "hide slides" if Hidden else "unhide slides")
            self.SlideIds = [int(I) for I in SlideIds]
            self.Hidden = Hidden
        except Exception as e:
            kS.GlobalErrorHandler(e, "kHideMutation.__init__")

    def Apply(self, Prs):
        """Set SlideShowTransition.Hidden on each slide."""
        if kS.ErrorMode:
            return {}
        try:
            for Sid in self.SlideIds:
                self.Live.Sorter.SlideById(Prs, Sid).SlideShowTransition.Hidden = -1 if self.Hidden else 0
            return {}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kHideMutation.Apply")
            return {}


class kFixMutation(kMutation):
    """powerpoint_fix: the automatic fix for one lint finding on one shape."""

    def __init__(self, Live, SlideId, ShapeId, Code, Floor):
        try:
            super().__init__(Live, f"fix {Code}")
            self.SlideId = int(SlideId)
            self.ShapeId = int(ShapeId)
            self.Code = Code
            self.Floor = Floor
        except Exception as e:
            kS.GlobalErrorHandler(e, "kFixMutation.__init__")

    def Apply(self, Prs):
        """Find the shape (it may be gone) and fix it; its slide becomes current."""
        if kS.ErrorMode:
            return {}
        try:
            Live = self.Live
            Slide = Live.Sorter.SlideById(Prs, self.SlideId)
            Shape = Live.FindShape(Slide, ShapeId=self.ShapeId)
            if Shape is None:
                raise ToolReportableException("that shape is gone - re-run the lint")
            Done = Live.Lint.FixShape(Shape, self.Code, self.Floor)
            Live.Current = Slide.SlideIndex
            return {"fixed": self.Code, "result": Done}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kFixMutation.Apply")
            return {}


class kThemeMutation(kMutation):
    """powerpoint_set_theme / powerpoint_apply_direction: theme colours and fonts on every slide master."""

    def __init__(self, Live, Colors, Major, Minor, Label):
        try:
            super().__init__(Live, Label)
            self.Colors = dict(Colors or {})
            self.Major = Major or ""
            self.Minor = Minor or ""
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeMutation.__init__")

    def Apply(self, Prs):
        """Write the colours and fonts through COM; every slide's signature changes so pictures are redrawn."""
        if kS.ErrorMode:
            return {}
        try:
            self.Live.ThemeEditor.SetLive(Prs, self.Colors, self.Major, self.Minor)
            self.Live.ThemeEpoch += 1
            return {"theme": {"colors": self.Colors, "major_font": self.Major, "minor_font": self.Minor}}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeMutation.Apply")
            return {}


class kChooseVariantMutation(kMutation):
    """powerpoint_choose_variant: the chosen variant takes the original's place; the rest are deleted."""

    def __init__(self, Live, SlideId, VariantId):
        try:
            super().__init__(Live, "use layout variant")
            self.SlideId = int(SlideId)
            self.VariantId = int(VariantId)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kChooseVariantMutation.__init__")

    def Apply(self, Prs):
        """Unhide (as the original was), untag, move into place, delete the original and the other variants."""
        if kS.ErrorMode:
            return {}
        try:
            Live = self.Live
            Sorter, Variants = Live.Sorter, Live.Variants
            Orig = Sorter.SlideById(Prs, self.SlideId)
            Chosen = Sorter.SlideById(Prs, self.VariantId)
            Others = [S.SlideID for S in Variants.Of(Prs, self.SlideId) if S.SlideID != self.VariantId]
            Chosen.SlideShowTransition.Hidden = Orig.SlideShowTransition.Hidden
            Variants.Untag(Chosen)
            Sorter.MoveOne(Prs, Chosen, self.SlideId, 0)
            for Sid in Others:
                Sorter.SlideById(Prs, Sid).Delete()
            Orig.Delete()
            Live.Current = Sorter.SlideById(Prs, self.VariantId).SlideIndex
            return {"chosen": self.VariantId, "replaced": self.SlideId, "discarded": Others}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kChooseVariantMutation.Apply")
            return {}


class kDiscardVariantsMutation(kMutation):
    """powerpoint_discard_variants: delete every variant of one slide."""

    def __init__(self, Live, SlideId):
        try:
            super().__init__(Live, "discard layout variants")
            self.SlideId = int(SlideId)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDiscardVariantsMutation.__init__")

    def Apply(self, Prs):
        """Delete the variants; the original becomes the current slide."""
        if kS.ErrorMode:
            return {}
        try:
            Live = self.Live
            Ids = [S.SlideID for S in Live.Variants.Of(Prs, self.SlideId)]
            for Sid in Ids:
                Live.Sorter.SlideById(Prs, Sid).Delete()
            Live.Current = Live.Sorter.SlideById(Prs, self.SlideId).SlideIndex
            return {"discarded": Ids}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDiscardVariantsMutation.Apply")
            return {}


# ---------------------------------------------------------------- sorter, outline, thumbnails

class kSorter:
    """Outline, sections, thumbnails and slide moves."""

    def __init__(self, Live):
        try:
            self.Live = Live
            self.Thumbs = {}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSorter.__init__")

    def SlidesById(self, Prs):
        """SlideID -> slide, read by walking the deck (FindBySlideID raises for a deleted slide)."""
        if kS.ErrorMode:
            return {}
        try:
            return {Prs.Slides(I).SlideID: Prs.Slides(I) for I in range(1, Prs.Slides.Count + 1)}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSorter.SlidesById")
            return {}

    def SlideById(self, Prs, SlideId):
        """The slide with this SlideID; a missing one is the caller's mistake."""
        if kS.ErrorMode:
            return None
        try:
            Slide = self.SlidesById(Prs).get(int(SlideId))
            if Slide is None:
                raise ToolInputException(f"no slide with SlideID {SlideId} - read deck_outline again")
            return Slide
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSorter.SlideById")
            return None

    def Title(self, Slide):
        """The text of the slide's title placeholder, or ""."""
        if kS.ErrorMode:
            return ""
        try:
            for Shape in Slide.Shapes:
                if Shape.Type == 14 and Shape.PlaceholderFormat.Type in (1, 3) and Shape.HasTextFrame \
                        and Shape.TextFrame.HasText:  # msoPlaceholder: title, ctrTitle
                    return Shape.TextFrame.TextRange.Text.strip()
            return ""
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSorter.Title")
            return ""

    def NotesText(self, Slide):
        """The speaker notes, or "" when the slide has no notes body."""
        if kS.ErrorMode:
            return ""
        try:
            if not Slide.HasNotesPage:
                return ""
            Placeholders = Slide.NotesPage.Shapes.Placeholders
            if Placeholders.Count < 2 or not Placeholders(2).HasTextFrame:
                return ""
            return Placeholders(2).TextFrame.TextRange.Text
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSorter.NotesText")
            return ""

    def Sections(self, Prs):
        """The deck's sections: index, name, first slide, slide count."""
        if kS.ErrorMode:
            return []
        try:
            Sp = Prs.SectionProperties
            return [{"index": I, "name": Sp.Name(I), "first": Sp.FirstSlide(I), "count": Sp.SlidesCount(I)}
                    for I in range(1, Sp.Count + 1)]
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSorter.Sections")
            return []

    def Outline(self):
        """Sections and slides (id, title, hidden, notes and visible words, talk time, variant of, change signature);
        talk time per section and in total counts the slides the show will play (not hidden ones)."""
        if kS.ErrorMode:
            return None
        try:
            Live = self.Live
            Prs = Live.Prs()
            Has = Prs.SectionProperties.Count > 0
            Sections = self.Sections(Prs) or [{"index": 0, "name": "", "first": 1, "count": Prs.Slides.Count}]
            Slides = []
            for N in range(1, Prs.Slides.Count + 1):
                Slide = Prs.Slides(N)
                Slides.append({"index": N, "id": Slide.SlideID, "section": Slide.sectionIndex if Has else 0,
                               "title": self.Title(Slide), "hidden": bool(Slide.SlideShowTransition.Hidden),
                               **Live.Talk.ForSlide(Slide), "variant_of": Live.Variants.VariantOf(Slide),
                               "signature": Live.Tracker.Signature(Prs, N)})
            for Sec in Sections:
                Sec["slides"] = [X["id"] for X in Slides if X["section"] == Sec["index"]]
                Sec["estimate_sec"] = sum(X["planned_sec"] for X in Slides
                                          if X["section"] == Sec["index"] and not X["hidden"])
            return {**Live.Summary(Prs), "sections": Sections, "slides": Slides, "has_sections": Has, "wpm": WPM,
                    "estimate_sec": sum(X["planned_sec"] for X in Slides if not X["hidden"])}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSorter.Outline")
            return None

    def ThumbsFor(self, Ids, Width):
        """Small PNGs of the given slides, cached by content; deleted slides are left out."""
        if kS.ErrorMode:
            return {"thumbs": []}
        try:
            Live = self.Live
            Prs = Live.Prs()
            ById = self.SlidesById(Prs)
            Out = []
            for Sid in Ids:
                Slide = ById.get(int(Sid))
                if Slide is None:  # deleted since the outline was read: nothing to draw
                    continue
                Sig = Live.Tracker.Signature(Prs, Slide.SlideIndex)
                Key = (Prs.FullName, Sid, Sig, Width)
                if Key not in self.Thumbs:
                    self.Thumbs[Key] = Live.Png(Slide, Width)
                    if len(self.Thumbs) > 400:
                        self.Thumbs.pop(next(iter(self.Thumbs)))
                Out.append({"id": Sid, "signature": Sig, "png": self.Thumbs[Key]})
            return {"thumbs": Out}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSorter.ThumbsFor")
            return {"thumbs": []}

    def MoveOne(self, Prs, Slide, BeforeId, Section):
        """Move one slide before BeforeId, to the end of Section, or to the end of the deck."""
        if kS.ErrorMode:
            return None
        try:
            Has = Prs.SectionProperties.Count > 0
            if BeforeId:
                if int(BeforeId) == Slide.SlideID:
                    return None
                B = self.SlideById(Prs, BeforeId)
                Bi, Bsec = B.SlideIndex, (B.sectionIndex if Has else 0)
                if Has and Prs.SectionProperties.FirstSlide(Bsec) == Bi:
                    Slide.MoveToSectionStart(Bsec)
                else:
                    Slide.MoveTo(Bi - 1 if Slide.SlideIndex < Bi else Bi)
            elif Has and Section:
                Sp = Prs.SectionProperties
                Sec = int(Section)
                if Sp.SlidesCount(Sec) == 0 or (Sp.SlidesCount(Sec) == 1 and Slide.sectionIndex == Sec):
                    Slide.MoveToSectionStart(Sec)
                else:
                    Last = Sp.FirstSlide(Sec) + Sp.SlidesCount(Sec) - 1
                    Slide.MoveTo(Last if Slide.SlideIndex <= Last else Last + 1)
                    if Slide.sectionIndex != Sec:  # landed on the far side of the boundary
                        Slide.MoveTo(Slide.SlideIndex - 1)
            else:
                Slide.MoveTo(Prs.Slides.Count)
            return None
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSorter.MoveOne")
            return None


# ---------------------------------------------------------------- lint + live fixes

class kLintBridge:
    """scripts/lint_deck.py on a copy of the live deck, with shape boxes; and the live fixes."""

    def __init__(self, Live):
        try:
            self.Live = Live
            self.Key = None
            self.Result = None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintBridge.__init__")

    def DeckKey(self, Prs):
        """Changes whenever any slide changes."""
        if kS.ErrorMode:
            return None
        try:
            Tracker = self.Live.Tracker
            return (Prs.FullName, Prs.Slides.Count,
                    tuple(Tracker.Signature(Prs, I) for I in range(1, Prs.Slides.Count + 1)))
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintBridge.DeckKey")
            return None

    def Lint(self):
        """Findings with slide ids and shape boxes, cached until the deck changes."""
        if kS.ErrorMode:
            return None
        try:
            Live = self.Live
            Prs = Live.Prs()
            Key = self.DeckKey(Prs)
            if self.Key == Key:
                return self.Result
            from lint_deck import kLintDeck
            Copy = os.path.join(Live.Work, "lint" + (os.path.splitext(Prs.FullName)[1] or ".pptx"))
            Prs.SaveCopyAs(Copy)
            _, Findings = kLintDeck.Lint(Copy, 18.0, 12)
            Out = []
            for I in Findings.Items:
                Item = {**I, "fixable": I["code"] in FIXABLE and bool(I["shape"]), "slide_id": None, "box": None}
                if I["slide"]:
                    Slide = Prs.Slides(I["slide"])
                    Item["slide_id"] = Slide.SlideID
                    Shape = Live.FindShape(Slide, I["shape"], I.get("shape_id")) if I["shape"] else None
                    if Shape is not None:
                        Item["shape_id"] = Shape.Id
                        Item["box"] = [round(Shape.Left, 1), round(Shape.Top, 1), round(Shape.Width, 1),
                                       round(Shape.Height, 1)]
                Out.append(Item)
            Counts = {}
            for I in Out:
                if I["slide_id"] and I["severity"] in ("error", "warn"):
                    Counts[I["slide_id"]] = Counts.get(I["slide_id"], 0) + 1
            self.Key, self.Result = Key, {"floor_pt": Findings.Floor, "findings": Out, "counts": Counts}
            return self.Result
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintBridge.Lint")
            return None

    def FloorPt(self):
        """The body-text floor of the last lint, or 18 pt scaled to the slide width."""
        if kS.ErrorMode:
            return 18.0
        try:
            return (self.Result or {}).get("floor_pt") or 18.0 * max(1.0, self.Live.Prs().PageSetup.SlideWidth / 960)
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintBridge.FloorPt")
            return 18.0

    def SeriesText(self, Series):
        """'name: cat value, ...' for one chart series, or None when its data cannot be read."""
        if kS.ErrorMode:
            return None
        try:
            try:
                Cats = [str(C) for C in (Series.XValues or [])]
                Vals = list(Series.Values or [])
            except Exception:  # ERROR-SUPPRESSED-JUSTIFIED: some chart kinds (stock, surface, linked data) expose no XValues/Values over COM; the alt text then names only the other series
                return None
            Pairs = ", ".join(f"{C} {V:g}" if isinstance(V, (int, float)) else f"{C} n/a" for C, V in zip(Cats, Vals))
            return f"{Series.Name}: {Pairs}"
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintBridge.SeriesText")
            return None

    def AltForChart(self, Chart):
        """Alt text for a chart, written from its title and data."""
        if kS.ErrorMode:
            return ""
        try:
            Parts = []
            for I in range(1, Chart.SeriesCollection().Count + 1):
                Text = self.SeriesText(Chart.SeriesCollection(I))
                if Text:
                    Parts.append(Text)
            Title = Chart.ChartTitle.Text + ". " if Chart.HasTitle else ""
            return f"Chart. {Title}" + "; ".join(Parts) + "."
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintBridge.AltForChart")
            return ""

    def FixShape(self, Shape, Code, Floor):
        """Apply the automatic fix for Code to Shape; returns what was done."""
        if kS.ErrorMode:
            return None
        try:
            if Code not in FIXABLE:
                raise ToolInputException(f"{Code} has no automatic fix - ask Claude")
            if Code == "unused_placeholder":
                Shape.Delete()
                return "deleted the empty placeholder"
            if Code == "body_below_floor":
                N = 0
                for P in Shape.TextFrame.TextRange.Paragraphs():
                    if 0 < P.Font.Size < Floor and len(P.Text.split()) >= 4:
                        P.Font.Size = Floor
                        N += 1
                return f"{N} paragraph(s) raised to {Floor:g} pt"
            if Code == "text_overflow":
                Tf, Tr = Shape.TextFrame, Shape.TextFrame.TextRange
                Room = Shape.Height - Tf.MarginTop - Tf.MarginBottom
                Steps = 0
                while Tr.BoundHeight > Room and Steps < 40:
                    Runs = [R for R in Tr.Runs() if R.Font.Size > Floor]
                    if not Runs:
                        break
                    for R in Runs:
                        R.Font.Size = max(Floor, R.Font.Size - 1)
                    Steps += 1
                Left = " (still too long - cut words)" if Tr.BoundHeight > Room else ""
                return f"text shrunk {Steps} pt to fit its box{Left}"
            if Code == "picture_stretched":
                L, T, W, H = Shape.Left, Shape.Top, Shape.Width, Shape.Height
                Shape.LockAspectRatio = 0
                Shape.ScaleHeight(1, -1)
                Shape.ScaleWidth(1, -1)
                Ratio = Shape.Width / Shape.Height
                Nw, Nh = (W, W / Ratio) if W / H < Ratio else (H * Ratio, H)
                Shape.Width, Shape.Height = Nw, Nh
                Shape.Left, Shape.Top = L + (W - Nw) / 2, T + (H - Nh) / 2
                Shape.LockAspectRatio = -1
                return "resized to its true proportions inside the same box"
            if Code == "a11y_missing_alt_text":
                if Shape.HasChart:
                    Shape.AlternativeText = self.AltForChart(Shape.Chart)
                    return "chart alt text written from its data"
                if Shape.HasTable:
                    Tb = Shape.Table
                    Head = [Tb.Cell(1, C).Shape.TextFrame.TextRange.Text for C in range(1, Tb.Columns.Count + 1)]
                    Shape.AlternativeText = f"Table: {', '.join(Head)}; {Tb.Rows.Count - 1} rows."
                    return "table alt text written from its header"
                raise ToolReportableException("only Claude can describe a picture - ask Claude")
            if not Shape.HasChart:
                raise ToolInputException(f"{Code}: not a chart")
            Chart = Shape.Chart
            if Code == "chart_default_palette":
                for I in range(1, Chart.SeriesCollection().Count + 1):
                    S = Chart.SeriesCollection(I)
                    Fmt = S.Format.Line if Chart.ChartType in LINE_CHARTS else S.Format.Fill
                    Fmt.ForeColor.ObjectThemeColor = 5  # msoThemeColorAccent1
                    Fmt.ForeColor.Brightness = RAMP[(I - 1) % len(RAMP)]
                return "series recoloured as shades of the theme accent"
            if Code == "chart_legend_steals_plot":
                Chart.Legend.Position = -4152  # xlLegendPositionRight
                Chart.Legend.IncludeInLayout = False
                return "legend moved to the right"
            Ax = Chart.Axes(2)  # chart_accounting_zero_dash
            Ax.TickLabels.NumberFormat = "$#,##0" if "$" in Ax.TickLabels.NumberFormat else "#,##0"
            return "Accounting axis format replaced"
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintBridge.FixShape")
            return None


# ---------------------------------------------------------------- talk length, slide show, designs, theme, variants

class kTalkTime:
    """Speaking time per slide: the timing marker in the notes, else an estimate from the words at WPM."""

    def __init__(self, Live):
        try:
            self.Live = Live
        except Exception as e:
            kS.GlobalErrorHandler(e, "kTalkTime.__init__")

    def MarkerSeconds(self, Notes):
        """Seconds of the timing marker that opens the notes ([15 sec], [2 min], [1:30]), or None."""
        if kS.ErrorMode:
            return None
        try:
            M = MARKER.match(Notes or "")
            if M is None:
                return None
            if M.group(1) is not None:
                return int(M.group(1)) * 60 + int(M.group(2))
            Value = float(M.group(3))
            return round(Value * 60 if M.group(4).lower().startswith("m") else Value)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kTalkTime.MarkerSeconds")
            return None

    def VisibleWords(self, Slide):
        """Words of text on the slide itself (group members included)."""
        if kS.ErrorMode:
            return 0
        try:
            Live = self.Live
            return sum(len(Live.Tracker.ShapeText(Shape).split()) for Shape in Live.Walk(Slide.Shapes))
        except Exception as e:
            kS.GlobalErrorHandler(e, "kTalkTime.VisibleWords")
            return 0

    def Seconds(self, Words):
        """Speaking time of Words words at WPM."""
        if kS.ErrorMode:
            return 0
        try:
            return round(Words * 60 / WPM)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kTalkTime.Seconds")
            return 0

    def ForSlide(self, Slide, Notes=None):
        """notes_words, visible_words, estimate_sec (from the notes, else the visible words), marker_sec and
        planned_sec (the marker when there is one, else the estimate)."""
        if kS.ErrorMode:
            return {}
        try:
            Notes = self.Live.Sorter.NotesText(Slide) if Notes is None else Notes
            NotesWords, Visible = len(Notes.split()), self.VisibleWords(Slide)
            Estimate = self.Seconds(NotesWords if NotesWords else Visible)
            Marker = self.MarkerSeconds(Notes)
            return {"notes_words": NotesWords, "visible_words": Visible, "estimate_sec": Estimate,
                    "estimate_from": "notes" if NotesWords else "slide", "marker_sec": Marker,
                    "planned_sec": Marker if Marker is not None else Estimate}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kTalkTime.ForSlide")
            return {}


class kSlideShow:
    """powerpoint_slideshow: drives the real slide show of the deck (SlideShowSettings.Run, SlideShowWindow.View)."""

    def __init__(self, Live):
        try:
            self.Live = Live
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlideShow.__init__")

    def Window(self, Prs):
        """The slide show window of this deck, or None when no show of it is running."""
        if kS.ErrorMode:
            return None
        try:
            for W in self.Live.App().SlideShowWindows:
                if W.Presentation.FullName == Prs.FullName:
                    return W
            return None
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlideShow.Window")
            return None

    def Require(self, Prs):
        """The running show's window; no show is an expected state the caller is told about."""
        if kS.ErrorMode:
            return None
        try:
            W = self.Window(Prs)
            if W is None:
                raise ToolReportableException("no slide show of this deck is running - call powerpoint_slideshow "
                                              "with action=start first")
            return W
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlideShow.Require")
            return None

    def Do(self, Action):
        """Run one action; every action answers with the show's state."""
        if kS.ErrorMode:
            return None
        try:
            if Action not in SHOW_ACTIONS:
                raise ToolInputException("action must be start, next, previous, black, end or state")
            Prs = self.Live.Prs()
            if Action == "start":
                self.Start(Prs)
            elif Action != "state":
                View = self.Require(Prs).View
                if Action != "end" and View.State == 5:
                    if Action == "previous":
                        View.Previous()   # from the end-of-show screen back onto the last slide
                    elif Action == "next":
                        raise ToolReportableException("the show is at its end - previous or end")
                    else:
                        raise ToolReportableException("the show is at its end - black needs a slide on screen")
                elif Action == "next":
                    View.Next()
                elif Action == "previous":
                    View.Previous()
                elif Action == "black":
                    View.State = 1 if View.State == 3 else 3   # ppSlideShowRunning <-> ppSlideShowBlackScreen
                else:
                    View.Exit()
            return {**self.State(Prs), "action": Action}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlideShow.Do")
            return None

    def Start(self, Prs):
        """SlideShowSettings.Run, then go to the current slide; a running show is reused."""
        if kS.ErrorMode:
            return None
        try:
            if Prs.Slides.Count == 0:
                raise ToolReportableException("the deck has no slides to show")
            W = self.Window(Prs)
            if W is None:
                W = Prs.SlideShowSettings.Run()
            N = self.Live.Clamp(Prs, self.Live.Current or 1)
            if W.View.State == 5 or W.View.Slide.SlideIndex != N:
                W.View.GotoSlide(N)
            return None
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlideShow.Start")
            return None

    def State(self, Prs):
        """Running or not, the slide on screen (the current slide when no show runs), its notes and timing,
        the next slide the show will reach and the planned time before this slide."""
        if kS.ErrorMode:
            return {}
        try:
            Live = self.Live
            W = self.Window(Prs)
            Show = SHOW_STATES.get(W.View.State, "running") if W is not None else "stopped"
            if W is not None and Show != "done":
                Live.Current = W.View.Slide.SlideIndex   # the Slide view follows the show
            N = Live.Clamp(Prs, Live.Current or 1)
            Out = {**Live.Summary(Prs), "running": W is not None, "state": Show, "slide": N, "slide_id": None,
                   "notes": "", "next_slide": 0, "next_slide_id": None, "planned_before_sec": 0, "planned_total_sec": 0}
            if not N:
                return Out
            Before = Total = 0
            for I in range(1, Prs.Slides.Count + 1):
                Slide = Prs.Slides(I)
                Hidden = bool(Slide.SlideShowTransition.Hidden)
                if I == N:
                    Notes = Live.Sorter.NotesText(Slide)
                    Out.update({"slide_id": Slide.SlideID, "notes": Notes, **Live.Talk.ForSlide(Slide, Notes)})
                if Hidden and I != N:
                    continue
                Planned = Live.Talk.ForSlide(Slide)["planned_sec"]
                Total += Planned
                if I < N:
                    Before += Planned
                elif I > N and not Out["next_slide"]:
                    Out["next_slide"], Out["next_slide_id"] = I, Slide.SlideID
            Out["planned_before_sec"], Out["planned_total_sec"] = Before, Total
            return Out
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlideShow.State")
            return {}


class kDesignStudio:
    """Design directions (scripts/directions.json, applied as scripts/build_deck.py kDeckDesign does): previews
    rendered from a copy of the deck, and the live theme change."""

    def __init__(self, Live):
        try:
            self.Live = Live
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDesignStudio.__init__")

    def All(self):
        """Every direction in directions.json."""
        if kS.ErrorMode:
            return []
        try:
            from build_deck import kSpecSchema
            return kSpecSchema.Directions()
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDesignStudio.All")
            return []

    def Resolve(self, Names):
        """The directions called Names (empty = six contrasting defaults); unknown names are the caller's mistake."""
        if kS.ErrorMode:
            return []
        try:
            ById = {D["id"]: D for D in self.All()}
            Names = [N for N in (Names or []) if N] or [N for N in DEFAULT_DIRECTIONS if N in ById]
            if len(Names) > MAX_PREVIEWS:
                raise ToolInputException(f"at most {MAX_PREVIEWS} directions at a time")
            Unknown = [N for N in Names if N not in ById]
            if Unknown:
                raise ToolInputException(f"unknown direction(s) {', '.join(Unknown)}; known: {', '.join(ById)}")
            return [ById[N] for N in Names]
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDesignStudio.Resolve")
            return []

    def Slots(self, D):
        """The theme colours and fonts a direction gives a deck - kDeckDesign.ApplyDirection on a blank deck,
        read back with kTheme, so they are exactly what build_deck writes."""
        if kS.ErrorMode:
            return {}
        try:
            from pptx import Presentation
            from build_deck import kDeckDesign
            from _theme import kTheme
            Blank = Presentation()
            kDeckDesign.ApplyDirection(Blank, D)
            Theme = kTheme(Blank.slide_master)
            return {"colors": {K: Theme.Colours[K] for K in DIRECTION_SLOTS}, "major_font": Theme.Major,
                    "minor_font": Theme.Minor}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDesignStudio.Slots")
            return {}

    def SpreadTheme(self, Pptx):
        """Give every slide master the first master's theme (ApplyDirection writes only the first)."""
        if kS.ErrorMode:
            return None
        try:
            from pptx.opc.constants import RELATIONSHIP_TYPE as RT
            Blob = Pptx.slide_master.part.part_related_by(RT.THEME).blob
            for Master in list(Pptx.slide_masters)[1:]:
                Master.part.part_related_by(RT.THEME)._blob = Blob
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDesignStudio.SpreadTheme")
            return None

    def RenderWith(self, Src, D, N, Width):
        """Slide N of the copy Src restyled in direction D: python-pptx writes the theme into another copy,
        a windowless read-only PowerPoint opens it and Slide.Export draws the slide. The open deck is untouched."""
        if kS.ErrorMode:
            return None
        try:
            from pptx import Presentation
            from build_deck import kDeckDesign
            Pptx = Presentation(Src)
            kDeckDesign.ApplyDirection(Pptx, D)
            self.SpreadTheme(Pptx)
            Path = os.path.join(self.Live.Work, f"design-{D['id']}.pptx")
            Pptx.save(Path)
            Copy = self.Live.App().Presentations.Open(Path, -1, 0, 0)   # read-only, untitled no, no window
            try:
                return self.Live.Png(Copy.Slides(N), Width)
            finally:
                Copy.Close()
                os.remove(Path)
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDesignStudio.RenderWith")
            return None

    def Previews(self, Slide, Names, Width):
        """design_previews: the slide in up to six directions, plus the list of every direction."""
        if kS.ErrorMode:
            return None
        try:
            Live = self.Live
            Ds = self.Resolve(Names)
            Prs = Live.Prs()
            N = Live.Clamp(Prs, Slide or Live.Current or 1)
            if not N:
                raise ToolReportableException("the deck has no slides to preview")
            Src = os.path.join(Live.Work, "design-src.pptx")
            Prs.SaveCopyAs(Src, PP_SAVE_AS_PPTX)
            Out = []
            for D in Ds:
                Out.append({"direction": D["id"], "mood": D.get("mood", ""), "tone": D.get("tone", ""),
                            "heading": D["heading"], "body": D["body"], "accent": D["accent"],
                            "background": D["background"], "text": D["text"], "png": self.RenderWith(Src, D, N, Width)})
            os.remove(Src)
            return {"slide": N, "slide_id": Prs.Slides(N).SlideID, "previews": Out,
                    "directions": [{"id": D["id"], "mood": D.get("mood", ""), "tone": D.get("tone", "")}
                                   for D in self.All()]}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDesignStudio.Previews")
            return None

    def ApplyLive(self, Name):
        """powerpoint_apply_direction: the direction's colours and fonts on the whole live deck, after a version."""
        if kS.ErrorMode:
            return None
        try:
            if not Name:
                raise ToolInputException("name a direction - see design_previews or scripts/directions.json")
            D = self.Resolve([Name])[0]
            Slots = self.Slots(D)
            Res = self.Live.Mutate(kThemeMutation(self.Live, Slots["colors"], Slots["major_font"], Slots["minor_font"],
                                                  f"design direction {Name}"))
            return {**(Res or {}), "direction": Name}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDesignStudio.ApplyLive")
            return None


class kThemeEditor:
    """theme_info and powerpoint_set_theme: theme colour slots, fonts and contrast pairs."""

    def __init__(self, Live):
        try:
            self.Live = Live
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeEditor.__init__")

    def Bgr(self, Hex):
        """RRGGBB to the BGR integer COM's .RGB expects."""
        if kS.ErrorMode:
            return 0
        try:
            return int(Hex[4:6] + Hex[2:4] + Hex[0:2], 16)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeEditor.Bgr")
            return 0

    def Clean(self, Colors, Major, Minor):
        """The colours as {slot: RRGGBB}; unknown slots, bad hex or nothing to change are the caller's mistake."""
        if kS.ErrorMode:
            return {}
        try:
            Out = {}
            for Slot, Value in (Colors or {}).items():
                if Slot not in SLOT_INDEX:
                    raise ToolInputException(f"unknown theme slot {Slot}; slots: {', '.join(SLOT_INDEX)}")
                M = HEX.match(str(Value).strip())
                if M is None:
                    raise ToolInputException(f"{Slot}: {Value!r} is not a RRGGBB hex colour")
                Out[Slot] = M.group(1).upper()
            if not Out and not (Major or "").strip() and not (Minor or "").strip():
                raise ToolInputException("nothing to change - give colors, major_font or minor_font")
            return Out
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeEditor.Clean")
            return {}

    def SetLive(self, Prs, Colors, Major, Minor):
        """Theme colours (ThemeColorScheme.Colors(i).RGB, BGR) and Latin fonts on every design's slide master."""
        if kS.ErrorMode:
            return None
        try:
            for I in range(1, Prs.Designs.Count + 1):
                Theme = Prs.Designs(I).SlideMaster.Theme
                for Slot, Hex in Colors.items():
                    Theme.ThemeColorScheme.Colors(SLOT_INDEX[Slot]).RGB = self.Bgr(Hex)
                if Major:
                    Theme.ThemeFontScheme.MajorFont.Item(1).Name = Major   # msoThemeLatin
                if Minor:
                    Theme.ThemeFontScheme.MinorFont.Item(1).Name = Minor
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeEditor.SetLive")
            return None

    def Pairs(self, Colours):
        """Contrast of text on backgrounds (4.5:1) and accents on the background (3:1, large text and graphics)."""
        if kS.ErrorMode:
            return []
        try:
            from _rules import kRules
            Out = []
            for Pairs, Need, Kind in ((TEXT_PAIRS, 4.5, "text"), (GRAPHIC_PAIRS, 3.0, "graphics")):
                for Fg, Bg in Pairs:
                    if Fg in Colours and Bg in Colours:
                        Ratio = kRules.ContrastRatio(Colours[Fg], Colours[Bg])
                        Out.append({"fg": Fg, "bg": Bg, "ratio": round(Ratio, 2), "min": Need, "kind": Kind,
                                    "pass": Ratio >= Need})
            return Out
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeEditor.Pairs")
            return []

    def Info(self):
        """theme_info: the current slide's theme read with kTheme from a Save As copy, and its contrast pairs."""
        if kS.ErrorMode:
            return None
        try:
            from pptx import Presentation
            from _theme import kTheme
            Live = self.Live
            Prs = Live.Prs()
            Copy = os.path.join(Live.Work, "theme.pptx")
            Prs.SaveCopyAs(Copy, PP_SAVE_AS_PPTX)
            Pptx = Presentation(Copy)
            N = Live.Clamp(Prs, Live.Current or 1)
            Master = Pptx.slides[N - 1].slide_layout.slide_master if N else Pptx.slide_master
            Theme = kTheme(Master)
            Colours = dict(Theme.Colours)
            os.remove(Copy)
            return {**Live.Summary(Prs), "colors": Colours, "major_font": Theme.Major or "",
                    "minor_font": Theme.Minor or "", "pairs": self.Pairs(Colours), "masters": Prs.Designs.Count}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeEditor.Info")
            return None

    def Set(self, Colors, Major, Minor):
        """powerpoint_set_theme: checked first, then a version, then the change."""
        if kS.ErrorMode:
            return None
        try:
            Clean = self.Clean(Colors, Major, Minor)
            return self.Live.Mutate(kThemeMutation(self.Live, Clean, (Major or "").strip(), (Minor or "").strip(),
                                                   "theme colours and fonts"))
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeEditor.Set")
            return None


class kVariants:
    """Layout variants: hidden copies of a slide, right after it, tagged PPTLIVE-VARIANT-OF=<SlideID>."""

    def __init__(self, Live):
        try:
            self.Live = Live
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVariants.__init__")

    def VariantOf(self, Slide):
        """The SlideID this slide is a variant of, or 0."""
        if kS.ErrorMode:
            return 0
        try:
            Value = str(Slide.Tags.Item(VARIANT_TAG) or "").strip()
            return int(Value) if Value.isdigit() else 0
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVariants.VariantOf")
            return 0

    def Of(self, Prs, SlideId):
        """The variants of one slide, in deck order."""
        if kS.ErrorMode:
            return []
        try:
            return [Prs.Slides(I) for I in range(1, Prs.Slides.Count + 1)
                    if self.VariantOf(Prs.Slides(I)) == int(SlideId)]
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVariants.Of")
            return []

    def Untag(self, Slide):
        """Remove the variant tags from a slide."""
        if kS.ErrorMode:
            return None
        try:
            for Tag in (VARIANT_TAG, VARIANT_NAME_TAG):
                if Slide.Tags.Item(Tag):
                    Slide.Tags.Delete(Tag)
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVariants.Untag")
            return None

    def Require(self, Prs, SlideId, VariantId=0):
        """The slide's variants; none (or a VariantId that is not one of them) is reported before any version."""
        if kS.ErrorMode:
            return []
        try:
            self.Live.Sorter.SlideById(Prs, SlideId)
            Found = self.Of(Prs, SlideId)
            if not Found:
                raise ToolReportableException(f"slide {SlideId} has no layout variants - see "
                                              "powerpoint_layout_variants")
            if VariantId and int(VariantId) not in [S.SlideID for S in Found]:
                raise ToolInputException(f"slide {VariantId} is not a variant of slide {SlideId} - call "
                                         "layout_variants")
            return Found
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVariants.Require")
            return []

    def Guide(self, Slide, Count):
        """powerpoint_layout_variants: what Claude needs to make the variants itself with powerpoint_run."""
        if kS.ErrorMode:
            return None
        try:
            Live = self.Live
            Prs = Live.Prs()
            N = Live.Clamp(Prs, Slide or Live.Current or 1)
            if not N:
                raise ToolReportableException("the deck has no slides")
            Live.Current = N
            Orig = Prs.Slides(N)
            if self.VariantOf(Orig):
                raise ToolInputException(f"slide {N} is itself a variant - ask for variants of slide "
                                         f"{self.VariantOf(Orig)}")
            Count = max(1, min(int(Count or 3), MAX_PREVIEWS))
            Existing = [{"id": S.SlideID, "index": S.SlideIndex, "name": S.Tags.Item(VARIANT_NAME_TAG)}
                        for S in self.Of(Prs, Orig.SlideID)]
            Code = (f"orig = prs.Slides({N})\n"
                    f"for k, name in enumerate([...{Count} short layout names...]):\n"
                    "    v = orig.Duplicate().Item(1)\n"
                    "    v.MoveTo(orig.SlideIndex + 1 + k)\n"
                    "    v.SlideShowTransition.Hidden = -1\n"
                    f"    v.Tags.Add(\"{VARIANT_TAG}\", str(orig.SlideID))\n"
                    f"    v.Tags.Add(\"{VARIANT_NAME_TAG}\", name)\n"
                    "    # now rearrange v's shapes into that layout (keep every word and number)\n")
            return {**Live.Summary(Prs), "slide": N, "slide_id": Orig.SlideID, "count": Count, "existing": Existing,
                    "how": (f"Make {Count} layout variants of slide {N} yourself with ONE powerpoint_run (label "
                            "\"layout variants\"): duplicate the slide, keep each copy HIDDEN, right after the "
                            f"original, tagged {VARIANT_TAG}=<original SlideID> and {VARIANT_NAME_TAG}=<a short "
                            "name>; then rearrange each copy into a genuinely different layout (reference/LAYOUT.md). "
                            "Never change the original. The view shows them in a Variants strip; the user picks one "
                            "(powerpoint_choose_variant) or discards them (powerpoint_discard_variants)."),
                    "code": Code}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVariants.Guide")
            return None

    def List(self, SlideId, Width):
        """layout_variants: the original and its variants with thumbnails."""
        if kS.ErrorMode:
            return None
        try:
            Live = self.Live
            Prs = Live.Prs()
            Orig = Live.Sorter.SlideById(Prs, SlideId)
            Found = self.Of(Prs, Orig.SlideID)
            Pngs = {T["id"]: T["png"] for T in
                    Live.Sorter.ThumbsFor([Orig.SlideID] + [S.SlideID for S in Found], Width)["thumbs"]}
            return {"slide_id": Orig.SlideID, "slide": Orig.SlideIndex,
                    "original": {"id": Orig.SlideID, "index": Orig.SlideIndex, "png": Pngs.get(Orig.SlideID)},
                    "variants": [{"id": S.SlideID, "index": S.SlideIndex,
                                  "name": S.Tags.Item(VARIANT_NAME_TAG) or f"Variant {K + 1}",
                                  "png": Pngs.get(S.SlideID)} for K, S in enumerate(Found)]}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVariants.List")
            return None

    def Choose(self, SlideId, VariantId):
        """powerpoint_choose_variant, checked before a version is saved."""
        if kS.ErrorMode:
            return None
        try:
            self.Require(self.Live.Prs(), SlideId, VariantId or -1)
            return self.Live.Mutate(kChooseVariantMutation(self.Live, SlideId, VariantId))
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVariants.Choose")
            return None

    def Discard(self, SlideId):
        """powerpoint_discard_variants, checked before a version is saved."""
        if kS.ErrorMode:
            return None
        try:
            self.Require(self.Live.Prs(), SlideId)
            return self.Live.Mutate(kDiscardVariantsMutation(self.Live, SlideId))
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kVariants.Discard")
            return None


# ---------------------------------------------------------------- the live deck (one COM thread)

class kPowerPointLive:
    """Owns the COM thread, the deck binding, the current slide and the change version the view polls."""

    def __init__(self):
        try:
            self.PrsName = None   # FullName of the deck we work on
            self.Current = 1
            self.Version = 0
            self.ThemeEpoch = 0   # bumped by every theme change: slide signatures (and so pictures) follow it
            self.Work = tempfile.mkdtemp(prefix="pptlive-")
            self.Pool = ThreadPoolExecutor(max_workers=1,
                                           initializer=self.ComInit if sys.platform == "win32" else None)
            self.Tracker = kChangeTracker(self)
            self.History = kVersionHistory(self)
            self.Lint = kLintBridge(self)
            self.Sorter = kSorter(self)
            self.Talk = kTalkTime(self)
            self.SlideShow = kSlideShow(self)
            self.Design = kDesignStudio(self)
            self.ThemeEditor = kThemeEditor(self)
            self.Variants = kVariants(self)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.__init__")

    def ComInit(self):
        """Initialise COM on the pool's one thread."""
        if kS.ErrorMode:
            return None
        try:
            import pythoncom
            pythoncom.CoInitialize()
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.ComInit")
            return None

    # DOCUMENTED EXCEPTION: builds the halt message; it must work exactly when kS.ErrorMode is set
    def HaltedError(self):
        First = kS.FirstError() or {}
        Ask = ""
        if kErrorReport.Pending and kErrorReport.EffectiveUrl():
            Ask = (f" ERROR-REPORT-PENDING: ask the user \"{kErrorReport.Question}\" (yes/no) - they can also answer"
                   " in the view. Only if they say yes, call powerpoint_send_error_report with send=true; if they say"
                   " no, call it with send=false.")
        return ToolReportableException(
            f"PowerPoint Live halted after an error in {First.get('location', '?')}: "
            f"{First.get('type', 'Error')}: {First.get('message', '')}. "
            "Press Resume in the view or call powerpoint_resume." + Ask)

    # DOCUMENTED EXCEPTION: the one re-arm (a person's Resume); it must work exactly when kS.ErrorMode is set
    def Resume(self):
        First = kS.FirstError()
        kS.Reset()
        return {"resumed": First is not None, "error": First}

    async def OnCom(self, Method, *Args):
        """Run Method on the single COM thread (COM objects must stay on the thread that created them).
        A method that hit the global handler returns its safe default: the caller gets the halt error instead."""
        if kS.ErrorMode:
            raise self.HaltedError()
        try:
            Loop = asyncio.get_running_loop()
            if Loop.get_exception_handler() is None:
                Loop.set_exception_handler(kS.OnAsyncioError)  # tasks nobody awaited report too
            Result = await Loop.run_in_executor(self.Pool, Method, *Args)
            if kS.ErrorMode:
                raise self.HaltedError()
            return Result
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.OnCom")
            raise self.HaltedError() from e

    def App(self):
        """The PowerPoint application; on another OS that is an expected state."""
        if kS.ErrorMode:
            return None
        try:
            if sys.platform != "win32":
                raise ToolReportableException("PowerPoint Live needs Windows with desktop PowerPoint")
            import win32com.client
            return win32com.client.Dispatch("PowerPoint.Application")
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.App")
            return None

    def Prs(self):
        """The deck we work on (bound by FullName), or the active deck the first time."""
        if kS.ErrorMode:
            return None
        try:
            App = self.App()
            if self.PrsName:
                for P in App.Presentations:
                    if P.FullName == self.PrsName:
                        return P
                raise ToolReportableException(f"{self.PrsName} is no longer open - call powerpoint_open")
            if App.Presentations.Count == 0:
                raise ToolReportableException("no deck is open - call powerpoint_open with a path")
            P = App.ActivePresentation
            self.PrsName = P.FullName
            return P
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Prs")
            return None

    def OpenNames(self):
        """FullName of every open presentation."""
        if kS.ErrorMode:
            return []
        try:
            return [P.FullName for P in self.App().Presentations]
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.OpenNames")
            return []

    def Rebind(self, OpenBefore):
        """After a change: the deck we work on, found again by name. Follows a Save As (the one new name);
        returns None and forgets the deck when the change closed it."""
        if kS.ErrorMode:
            return None
        try:
            Open = list(self.App().Presentations)
            for P in Open:
                if P.FullName == self.PrsName:
                    return P
            Added = [P for P in Open if P.FullName not in OpenBefore]
            if len(Added) == 1:
                self.PrsName = Added[0].FullName
                return Added[0]
            self.PrsName = None
            self.Current = 1
            self.Tracker.Clear()
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Rebind")
            return None

    def Clamp(self, Prs, N):
        """Slide number N kept inside the deck (0 for an empty deck)."""
        if kS.ErrorMode:
            return 0
        try:
            return max(1, min(int(N), Prs.Slides.Count)) if Prs.Slides.Count else 0
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Clamp")
            return 0

    def SyncWindow(self, Prs):
        """Let the PowerPoint window follow the current slide, when its view can (sorter/reading view cannot)."""
        if kS.ErrorMode:
            return None
        try:
            if Prs.Windows.Count and self.Current and Prs.Windows(1).ViewType in GOTO_VIEWS:
                Prs.Windows(1).View.GotoSlide(self.Current)
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.SyncWindow")
            return None

    def Summary(self, Prs):
        """Deck, current slide, count, version and size: the keys every tool result starts with."""
        if kS.ErrorMode:
            return {}
        try:
            return {"deck": os.path.basename(Prs.FullName), "path": Prs.FullName, "slide": self.Current,
                    "count": Prs.Slides.Count, "version": self.Version,
                    "width_pt": Prs.PageSetup.SlideWidth, "height_pt": Prs.PageSetup.SlideHeight}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Summary")
            return {}

    def Png(self, Slide, Width):
        """Base64 PNG of one slide through Slide.Export."""
        if kS.ErrorMode:
            return None
        try:
            Fd, Path = tempfile.mkstemp(suffix=".png", dir=self.Work)
            os.close(Fd)
            try:
                Prs = Slide.Parent
                H = round(Width * Prs.PageSetup.SlideHeight / Prs.PageSetup.SlideWidth)
                Slide.Export(Path, "PNG", int(Width), H)
                with open(Path, "rb") as F:
                    return base64.b64encode(F.read()).decode()
            finally:
                os.remove(Path)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Png")
            return None

    def Walk(self, Shapes):
        """Every shape, group members included."""
        if kS.ErrorMode:
            return
        try:
            for Shape in Shapes:
                yield Shape
                if Shape.Type == 6:  # msoGroup
                    yield from self.Walk(Shape.GroupItems)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Walk")
            return

    def FindShape(self, Slide, Name=None, ShapeId=None):
        """The shape with this id or name, or None."""
        if kS.ErrorMode:
            return None
        try:
            for Shape in self.Walk(Slide.Shapes):
                if (ShapeId and Shape.Id == int(ShapeId)) or (Name and Shape.Name == Name):
                    return Shape
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.FindShape")
            return None

    def Mutate(self, Mutation):
        """Version first, then change, then record what changed on every slide."""
        if kS.ErrorMode:
            return None
        try:
            Prs = self.Prs()
            Tracker = self.Tracker
            self.History.Snapshot(Mutation.Label)
            Before = {Prs.Slides(I).SlideID: Tracker.ShapeMap(Prs.Slides(I)) for I in range(1, Prs.Slides.Count + 1)}
            OpenBefore = self.OpenNames()
            Result = Mutation.Apply(Prs)
            Prs = self.Rebind(OpenBefore)
            if Prs is None:  # the change itself closed the deck: a defined outcome, not an error
                self.Version += 1
                return {**(Result or {}), "deck_closed": True, "slide": 0, "count": 0, "version": self.Version,
                        "note": "The deck was closed by this change. Call powerpoint_open to work on a deck again."}
            for I in range(1, Prs.Slides.Count + 1):
                Slide = Prs.Slides(I)
                if Slide.SlideID in Before:
                    Tracker.NoteChanges(Slide, Before[Slide.SlideID])
                else:
                    Tracker.NoteNewSlide(Prs, Slide)
            self.Current = self.Clamp(Prs, self.Current or 1)
            self.Version += 1
            self.SyncWindow(Prs)
            return {**self.Summary(Prs), **(Result or {})}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Mutate")
            return None

    # ------------------------------------------------ COM work (runs on the COM thread)

    def Open(self, Path, Slide):
        """Open or attach to a deck and pick the current slide."""
        if kS.ErrorMode:
            return None
        try:
            App = self.App()
            if Path:
                Full = os.path.abspath(Path)
                Hit = next((P for P in App.Presentations if P.FullName.lower() == Full.lower()), None)
                Hit = Hit or next((P for P in App.Presentations if P.Name.lower() == os.path.basename(Path).lower()),
                                  None)
                if Hit is None:
                    if not os.path.exists(Full):
                        raise ToolInputException(f"not found: {Full}")
                    Hit = App.Presentations.Open(Full)
                if Hit.FullName != self.PrsName:
                    self.Tracker.Clear()
                self.PrsName = Hit.FullName
            Prs = self.Prs()
            self.Current = self.Clamp(Prs, Slide or self.Current)
            self.Version += 1
            self.SyncWindow(Prs)
            return self.Summary(Prs)
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Open")
            return None

    def Run(self, Code, Slide, Label):
        """powerpoint_run: the code's print() output and its error text come back with the deck summary."""
        if kS.ErrorMode:
            return None
        try:
            Mutation = kRunMutation(self, Code, Slide, Label)
            Res = self.Mutate(Mutation)
            return {**(Res or {}), "output": Mutation.Output.getvalue()[-8000:], "error": Mutation.Error}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Run")
            return None

    def Show(self, Slide):
        """Make slide n current."""
        if kS.ErrorMode:
            return None
        try:
            Prs = self.Prs()
            self.Current = self.Clamp(Prs, Slide)
            self.Version += 1
            self.SyncWindow(Prs)
            return self.Summary(Prs)
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Show")
            return None

    def State(self):
        """Current slide, its shapes, recent changes and a change signature."""
        if kS.ErrorMode:
            return None
        try:
            Prs = self.Prs()
            if self.Current > Prs.Slides.Count:
                self.Current = Prs.Slides.Count
            Out = {**self.Summary(Prs), "signature": self.Tracker.Signature(Prs, self.Current), "shapes": [],
                   "changes": None}
            if self.Current:
                Slide = Prs.Slides(self.Current)
                Shapes = self.Tracker.NoteChanges(Slide)
                Out["shapes"] = [{"id": K, "name": V["name"], "box": V["box"], "text": V["snippet"]}
                                 for K, V in Shapes.items()]
                Out["changes"] = self.Tracker.Changes.get(Slide.SlideID)
                Out["slide_id"] = Slide.SlideID
                Out["variant_of"] = self.Variants.VariantOf(Slide)
                Out["variants"] = len(self.Variants.Of(Prs, Slide.SlideID))
            return Out
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.State")
            return None

    def AccuratePng(self, Prs, N, Width):
        """Render one slide through Save As JPEG (true embedded fonts), not Slide.Export."""
        if kS.ErrorMode:
            return None
        try:
            Src = os.path.join(self.Work, "acc-src" + (os.path.splitext(Prs.FullName)[1] or ".pptx"))
            Prs.SaveCopyAs(Src)
            Tmp = self.App().Presentations.Open(Src, -1, 0, 0)  # read-only, no window
            Out = os.path.join(self.Work, "acc-out")
            try:
                for I in range(Tmp.Slides.Count, 0, -1):
                    if I != N:
                        Tmp.Slides(I).Delete()
                shutil.rmtree(Out, ignore_errors=True)
                Tmp.SaveCopyAs(Out, 17)  # ppSaveAsJPG -> folder with Slide1.JPG
            finally:
                Tmp.Close()
            Jpg = next(os.path.join(Out, F) for F in os.listdir(Out) if F.lower().endswith(".jpg"))
            from PIL import Image
            Im = Image.open(Jpg)
            Im = Im.resize((int(Width), round(int(Width) * Im.height / Im.width)))
            Buf = io.BytesIO()
            Im.save(Buf, "PNG")
            shutil.rmtree(Out, ignore_errors=True)
            return base64.b64encode(Buf.getvalue()).decode()
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.AccuratePng")
            return None

    def Image(self, Slide, Width, Accurate):
        """PNG of a slide, fast or accurate."""
        if kS.ErrorMode:
            return None
        try:
            Prs = self.Prs()
            N = self.Clamp(Prs, Slide or self.Current)
            if not N:
                return {"slide": 0, "png": None}
            Sl = Prs.Slides(N)
            Png = self.AccuratePng(Prs, N, Width) if Accurate else self.Png(Sl, Width)
            return {"slide": N, "slide_id": Sl.SlideID, "count": Prs.Slides.Count, "png": Png,
                    "accurate": bool(Accurate), "signature": self.Tracker.Signature(Prs, N)}
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Image")
            return None

    def Move(self, SlideIds, BeforeId, Section):
        """powerpoint_move."""
        if kS.ErrorMode:
            return None
        try:
            return self.Mutate(kMoveMutation(self, SlideIds, BeforeId, Section))
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Move")
            return None

    def SectionsEdit(self, Action, Section, Name, BeforeId):
        """powerpoint_sections; an unknown action is refused before a version is saved."""
        if kS.ErrorMode:
            return None
        try:
            if Action not in SECTION_ACTIONS:
                raise ToolInputException("action must be add, rename or delete")
            return self.Mutate(kSectionMutation(self, Action, Section, Name, BeforeId))
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.SectionsEdit")
            return None

    def Hide(self, SlideIds, Hidden):
        """powerpoint_hide."""
        if kS.ErrorMode:
            return None
        try:
            return self.Mutate(kHideMutation(self, SlideIds, Hidden))
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Hide")
            return None

    def Fix(self, SlideId, ShapeId, Code):
        """powerpoint_fix; an unknown code is refused before a version is saved."""
        if kS.ErrorMode:
            return None
        try:
            if Code not in FIXABLE:
                raise ToolInputException(f"{Code} has no automatic fix - ask Claude")
            return self.Mutate(kFixMutation(self, SlideId, ShapeId, Code, self.Lint.FloorPt()))
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPowerPointLive.Fix")
            return None


Live = kPowerPointLive()


# ---------------------------------------------------------------- MCP surface

@mcp.resource(URI, name="PowerPoint Live", mime_type=MIME, meta={"ui": {"prefersBorder": True}})
def SlideView() -> str:
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        with open(os.path.join(HERE, "view.html"), encoding="utf-8") as F:
            return F.read()
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.SlideView")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_open", meta=UI)
async def PowerpointOpen(path: str = "", slide: int = 0) -> dict:
    """Open a .pptx in desktop PowerPoint (or attach to it if already open; empty path = the active deck)
    and make `slide` (1-based) the current slide. The live view shows the deck as it changes."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Open, path, slide)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointOpen")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_run", meta=UI)
async def PowerpointRun(code: str, slide: int = 0, label: str = "") -> dict:
    """Run Python against the open deck through COM. In scope: `app`, `prs` (the deck), `slide` (the current slide),
    `goto(n)` (change the current slide - the live view follows), `win32com`. print() output is returned.
    If the code closes the deck the result has "deck_closed": true (call powerpoint_open again).
    `label` names the change in the History view. A version is saved first, so the user can undo it.
    Changed shapes are highlighted in the live view. Pin to `prs`, never app.ActivePresentation."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Run, code, slide, label)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointRun")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_show", meta=UI)
async def PowerpointShow(slide: int) -> dict:
    """Make slide n (1-based) the current slide in the live view and the PowerPoint window."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Show, slide)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointShow")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_move", meta=UI)
async def PowerpointMove(slide_ids: list[int], before_slide_id: int = 0, section: int = 0) -> dict:
    """Move slides (by SlideID, kept in deck order) to just before `before_slide_id`, or (with `section`, 1-based)
    to the end of that section; neither = to the end of the deck. SlideIDs: prs.Slides(n).SlideID."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Move, slide_ids, before_slide_id, section)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointMove")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_sections", meta=UI)
async def PowerpointSections(action: str, section: int = 0, name: str = "", before_slide_id: int = 0) -> dict:
    """Edit sections. action: "add" (a section named `name` starting at `before_slide_id`, or at the end),
    "rename" (section n to `name`), "delete" (section n; its slides join the section before)."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.SectionsEdit, action, section, name, before_slide_id)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointSections")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_hide", meta=UI)
async def PowerpointHide(slide_ids: list[int], hidden: bool = True) -> dict:
    """Hide (or unhide) slides in the slide show, by SlideID."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Hide, slide_ids, hidden)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointHide")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_fix", meta=UI)
async def PowerpointFix(slide_id: int, shape_id: int, code: str) -> dict:
    """Apply the automatic fix for a lint finding to one shape, live: unused_placeholder, body_below_floor,
    text_overflow (shrinks using PowerPoint's own layout, never below the floor), picture_stretched,
    a11y_missing_alt_text (charts and tables), chart_default_palette, chart_legend_steals_plot,
    chart_accounting_zero_dash. A version is saved first."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Fix, slide_id, shape_id, code)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointFix")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_history", meta=UI)
async def PowerpointHistory() -> dict:
    """Versions saved before each change made through these tools (newest first)."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.History.List)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointHistory")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_restore", meta=UI)
async def PowerpointRestore(version: int) -> dict:
    """Put a saved version back: the deck file is replaced with it and reopened. The current state is saved
    as a version first, so a restore can itself be undone. Needs a deck that has been saved to disk."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.History.Restore, version)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointRestore")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_resume", meta=UI)
async def PowerpointResume() -> dict:
    """Re-arm PowerPoint Live after an unexpected error halted it (the view's Resume button calls this).
    Returns the error that caused the halt, if there was one."""
    if kS.ErrorMode:
        return Live.Resume()
    try:
        return Live.Resume()
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointResume")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_send_error_report", meta=UI)
async def PowerpointSendErrorReport(send: bool) -> dict:
    """Answer a waiting error report. ONLY after asking the user "Do you want to send this error message?" and
    getting their explicit answer: send=true if they said yes (it goes to the support flow), send=false if they
    said no (it is discarded). Never call this with send=true on your own."""
    if kS.ErrorMode:
        return kErrorReport.Decide(send)
    try:
        return kErrorReport.Decide(send)
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointSendErrorReport")
        raise Live.HaltedError() from e


@mcp.tool(name="error_reports", meta=APP_ONLY)
async def ErrorReports() -> dict:
    """(View only) Error reports waiting for the person's answer to "Do you want to send this error message?".
    Works while halted: the question must be answerable."""
    if kS.ErrorMode:
        return kErrorReport.PendingForView()
    try:
        return kErrorReport.PendingForView()
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.ErrorReports")
        raise Live.HaltedError() from e


@mcp.tool(name="error_report_answer", meta=APP_ONLY)
async def ErrorReportAnswer(send: bool) -> dict:
    """(View only - never the model) The person's answer: send=True sends every waiting report, False discards
    them. Only the view's Yes/No buttons call this."""
    if kS.ErrorMode:
        return kErrorReport.Decide(send)
    try:
        return kErrorReport.Decide(send)
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.ErrorReportAnswer")
        raise Live.HaltedError() from e


@mcp.tool(name="slide_state", meta=APP_ONLY)
async def SlideState() -> dict:
    """(View only) Current slide, its shapes, recent changes and a change signature."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.State)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.SlideState")
        raise Live.HaltedError() from e


@mcp.tool(name="slide_image", meta=APP_ONLY)
async def SlideImage(slide: int = 0, width: int = 1280, accurate: bool = False) -> dict:
    """(View only) PNG of a slide. accurate=True renders through Save As JPEG (true embedded fonts), slower."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Image, slide, max(320, min(width, 2560)), accurate)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.SlideImage")
        raise Live.HaltedError() from e


@mcp.tool(name="deck_outline", meta=APP_ONLY)
async def DeckOutline() -> dict:
    """(View only) Sections and slides (id, title, hidden, notes words, change signature)."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Sorter.Outline)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.DeckOutline")
        raise Live.HaltedError() from e


@mcp.tool(name="slide_thumbs", meta=APP_ONLY)
async def SlideThumbs(slide_ids: list[int], width: int = 320) -> dict:
    """(View only) Small PNGs of the given slides, cached by content."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Sorter.ThumbsFor, slide_ids[:12], max(120, min(width, 640)))
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.SlideThumbs")
        raise Live.HaltedError() from e


@mcp.tool(name="deck_lint", meta=APP_ONLY)
async def DeckLint() -> dict:
    """(View only) lint_deck.py findings with shape boxes, cached until the deck changes."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Lint.Lint)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.DeckLint")
        raise Live.HaltedError() from e


@mcp.tool(name="history_thumb", meta=APP_ONLY)
async def HistoryThumb(version: int) -> dict:
    """(View only) Thumbnail of the current slide when a version was saved."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.History.Thumb, version)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.HistoryThumb")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_slideshow", meta=UI)
async def PowerpointSlideshow(action: str = "state") -> dict:
    """Drive the real slide show. action: "start" (SlideShowSettings.Run, from the current slide; a running show is
    reused), "next", "previous", "black" (toggle a black screen), "end", "state". Every answer has running, state
    (running/paused/black/white/done/stopped), the slide on screen, its notes, timing marker and talk estimate, the
    next slide and the planned time before this slide. Next/previous/black/end need a running show."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.SlideShow.Do, (action or "state").strip().lower())
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointSlideshow")
        raise Live.HaltedError() from e


@mcp.tool(name="design_previews", meta=APP_ONLY)
async def DesignPreviews(slide: int = 0, directions: list[str] | None = None, width: int = 480) -> dict:
    """(View only) The slide restyled in up to six design directions (empty = six contrasting ones), rendered from a
    copy - the open deck is not touched. Also lists every direction."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Design.Previews, slide, directions or [], max(160, min(width, 960)))
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.DesignPreviews")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_apply_direction", meta=UI)
async def PowerpointApplyDirection(direction: str) -> dict:
    """Apply a design direction (an id from scripts/directions.json, e.g. "editorial-serif") to the whole deck:
    its theme colours and heading/body fonts on every slide master, exactly as build_deck.py writes them.
    Shapes with hard-coded colours keep them. A version is saved first."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Design.ApplyLive, (direction or "").strip())
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointApplyDirection")
        raise Live.HaltedError() from e


@mcp.tool(name="theme_info", meta=APP_ONLY)
async def ThemeInfo() -> dict:
    """(View only) The current slide's theme: colour slots as RRGGBB, major/minor fonts and contrast pairs
    (text 4.5:1, graphics and large text 3:1) with pass/fail."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.ThemeEditor.Info)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.ThemeInfo")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_set_theme", meta=UI)
async def PowerpointSetTheme(colors: dict[str, str] | None = None, major_font: str = "", minor_font: str = "") -> dict:
    """Change the deck's theme on every slide master: `colors` maps slots (dk1, lt1, dk2, lt2, accent1-accent6,
    hlink, folHlink) to RRGGBB hex; `major_font` / `minor_font` set the heading / body typeface. Only what is given
    changes. Check contrast after (text 4.5:1 on its background). A version is saved first."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.ThemeEditor.Set, colors or {}, major_font, minor_font)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointSetTheme")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_layout_variants", meta=UI)
async def PowerpointLayoutVariants(slide: int = 0, count: int = 3) -> dict:
    """Start layout alternatives for slide n (0 = current). Returns the original's SlideID, any variants that already
    exist, and HOW to make them: you create `count` variants yourself with one powerpoint_run - duplicate the slide,
    keep each copy hidden right after the original, tag it PPTLIVE-VARIANT-OF=<original SlideID> and
    PPTLIVE-VARIANT-NAME=<short name>, then rearrange it into a different layout. The user compares them in the
    view and picks one (powerpoint_choose_variant) or drops them (powerpoint_discard_variants)."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Variants.Guide, slide, count)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointLayoutVariants")
        raise Live.HaltedError() from e


@mcp.tool(name="layout_variants", meta=APP_ONLY)
async def LayoutVariants(slide_id: int, width: int = 320) -> dict:
    """(View only) The original slide and its layout variants, with thumbnails."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Variants.List, slide_id, max(120, min(width, 640)))
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.LayoutVariants")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_choose_variant", meta=UI)
async def PowerpointChooseVariant(slide_id: int, variant_id: int) -> dict:
    """Replace slide `slide_id` with its layout variant `variant_id` (both SlideIDs): the variant is unhidden
    (unless the original was hidden), untagged and moved into place; the original and the other variants are
    deleted. A version is saved first."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Variants.Choose, slide_id, variant_id)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointChooseVariant")
        raise Live.HaltedError() from e


@mcp.tool(name="powerpoint_discard_variants", meta=UI)
async def PowerpointDiscardVariants(slide_id: int) -> dict:
    """Delete every layout variant of slide `slide_id` (SlideID); the original stays. A version is saved first."""
    if kS.ErrorMode:
        raise Live.HaltedError()
    try:
        return await Live.OnCom(Live.Variants.Discard, slide_id)
    except kToolException:
        raise
    except Exception as e:
        kS.GlobalErrorHandler(e, "server.PowerpointDiscardVariants")
        raise Live.HaltedError() from e


if __name__ == "__main__":  # the MCP entry point: FastMCP owns the loop; every tool above carries the pattern
    mcp.run()
