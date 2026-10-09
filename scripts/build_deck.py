"""Build a deck from a JSON (or YAML) spec - any OS, python-pptx only.

    uvx --with python-pptx --with pillow python scripts/build_deck.py spec.json --out deck.pptx
    uvx --with python-pptx --with pillow python scripts/build_deck.py spec.json --out deck.pptx --lint
    python scripts/build_deck.py --list-directions

The spec names a design direction (scripts/directions.json) or a template (.pptx/.potx), then lists
slides by pattern. The builder applies the skill's rules by construction: 1440 x 810 pt, a real title
placeholder on every slide, theme colours and fonts (so the deck re-themes), named shapes, stable slide
ids, cover-cropped pictures, native charts with one highlighted finding, alt text and speaker notes.
Inputs are checked against each pattern's limits first; nothing is written if a limit is broken
(--force builds anyway). The spec format is in reference/BUILDER.md.

Patterns: title, section, statement, big_number, kpi, bullets, compare, process, timeline, quote,
chart, table, image, matrix, email, kpi_chart, cost_table, quiz, risks, metrics, next_steps. Every pattern lays its content out over
the whole body area and grows text up to a ceiling per role (never below the 27 pt floor for sentences).
"""
import argparse
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import sys
import time

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Pt

from kShared import ToolInputException, ToolReportableException, kRun, kS, kToolException

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from _figures import kFigures  # noqa: E402
from _measure import LINE_HEIGHT, kMeasure  # noqa: E402  (the skill's own helpers sit beside this script)
from _rules import kRules  # noqa: E402
from _theme import kTheme  # noqa: E402

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": A, "p": P}

# Canvas and type scale for a 1440 x 810 pt slide (1920 x 1080 px). Body never below 18 pt.
W, H, M = 1440, 810, 80
TITLE_TOP, TITLE_H = 56, 100
BODY_TOP, BODY_BOTTOM = 196, 730
SIZE = {"title": 54, "section": 84, "statement": 72, "hero": 220, "kpi": 84, "h2": 46, "body": 34,
        "small": 30, "caption": 28, "label": 24}  # 1.5x a 960-pt slide: body 34 ~ 23 pt, floor 27 ~ 18 pt
GROW = {"body": 44, "point": 40, "detail": 36, "heading": 44, "value": 120, "note": 32}  # ceilings text grows to
GAP = 24
EDGE = 8                        # the accent edge every card carries (card language: tint + edge, highlight = accent fill)
FOOTER_TOP, FOOTER_H, FOOTER_SIZE = 760, 30, 18  # deck footer and page number, below BODY_BOTTOM (caption tier)
KICKER_SIZE, KICKER_H = 24, 34  # the small accent label above a title
KICKER_TOP, KICKER_GAP = 14, 6  # the kicker never starts above this; the gap between it and the title's ink
TITLE_DROP = 16                 # a two-line title's box reaches this much lower (the body starts at BODY_TOP)
LOOSE_LINE = 1.28               # line height of the loosest renderer (LibreOffice, serif faces) for title ink
NBSP = "\u00a0"

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
FOOTER_TYPES = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
TEXT, MUTED, ACCENT, BG, QUIET = (MSO_THEME_COLOR.TEXT_1, MSO_THEME_COLOR.TEXT_2, MSO_THEME_COLOR.ACCENT_1,
                                  MSO_THEME_COLOR.BACKGROUND_1, MSO_THEME_COLOR.BACKGROUND_2)
SOFT = MSO_THEME_COLOR.ACCENT_2  # the accent mixed towards the background: card edges that are not the highlight
FAINT = MSO_THEME_COLOR.ACCENT_6  # the accent most of the way to the background: the empty part of a share

LIMITS = {  # pattern: {field: max characters} plus list-length ranges, checked before building
    "title": {"title": 70, "subtitle": 120},
    "section": {"title": 60, "eyebrow": 30},
    "statement": {"title": 90, "support": 140, "decision": 140, "decision_label": 30, "owner": 40, "date": 30,
                  "points": (1, 3), "points.*": 60, "figure.value": 12, "figure.label": 28, "figure.note": 50,
                  "figure.trend": (2, 12), "figure.trend_labels": (2, 12), "figure.trend_labels.*": 16},
    "big_number": {"title": 70, "number": 12, "unit": 10, "caption": 90, "points": (1, 3), "points.*": 70},
    "kpi": {"title": 70, "metrics": (2, 6), "metrics.value": 12, "metrics.label": 28, "metrics.note": 50,
            "metrics.trend": (2, 12), "metrics.trend_labels": (2, 12), "metrics.trend_labels.*": 16,
            "decision": 140, "decision_label": 30},
    "bullets": {"title": 70, "items": (1, 7), "items.*": 100},
    "compare": {"title": 70, "columns": (2, 3), "columns.heading": 40, "columns.points": (1, 4),
                "columns.points.*": 70},
    "process": {"title": 70, "steps": (3, 8), "steps.label": 30, "steps.detail": 60},
    "timeline": {"title": 70, "events": (3, 7), "events.date": 16, "events.label": 40, "window": 30},
    "quote": {"quote": 240, "attribution": 40, "role": 60},
    "chart": {"title": 70, "categories": (2, 24), "series": (1, 6)},
    "table": {"title": 70, "header": (2, 6), "rows": (1, 8)},
    "image": {"title": 70, "caption": 120},
    "matrix": {"title": 70, "quadrants": (4, 4), "quadrants.heading": 30, "quadrants.text": 90,
               "x_axis": 30, "y_axis": 30},
    "email": {"title": 70, "from": 70, "to": 70, "subject": 90, "body": (1, 6), "body.*": 160, "attachment": 40,
              "callouts": (0, 5), "callouts.note": 70},
    "kpi_chart": {"title": 70, "metrics": (1, 4), "metrics.value": 12, "metrics.label": 28, "metrics.note": 40,
                  "metrics.trend": (2, 12), "metrics.trend_labels": (2, 12), "metrics.trend_labels.*": 16,
                  "categories": (2, 12), "series": (1, 4), "caption": 90},
    "cost_table": {"title": 70, "rows": (1, 7), "rows.item": 40, "rows.detail": 70, "unit": 12, "total_label": 20,
                   "note": 140},
    "quiz": {"title": 90, "options": (2, 4), "options.text": 80, "explain": 160, "answer_title": 55},
    "risks": {"title": 70, "risks": (2, 4), "risks.risk": 50, "risks.mitigation": 120},
    "metrics": {"title": 70, "rows": (1, 6), "rows.metric": 40, "rows.baseline": 24, "rows.target": 24,
                "rows.owner": 24, "rows.date": 16, "rule": 140},
    "next_steps": {"title": 70, "steps": (2, 6), "steps.action": 80, "steps.owner": 30, "steps.date": 20,
                   "decision": 140, "decision_label": 30},
}
KICKER_MAX = 30  # characters; at most three words, so it stays a label
COUNT_HINTS = {  # pattern.list: the pattern that fits when the count is outside the range
    "kpi.metrics": "one number is a 'big_number' slide; a metric with its series is a 'kpi_chart'; more than six "
                   "is a 'table'",
    "kpi_chart.metrics": "more than four metrics: a 'kpi' slide, or a 'table'",
    "statement.points": "more than three: a 'bullets' slide, or 'compare'",
    "big_number.points": "more than three: a 'kpi' slide (each number its own tile) or 'bullets'",
    "bullets.items": "more than seven: split the slide, or a 'table'",
    "compare.columns": "one column is a 'statement' with 'points'; four or more is a 'table'",
    "process.steps": "two steps are a 'compare'; more than eight: split into two slides",
    "timeline.events": "two events are a 'compare'; more than seven: split, or a 'table'",
    "risks.risks": "one risk is a 'statement'; more than four: a 'table' with likelihood and impact columns",
    "next_steps.steps": "one step is a 'statement' with a 'decision'",
    "metrics.rows": "more than six: a 'table'",
}
COMPARE_HEADING = {2: 40, 3: 24}  # compare: heading characters per column count (3 narrow columns hold less)


# ---------------------------------------------------------------- JSON Schema (generated from LIMITS)

FIELDS = {  # pattern: {field: kind}; kinds: text, int, number, list[text], list[obj:{...}], obj
    "title": {"subtitle": "text"},
    "section": {"eyebrow": "text"},
    "statement": {"support": "text", "decision": "text", "decision_label": "text", "owner": "text", "date": "text",
                  "points": "list[text]",
                  "figure": ("object", {"value": "text", "label": "text", "note": "text", "trend": "list[number]",
                                        "trend_labels": "list[text]"})},
    "big_number": {"number": "text", "unit": "text", "caption": "text", "visual": "enum:auto,dots,bar,none",
                   "points": "list[text]"},
    "kpi": {"metrics": {"value": "text", "label": "text", "note": "text", "trend": "list[number]",
                        "trend_labels": "list[text]"},
            "highlight": "int", "decision": "text", "decision_label": "text", "trend_chart": "bool"},
    "bullets": {"items": "list[text]", "allow_split": "bool"},
    "compare": {"columns": {"heading": "text", "points": "list[text]"}, "highlight": "int"},
    "process": {"steps": {"label": "text", "detail": "text"}, "highlight": "int"},
    "timeline": {"events": {"date": "text", "label": "text"}, "highlight": "int", "window": "text"},
    "quote": {"quote": "text", "attribution": "text", "role": "text"},
    "chart": {"type": "enum:column,bar,line,pie", "categories": "list[text]",
              "series": {"name": "text", "values": "list[number]"}, "highlight": "int|text",
              "number_format": "text", "caption": "text", "alt": "text"},
    "table": {"header": "list[text]", "rows": "list[list[text]]", "highlight_row": "int", "alt": "text"},
    "image": {"image": "text", "caption": "text", "alt": "text", "focus_x": "number", "focus_y": "number"},
    "matrix": {"quadrants": {"heading": "text", "text": "text"}, "x_axis": "text", "y_axis": "text",
               "highlight": "int"},
    "email": {"from": "text", "to": "text", "subject": "text", "body": "list[text]", "attachment": "text",
              "callouts": {"target": "enum:from,to,subject,body,attachment", "line": "int", "note": "text"}},
    "kpi_chart": {"metrics": {"value": "text", "label": "text", "note": "text", "trend": "list[number]",
                              "trend_labels": "list[text]"},
                  "highlight_metric": "int",
                  "type": "enum:column,bar,line,pie", "categories": "list[text]",
                  "series": {"name": "text", "values": "list[number]"}, "highlight": "int|text",
                  "number_format": "text", "caption": "text", "alt": "text"},
    "cost_table": {"rows": {"item": "text", "amount": "number", "detail": "text"}, "unit": "text",
                   "total_label": "text", "note": "text", "highlight": "int"},
    "quiz": {"options": {"text": "text", "correct": "bool"}, "explain": "text", "reveal": "enum:notes,slide",
             "answer_title": "text"},
    "risks": {"risks": {"risk": "text", "likelihood": "enum:low,medium,high", "impact": "enum:low,medium,high",
                        "mitigation": "text"}, "highlight": "int"},
    "metrics": {"rows": {"metric": "text", "baseline": "text", "target": "text", "owner": "text", "date": "text"},
                "rule": "text", "highlight": "int"},
    "next_steps": {"steps": {"action": "text", "owner": "text", "date": "text"}, "decision": "text",
                   "decision_label": "text"},
}
OPTIONAL_SUB = {"process.detail", "matrix.text", "kpi.note", "kpi_chart.note", "cost_table.detail",
                "email.line", "kpi.trend", "kpi_chart.trend", "metrics.owner", "metrics.date", "next_steps.owner",
                "next_steps.date", "timeline.date", "kpi.trend_labels", "kpi_chart.trend_labels", "statement.note", "statement.trend",
                "statement.trend_labels"}  # sub-fields of list items that may be left out
REQUIRED = {"title": ["title"], "section": ["title"], "statement": ["title"], "quote": ["quote"],
            "big_number": ["title", "number"], "kpi": ["title", "metrics"], "bullets": ["title", "items"],
            "compare": ["title", "columns"], "process": ["title", "steps"], "timeline": ["title", "events"],
            "chart": ["title", "categories", "series"], "table": ["title", "header", "rows"],
            "image": ["title", "image"], "matrix": ["title", "quadrants"],
            "email": ["title", "from", "subject", "body"], "kpi_chart": ["title", "metrics"],
            "cost_table": ["title", "rows"], "quiz": ["title", "options"], "risks": ["title", "risks"],
            "metrics": ["title", "rows"], "next_steps": ["title", "steps"]}
_STRS = {"oneOf": [{"type": "string"}, {"type": "array", "items": {"type": "string"}}]}
NOTES_SCHEMA = {"description": "Speaker notes: a string, or key_fact / facts / assumptions / pitfalls / sources / "
                               "qa (lists may be a single string; a Q&A item may be a string).", "oneOf": [
    {"type": "string"},
    {"type": "object", "additionalProperties": False, "properties": {
        "key_fact": {"type": "string"}, "facts": _STRS, "pitfalls": _STRS, "sources": _STRS,
        "assumptions": dict(_STRS, description="What the slide assumes beyond the brief (ratings, targets, dates, "
                                               "owners); written into the notes under ASSUMPTIONS:."),
        "qa": {"oneOf": [{"type": "object"}, {"type": "array", "items": {"oneOf": [
            {"type": "string"},
            {"type": "object", "additionalProperties": False, "required": ["q"],
             "properties": {"q": {"type": "string"}, "a": {"type": "string"}}}]}}]}}}]}

FLOOR = 27        # body text never shrinks below this (1440-pt grid; scaled by F())
LABEL_MIN = 24    # short labels (< 4 words) may go this small
FIGURE_W, FIGURE_H, FIGURE_TREND_H = 440, 320, 400  # the figure tile beside a decision (with a trend: taller)
POINTS_W = 520    # the column of point cards beside a statement
SHARE_W = 470     # the dot grid or bar beside a big number that is a share
TREND_MAX = 200   # the tallest a tile's trend (figures, bars, periods) grows
TREND_LABEL = 24  # the figures and periods on a tile's trend bars: the label floor, never smaller
TREND_BARS = 72   # the shortest a tile's bars may be; with less room the trend moves to the note, never tiny
ROW_PAD, ROW_CARD = 20, 170  # timeline/stage label cards in one row: inner padding, least height
RAMP = [0, 0.35, -0.3, 0.6, -0.5, 0.75]  # shades of the accent for series 1..6
CHART_TYPES = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
               "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}
PATTERNS = {"title": "Title", "section": "Section", "statement": "Statement", "big_number": "BigNumber",
            "kpi": "Kpi", "bullets": "Bullets", "compare": "Compare", "process": "Process",
            "timeline": "Timeline", "quote": "Quote", "chart": "Chart", "table": "Table", "image": "Image",
            "matrix": "Matrix", "email": "Email", "kpi_chart": "KpiChart", "cost_table": "CostTable",
            "quiz": "Quiz", "risks": "Risks", "metrics": "Metrics", "next_steps": "NextSteps"}  # pattern -> kSlidePatterns method
NUMERIC_KEYS = {"values", "trend", "highlight", "highlight_row", "highlight_metric", "focus_x", "focus_y", "amount", "line",
                "correct"}
HIGHLIGHT_ITEMS = {"kpi": "metrics", "compare": "columns", "process": "steps", "timeline": "events",
                   "matrix": "quadrants", "chart": "categories", "kpi_chart": "categories", "cost_table": "rows",
                   "risks": "risks", "metrics": "rows"}
NEEDED_LIST = {"kpi": "metrics", "process": "steps", "timeline": "events", "compare": "columns",
               "matrix": "quadrants", "chart": "series", "table": "rows", "bullets": "items", "email": "body", "cost_table": "rows", "quiz": "options", "risks": "risks", "metrics": "rows",
               "next_steps": "steps"}
LAYOUT_CODES = {"text_overflow", "kicker_title_overlap", "tile_text_below_floor", "unwanted_wrap",
                "body_below_floor"}  # lint findings = unfit text (and every lint error: shape_overlap etc.)
AUTO_PASSES = 6  # rounds of automatic fixes (filler words, units, detail to the notes, a split) before exit 3
SEVERITY_ORDER = {"error": 0, "warn": 1, "info": 2}
PLAN_SKIP = {"notes", "id", "pattern", "image", "type", "reveal", "visual", "likelihood", "impact", "categories",
             "series", "header", "rows", "number_format", "alt", "highlight", "trend", "trend_labels", "kicker",
             "allow_split", "decision_label", "window"}
NO_CHROME = {"title", "section"}            # no footer or page number on covers and dividers
NO_KICKER = {"title", "section", "quote"}   # their own title treatment
LEVELS = {"high": "HIGH", "medium": "MEDIUM", "low": "LOW"}
MONTHS = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
UNIT_GLUE = re.compile(r"(?<=\d) (?=(?:%|\u2030|pp\b|pt\b|[kKMB]\b|bn\b|[kM]?(?:USD|EUR|SEK|GBP)\b|x\b))")


class kSpecSchema:
    """The spec's JSON Schema, generated from FIELDS / LIMITS / REQUIRED."""

    @staticmethod
    def Directions():
        """Every design direction in directions.json."""
        if kS.ErrorMode:
            return []
        try:
            with open(os.path.join(HERE, "directions.json"), encoding="utf-8") as Fh:
                return json.load(Fh)["directions"]
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecSchema.Directions")
            return []

    @staticmethod
    def KindSchema(Kind, MaxLen=None):
        """The schema of one field kind."""
        if kS.ErrorMode:
            return {}
        try:
            Text = {"type": ["string", "number"]} if MaxLen is None else {"type": ["string", "number"],
                                                                          "maxLength": MaxLen}
            if isinstance(Kind, dict):
                return {"type": "object", "additionalProperties": False, "properties": {}}
            if Kind == "text":
                return Text
            if Kind == "int":
                return {"type": "integer", "minimum": 0}
            if Kind == "number":
                return {"type": "number"}
            if Kind == "bool":
                return {"type": "boolean"}
            if Kind == "int|text":
                return {"type": ["integer", "string"]}
            if Kind.startswith("enum:"):
                return {"enum": Kind[5:].split(",")}
            if Kind == "list[text]":
                return {"type": "array", "items": Text}
            if Kind == "list[number]":
                return {"type": "array", "items": {"type": ["number", "null"]}}
            if Kind == "list[list[text]]":
                return {"type": "array", "items": {"type": "array", "items": {"type": ["string", "number"]}}}
            raise ValueError(Kind)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecSchema.KindSchema")
            return {}

    @staticmethod
    def SpecSchema():
        """JSON Schema for a build_deck spec, generated from FIELDS / LIMITS / REQUIRED."""
        if kS.ErrorMode:
            return {}
        try:
            Variants = []
            for Pat, Fields in FIELDS.items():
                Lim = LIMITS[Pat]
                Props = {"pattern": {"const": Pat}, "id": {"type": "string", "pattern": "^[A-Za-z0-9_.-]+$",
                                                           "description": "Stable slide id; keep it when the content changes."},
                         "title": {"type": ["string", "number"], "maxLength": Lim.get("title", 90),
                                   "description": "Write it as a claim, not a topic."},
                         "notes": NOTES_SCHEMA,
                         "kicker": {"type": "string", "maxLength": KICKER_MAX,
                                    "description": "Small label above the title (at most three words); defaults "
                                                   "to the current section's eyebrow; \"\" for none."}}
                for Name, Kind in Fields.items():
                    Single = isinstance(Kind, tuple)  # ("object", {...}): one object, not a list of them
                    if Single:
                        Kind = Kind[1]
                    if isinstance(Kind, dict):
                        Item = {"type": "object", "additionalProperties": False, "properties": {}, "required": []}
                        for Sub, Sk in Kind.items():
                            SubLim = Lim.get(f"{Name}.{Sub}")
                            if isinstance(SubLim, tuple):
                                Item["properties"][Sub] = dict(kSpecSchema.KindSchema(Sk, Lim.get(f"{Name}.{Sub}.*")),
                                                               minItems=SubLim[0], maxItems=SubLim[1])
                            else:
                                Item["properties"][Sub] = kSpecSchema.KindSchema(Sk, SubLim)
                            if Sk in ("text", "list[text]", "list[number]", "number") \
                                    and f"{Pat}.{Sub}" not in OPTIONAL_SUB:
                                Item["required"].append(Sub)
                        Schema = Item if Single else {"type": "array", "items": Item}
                    else:
                        Star = Lim.get(f"{Name}.*")
                        Schema = kSpecSchema.KindSchema(Kind, Star if Kind.startswith("list") else Lim.get(Name))
                    Rng = Lim.get(Name)
                    if isinstance(Rng, tuple):
                        Schema["minItems"], Schema["maxItems"] = Rng
                    Props[Name] = Schema
                Variants.append({"type": "object", "additionalProperties": False, "properties": Props,
                                 "required": ["pattern"] + REQUIRED[Pat]})
            return {
                "$schema": "https://json-schema.org/draft/2020-12/schema",
                "$id": "https://github.com/PeterKalmstrom/claude-powerpoint-skill/scripts/spec.schema.json",
                "title": "build_deck.py spec",
                "description": "A deck as data: a design direction or template, then slides by pattern. "
                               "Generated by `build_deck.py --print-schema`; see reference/BUILDER.md.",
                "type": "object", "additionalProperties": False, "required": ["slides"],
                "properties": {
                    "$schema": {"type": "string"},
                    "direction": {"enum": [D["id"] for D in kSpecSchema.Directions()]},
                    "template": {"type": "string", "description": ".pptx or .potx, relative to the spec"},
                    "footer": {"type": "string", "maxLength": 70,
                               "description": "Footer text on every content slide (not on title and section slides)."},
                    "page_numbers": {"type": "boolean", "description": "Page numbers on content slides (default true)."},
                    "kickers": {"type": "boolean", "description": "Kicker labels above titles (default true)."},
                    "sources": dict(_STRS, description="Deck-level sources (e.g. [\"the brief\"]): written into "
                                                       "the notes of every content slide whose notes name none."),
                    "facts": {"type": "object", "description": "Named numbers from the brief (a number or a list of "
                              "numbers, e.g. {\"revenue_q\": [4.1, 4.6, 5.2, 5.9]}): derived figures in the text are "
                              "checked against them, and {sum:revenue_q}, {average:...}, {change:...}, {first:...}, "
                              "{last:...}, {count:...} or {name} (a single number) in any text are computed from them.",
                              "additionalProperties": {"oneOf": [{"type": "number"}, {"type": "array", "minItems": 1,
                                                                                       "items": {"type": "number"}}]}},
                    "slides": {"type": "array", "minItems": 1, "items": {"oneOf": Variants}}}}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecSchema.SpecSchema")
            return {}


class kSpecCheck:
    """Spec checks: pattern limits, cross-field rules and (when jsonschema is installed) the schema."""

    @staticmethod
    def Values(Obj, Parts, Path=""):
        """Every (path, value) that a dotted LIMITS key ('columns.points.*') reaches in Obj."""
        if kS.ErrorMode:
            return []
        try:
            Head, Rest = Parts[0], Parts[1:]
            if Head == "*":
                Items = list(enumerate(Obj)) if isinstance(Obj, list) else []
            else:
                Items = [(Head, Obj.get(Head))] if isinstance(Obj, dict) and Head in Obj else []
            Out = []
            for K, V in Items:
                Pth = f"{Path}.{K}" if Path else str(K)
                if not Rest:
                    Out.append((Pth, V))
                elif isinstance(V, list) and Rest[0] != "*":
                    for J, El in enumerate(V):
                        Out += kSpecCheck.Values(El, Rest, f"{Pth}[{J}]")
                else:
                    Out += kSpecCheck.Values(V, Rest, Pth)
            return Out
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecCheck.Values")
            return []

    @staticmethod
    def CheckSpec(Spec):
        """Every pattern-limit and cross-field mistake in the spec, as messages."""
        if kS.ErrorMode:
            return []
        try:
            Errors = []
            for I, Sl in enumerate(Spec.get("slides", []), 1):
                Pat = Sl.get("pattern")
                if Pat not in LIMITS:
                    Errors.append(f"slide {I}: unknown pattern '{Pat}' (one of: {', '.join(LIMITS)})")
                    continue
                if Pat not in ("title", "quote", "section") and not Sl.get("title"):
                    Errors.append(f"slide {I} ({Pat}): needs a 'title' - write it as a claim")
                Kick = str(Sl.get("kicker") or "")
                if len(Kick) > KICKER_MAX or len(Kick.split()) > 3:
                    Errors.append(f"slide {I} ({Pat}): 'kicker' is a label - at most 3 words and {KICKER_MAX} "
                                  f"characters: '{Kick[:40]}'")
                for Key, Lim in LIMITS[Pat].items():
                    for Pth, Val in kSpecCheck.Values(Sl, Key.split(".")):
                        if isinstance(Lim, tuple):
                            N = len(Val) if isinstance(Val, list) else 0
                            if not Lim[0] <= N <= Lim[1]:
                                Hint = COUNT_HINTS.get(f"{Pat}.{Key}")
                                Errors.append(f"slide {I} ({Pat}): '{Pth}' has {N} items, needs {Lim[0]}-{Lim[1]}"
                                              + (f" - {Hint}" if Hint else ""))
                        elif isinstance(Val, str) and len(Val) > Lim:
                            Errors.append(f"slide {I} ({Pat}): '{Pth}' is {len(Val)} characters, max {Lim}: '{Val[:40]}…'")
                if Pat in NEEDED_LIST and NEEDED_LIST[Pat] not in Sl:
                    Errors.append(f"slide {I} ({Pat}): missing '{NEEDED_LIST[Pat]}'")
                if Pat == "image" and not os.path.exists(Sl.get("image", "")):
                    Errors.append(f"slide {I} (image): image file not found: {Sl.get('image')}")
                if Pat == "table":
                    Width = len(Sl.get("header", []))
                    for R, Row in enumerate(Sl.get("rows", [])):
                        if len(Row) != Width:
                            Errors.append(f"slide {I} (table): row {R + 1} has {len(Row)} cells, the header has {Width}")
                    Hr = Sl.get("highlight_row")
                    if Hr is not None and not (isinstance(Hr, int) and 0 <= Hr < len(Sl.get("rows", []))):
                        Errors.append(f"slide {I} (table): highlight_row {Hr!r} is not a row index (0-based)")
                if Pat == "compare":
                    Cols = Sl.get("columns", [])
                    Max = COMPARE_HEADING.get(len(Cols), 40)
                    for C, Col in enumerate(Cols):
                        Head = str(Col.get("heading", "")) if isinstance(Col, dict) else ""
                        if len(Head) > Max:
                            Errors.append(f"slide {I} (compare): 'columns.{C}.heading' is {len(Head)} characters; "
                                          f"with {len(Cols)} columns the limit is {Max}: '{Head[:40]}'")
                if Pat == "timeline":
                    Dated = [isinstance(E, dict) and bool(E.get("date")) for E in Sl.get("events", [])]
                    if any(Dated) and not all(Dated):
                        Errors.append(f"slide {I} (timeline): give every event a 'date', or none (undated stages are "
                                      "drawn as numbered steps)")
                if Pat == "quiz" and not any(isinstance(O, dict) and O.get("correct") is True
                                             for O in Sl.get("options", [])):
                    Errors.append(f"slide {I} (quiz): mark at least one option \"correct\": true")
                if Pat == "email":
                    for C, Co in enumerate(Sl.get("callouts", [])):
                        if not isinstance(Co, dict):
                            continue
                        if Co.get("target", "body") == "attachment" and not Sl.get("attachment"):
                            Errors.append(f"slide {I} (email): callout {C + 1} points at the attachment, but there is none")
                        Ln = Co.get("line", 0)
                        if Co.get("target", "body") == "body" and not (isinstance(Ln, int) and
                                                                       0 <= Ln < len(Sl.get("body", []))):
                            Errors.append(f"slide {I} (email): callout {C + 1} line {Ln!r} is not a body paragraph "
                                          f"(0-based, {len(Sl.get('body', []))} paragraphs)")
                for Mi, Mt in enumerate(Sl.get("metrics", []) + ([Sl["figure"]] if isinstance(Sl.get("figure"), dict)
                                                                   else [])):
                    Tl = Mt.get("trend_labels") if isinstance(Mt, dict) else None
                    if Tl and len(Tl) not in (2, len(Mt.get("trend") or [])):
                        Errors.append(f"slide {I} ({Pat}): trend_labels of metric {Mi} has {len(Tl)} labels; give one "
                                      f"per trend value ({len(Mt.get('trend') or [])}) or two (first and last)")
                    if Tl and not Mt.get("trend"):
                        Errors.append(f"slide {I} ({Pat}): trend_labels of metric {Mi} need a 'trend'")
                if Pat == "statement" and Sl.get("figure") and not Sl.get("decision"):
                    Errors.append(f"slide {I} (statement): 'figure' goes beside a 'decision'; without one use "
                                  "'points' or a big_number slide")
                if Pat == "kpi_chart" and "series" not in Sl and not any(
                        isinstance(Mt, dict) and Mt.get("trend") for Mt in Sl.get("metrics", [])):
                    Errors.append(f"slide {I} (kpi_chart): give 'categories' and 'series', or a 'trend' on a metric "
                                  "for the chart to draw")
                if Pat == "kpi_chart" and "series" in Sl and "categories" not in Sl:
                    Errors.append(f"slide {I} (kpi_chart): 'series' needs 'categories'")
                if Pat == "kpi_chart":
                    Hm = Sl.get("highlight_metric")
                    if Hm is not None and not (isinstance(Hm, int) and 0 <= Hm < len(Sl.get("metrics", []))):
                        Errors.append(f"slide {I} (kpi_chart): highlight_metric {Hm!r} is not a metric index (0-based)")
                if Pat == "cost_table":
                    for R, Row in enumerate(Sl.get("rows", [])):
                        if not isinstance(Row, dict) or isinstance(Row.get("amount"), bool) \
                                or not isinstance(Row.get("amount"), (int, float)):
                            Errors.append(f"slide {I} (cost_table): row {R + 1} needs a numeric 'amount'")
                if Pat in ("chart", "kpi_chart"):
                    Hi = Sl.get("highlight")
                    if isinstance(Hi, str) and Hi not in Sl.get("categories", []):
                        Errors.append(f"slide {I} ({Pat}): highlight '{Hi}' is not one of the categories")
                    if Sl.get("type", "column") not in CHART_TYPES:
                        Errors.append(f"slide {I} ({Pat}): type '{Sl.get('type')}' (one of: {', '.join(CHART_TYPES)})")
                ItemsKey = HIGHLIGHT_ITEMS.get(Pat)
                Hi = Sl.get("highlight")
                if ItemsKey and isinstance(Hi, int) and not 0 <= Hi < len(Sl.get(ItemsKey, [])):
                    Errors.append(f"slide {I} ({Pat}): highlight {Hi} is out of range (0-based, "
                                  f"{len(Sl.get(ItemsKey, []))} items)")
                if Pat in ("chart", "kpi_chart") and "series" in Sl:
                    Cats = len(Sl.get("categories", []))
                    for S in Sl.get("series", []):
                        if len(S.get("values", [])) != Cats:
                            Errors.append(f"slide {I} ({Pat}): series '{S.get('name')}' has {len(S.get('values', []))} "
                                          f"values for {Cats} categories")
            if len(str(Spec.get("footer") or "")) > 70:
                Errors.append("spec: 'footer' is over 70 characters; keep it to the deck name and audience")
            return Errors
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecCheck.CheckSpec")
            return []

    @staticmethod
    def SlideSet(Text, Count):
        """The spec slide numbers (1-based) that '--slides 1,3,5-7' names; a number outside the deck, or a part
        that is not a number or a range, is an input mistake."""
        if kS.ErrorMode:
            return set()
        try:
            Out = set()
            for Part in str(Text).replace(" ", "").split(","):
                Range = re.fullmatch(r"(\d+)(?:-(\d+))?", Part)
                if not Range:
                    raise ToolInputException(f"--slides: '{Part}' is not a slide number or a range (e.g. 1,3,5-7)")
                Lo, Hi = int(Range.group(1)), int(Range.group(2) or Range.group(1))
                if not 1 <= Lo <= Hi <= Count:
                    raise ToolInputException(f"--slides: '{Part}' is outside the deck (slides 1-{Count})")
                Out |= set(range(Lo, Hi + 1))
            return Out
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecCheck.SlideSet")
            return set()

    @staticmethod
    def ForSlides(Messages, Only):
        """The messages that concern the slides being built: all of them, or (with --slides) those about the
        chosen slides and the deck as a whole."""
        if kS.ErrorMode:
            return []
        try:
            if not Only:
                return Messages
            Out = []
            for Msg in Messages:
                Found = re.search(r"\bslide (\d+)\b|^slides/(\d+)", Msg)
                No = int(Found.group(1) or Found.group(2)) + (1 if Found and Found.group(2) else 0) if Found else 0
                if not Found or No in Only:
                    Out.append(Msg)
            return Out
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecCheck.ForSlides")
            return Messages

    @staticmethod
    def MonthIndex(Text):
        """0-11 for a date that is just a month name ('Mar', 'April', 'Mar 2027'), else None."""
        if kS.ErrorMode:
            return None
        try:
            Word = str(Text).strip().split(" ")[0].lower()[:3] if str(Text).strip() else ""
            return MONTHS.index(Word) if Word in MONTHS else None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecCheck.MonthIndex")
            return None

    @staticmethod
    def MonthRanges(Text):
        """Every month range in Text ('Jan-Jun', 'January to June', 'Mar–May 2027') as a set of (start, end)
        month indexes."""
        if kS.ErrorMode:
            return set()
        try:
            Month = r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?(?:\s+\d{4})?"
            Found = re.findall(r"\b" + Month + r"\s*(?:-|–|—|to|through|until|till)\s*" + Month + r"\b",
                               str(Text), re.I)
            return {(MONTHS.index(A.lower()), MONTHS.index(B.lower())) for A, B in Found}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecCheck.MonthRanges")
            return set()

    @staticmethod
    def WindowWarnings(Sl, I):
        """A dated timeline whose month falls outside its window: the slide's 'window' ('Jan-Jun'), else the one
        month range the slide itself states in its title or notes (a pilot 'January to June'; a range elsewhere in
        the deck, such as the quarter a review covers, is not this timeline's window). An event labelled
        'after ...' / 'before ...' is outside on purpose and passes."""
        if kS.ErrorMode:
            return []
        try:
            Ranges = kSpecCheck.MonthRanges(Sl.get("window", ""))
            Where = "its 'window'"
            if not Sl.get("window"):
                Ranges = kSpecCheck.MonthRanges(json.dumps(Sl, ensure_ascii=False))
                Where = "the range the slide states"
            if len(Ranges) != 1:
                return []
            Start, End = next(iter(Ranges))
            Out = []
            for E in Sl.get("events", []):
                Mi = kSpecCheck.MonthIndex(E.get("date", "")) if isinstance(E, dict) else None
                Label = str(E.get("label", "")) if isinstance(E, dict) else ""
                if Mi is None or re.search(r"\b(after|before|post|pre)\b", Label, re.I):
                    continue
                if (Mi - Start) % 12 > (End - Start) % 12:
                    Out.append(f"slide {I} (timeline): '{E.get('date')}: {Label}' falls outside "
                               f"{MONTHS[Start].title()}-{MONTHS[End].title()} ({Where}) - move it inside, or say "
                               "it is outside on purpose in its label ('After the pilot: ...')")
            return Out
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecCheck.WindowWarnings")
            return []

    @staticmethod
    def Warnings(Spec):
        """Content the spec can build but a reader will question, as messages: a success target that is not
        measurable, a monthly timeline that skips a month. They do not stop the build."""
        if kS.ErrorMode:
            return []
        try:
            Out = []
            for I, Sl in enumerate(Spec.get("slides", []), 1):
                Pat = Sl.get("pattern") if isinstance(Sl, dict) else None
                if Pat == "metrics":
                    for R, Row in enumerate(Sl.get("rows", [])):
                        Target = str(Row.get("target", "")) if isinstance(Row, dict) else ""
                        Metric = str(Row.get("metric", "")) if isinstance(Row, dict) else ""
                        if Target and not kRules.Measurable(Target, Metric):
                            Out.append(f"slide {I} (metrics): target '{Target}' of '{Metric or R}' is not "
                                       "measurable - give a number with a unit or comparator and say what is "
                                       "measured (e.g. 'below 30 %', '>= 95 % of Q4', '< 2 days lead time'); a time "
                                       "window alone ('<= last 6 months', 'by Q3') is not a target")
                if Pat == "statement" and not Sl.get("decision") and not Sl.get("points") and I > 1:
                    Out.append(f"slide {I} (statement): a claim with " + ("only a support line" if Sl.get("support")
                               else "nothing under it") + " reads as sparse - add 'points' (2-3 short facts, "
                               "steps or reasons) beside it")
                if Pat == "big_number" and not Sl.get("points") and not kSlidePatterns.Share(Sl) and I > 1:
                    Out.append(f"slide {I} (big_number): one number and one caption - add 'points' (what it costs, "
                               "what drives it) or state it as a share ('41' + '%', '14/20') for a dot grid")
                if Pat == "quiz" and Sl.get("reveal") == "slide":
                    Answer = str(Sl.get("answer_title") or "")
                    if not Answer or re.match(r"^\s*(?:the\s+)?(?:right\s+|correct\s+)?answers?\b\s*(?:is|are)?\s*[:\-"
                                              r"\u2013\u2014]?\s*[A-D](?:\s*(?:,|and|&)\s*[A-D])*\s*[.!]?\s*$",
                                              Answer, re.I) or kRules.LooksLikeLabel(Answer):
                        Out.append(f"slide {I} (quiz): the answer slide's title is " + (f"'{Answer}'" if Answer else
                                   "the default 'Answer: <letter>'") + " - a label, not a claim; set 'answer_title' "
                                   "to what the answer teaches (e.g. 'A look-alike sender is the first red flag')")
                if Pat in ("statement", "kpi", "next_steps") and Sl.get("decision"):
                    Ask = " ".join(str(X) for X in [Sl.get("decision"), Sl.get("support"), Sl.get("points"),
                                                    Sl.get("figure")] if X)
                    if not kRules.StatesCost(Ask):
                        Out.append(f"slide {I} ({Pat}): the ask does not state its cost - say what it costs or "
                                   "takes (money, time, FTE) in 'decision', a point or a 'figure' (e.g. '200 kSEK', "
                                   "'2 FTE for 6 months'); if the brief gives none, say the figure comes from "
                                   "Finance (reference/CONTENT.md)")
                if Pat == "bullets" and any(len(str(X)) > 80 for X in Sl.get("items", [])):
                    Out.append(f"slide {I} (bullets): a point over 80 characters keeps the slide a plain bulleted "
                               "list - shorten each point to a phrase (detail to the notes) for numbered bands, or "
                               "use 'compare', 'process' or 'statement' with 'points'")
                Out.extend(kSpecCheck.WindowWarnings(Sl, I) if Pat == "timeline" else [])
                if Pat == "timeline":
                    Months = [kSpecCheck.MonthIndex(E.get("date", "")) for E in Sl.get("events", [])
                              if isinstance(E, dict)]
                    if len(Months) >= 3 and None not in Months:
                        Gaps = [(B - A) % 12 for A, B in zip(Months, Months[1:])]
                        if Gaps.count(1) >= len(Gaps) - 1 and max(Gaps) > 1:  # monthly, but for one jump
                            J = Gaps.index(max(Gaps))
                            Out.append(f"slide {I} (timeline): the months jump from {MONTHS[Months[J]].title()} to "
                                       f"{MONTHS[Months[J + 1]].title()}; add the missing month or say why in "
                                       "the notes")
            return Out
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecCheck.Warnings")
            return []

    @staticmethod
    def ErrorPath(Error):
        """'a/b/0' for a jsonschema error's absolute path."""
        if kS.ErrorMode:
            return ""
        try:
            return "/".join(str(Part) for Part in Error.absolute_path)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecCheck.ErrorPath")
            return ""

    @staticmethod
    def SchemaErrors(Spec):
        """Validate against the JSON Schema when the jsonschema package is installed (optional). Each slide is
        checked against its own pattern's schema, so every misspelt or misplaced field is named."""
        if kS.ErrorMode:
            return []
        try:
            if importlib.util.find_spec("jsonschema") is None:
                return []  # optional dependency: without it check_spec alone validates
            import jsonschema
            Schema = kSpecSchema.SpecSchema()
            Variants = {V["properties"]["pattern"]["const"]: V
                        for V in Schema["properties"]["slides"]["items"]["oneOf"]}
            Top = dict(Schema, properties=dict(Schema["properties"], slides={"type": "array", "minItems": 1}))
            Out = [f"{kSpecCheck.ErrorPath(E) or 'spec'}: {E.message}"
                   for E in jsonschema.Draft202012Validator(Top).iter_errors(Spec)]
            for I, Sl in enumerate(Spec.get("slides", []), 1):
                V = Variants.get(Sl.get("pattern")) if isinstance(Sl, dict) else None
                if V is None:
                    continue  # unknown pattern: CheckSpec reports it
                for E in jsonschema.Draft202012Validator(V).iter_errors(Sl):
                    Where = kSpecCheck.ErrorPath(E)
                    if E.validator == "additionalProperties":
                        Extra = sorted(set(E.instance) - set(E.schema.get("properties", {})))
                        Out.append(f"slide {I} ({Sl['pattern']}): unknown field {', '.join(repr(X) for X in Extra)}"
                                   + (f" in {Where}" if Where else "") + " - see reference/BUILDER.md")
                    elif E.validator in ("maxLength", "minItems", "maxItems"):
                        continue  # CheckSpec already reports limits, with friendlier wording
                    else:
                        Out.append(f"slide {I} ({Sl['pattern']}): {Where or 'slide'}: {E.message}")
            return sorted(set(Out))
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecCheck.SchemaErrors")
            return []


class kSpecLoader:
    """Reads a spec file (JSON or YAML), normalises text fields and resolves relative paths."""

    @staticmethod
    def Texts(Obj, Key=None):
        """YAML turns 42 and 2024 into numbers; every text field must be a string."""
        if kS.ErrorMode:
            return Obj
        try:
            if Key in NUMERIC_KEYS:
                return Obj
            if isinstance(Obj, dict):
                return {K: kSpecLoader.Texts(V, K) for K, V in Obj.items()}
            if isinstance(Obj, list):
                return [kSpecLoader.Texts(V, Key) for V in Obj]
            if isinstance(Obj, (int, float)) and not isinstance(Obj, bool):
                return f"{Obj:g}" if isinstance(Obj, float) else str(Obj)
            return Obj
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecLoader.Texts")
            return Obj

    @staticmethod
    def Load(Path):
        """The spec at Path with image and template paths made absolute (they are relative to the spec)."""
        if kS.ErrorMode:
            return None
        try:
            if not os.path.isfile(Path):  # a mistyped spec path is the caller's input, not a bug: exit 2, no report
                raise ToolInputException(f"spec not found: {os.path.abspath(Path)}")
            with open(Path, encoding="utf-8") as Fh:
                if Path.lower().endswith((".yml", ".yaml")):
                    import yaml  # uvx --with pyyaml
                    Spec = yaml.safe_load(Fh)
                else:
                    Spec = json.load(Fh)
            Spec["slides"] = [kSpecLoader.Texts(Sl) for Sl in Spec.get("slides", [])]
            Base = os.path.dirname(os.path.abspath(Path))
            for Sl in Spec.get("slides", []):
                if Sl.get("image") and not os.path.isabs(Sl["image"]):
                    Sl["image"] = os.path.join(Base, Sl["image"])
            if Spec.get("template") and not os.path.isabs(Spec["template"]):
                Spec["template"] = os.path.join(Base, Spec["template"])
            return Spec
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kSpecLoader.Load(path={Path})")
            return None

    @staticmethod
    def PrintPlan(Spec):
        """The deck as a story: one line per slide - the claim, then the key fact from the notes."""
        if kS.ErrorMode:
            return
        try:
            Look = Spec.get("template") or Spec.get("direction", "clean-corporate")
            print(f"Deck plan ({len(Spec['slides'])} slides, look: {Look})")
            if Spec.get("footer"):
                print(f"Footer: {Spec['footer']}" + ("" if Spec.get("page_numbers", True) else " (no page numbers)"))
            if Spec.get("sources"):
                Src = [Spec["sources"]] if isinstance(Spec["sources"], str) else Spec["sources"]
                print(f"Sources (every slide without its own): {'; '.join(str(X) for X in Src)}")
            for Name, Val in kFigures.Facts(Spec).items():
                Shown = ", ".join(f"{X:g}" for X in Val) + f" (sum {kFigures.Compute('sum', Val, None)})" \
                    if isinstance(Val, list) else f"{Val:g}"
                print(f"Fact {Name}: {Shown}")
            print()
            Section = ""
            for N, Sl in enumerate(Spec["slides"], 1):
                Title = Sl.get("title") or Sl.get("quote", "")[:60]
                Notes = Sl.get("notes")
                Key = Notes.get("key_fact") if isinstance(Notes, dict) else (Notes or "").split("\n")[0]
                if Sl.get("pattern") == "section":
                    Section = Sl.get("kicker") or Sl.get("eyebrow") or ""
                Kick = Sl.get("kicker", Section) if Spec.get("kickers", True) else ""
                Kick = f"{Kick.upper()} / " if Kick and Sl.get("pattern") not in NO_KICKER else ""
                print(f"{N:>2}. [{Sl.get('pattern')}] {Kick}{Title}")
                if Sl.get("decision"):
                    print(f"      DECISION: {Sl['decision'][:100]}")
                if Key:
                    print(f"      {Key[:110]}")
                else:
                    print("      (no speaker notes yet)")
                Assume = Notes.get("assumptions") if isinstance(Notes, dict) else None
                if Assume:
                    Assume = [Assume] if isinstance(Assume, str) else Assume
                    print(f"      ASSUMES: {'; '.join(str(X) for X in Assume)[:104]}")
                if Sl.get("pattern") == "quiz" and Sl.get("reveal") == "slide":
                    print("      + an answer slide after it")
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSpecLoader.PrintPlan")
            return


class kDeckDesign:
    """Design directions and templates: the theme a deck is built on."""

    @staticmethod
    def LoadDirection(Name):
        """The direction called Name; an unknown name stops the build (exit 1)."""
        if kS.ErrorMode:
            return None
        try:
            for D in kSpecSchema.Directions():
                if D["id"] == Name:
                    return D
            raise ToolReportableException(f"unknown direction '{Name}'; see --list-directions")
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckDesign.LoadDirection")
            return None

    @staticmethod
    def Mix(Ca, Cb, T):
        """RRGGBB a fraction T of the way from Ca to Cb."""
        if kS.ErrorMode:
            return "000000"
        try:
            La, Lb = [int(Ca[I:I + 2], 16) for I in (0, 2, 4)], [int(Cb[I:I + 2], 16) for I in (0, 2, 4)]
            return "".join(f"{round(X + (Y - X) * T):02X}" for X, Y in zip(La, Lb))
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckDesign.Mix")
            return "000000"

    @staticmethod
    def QuietAndMuted(D):
        """The quiet card colour (the background with a breath of the accent and the text) and the muted text
        colour, chosen so muted text stays at 4.5:1 or better on both the background and the card, and accent
        figures at 3:1 or better on the card."""
        if kS.ErrorMode:
            return None, None
        try:
            Muted = D["muted"]
            Tint = kDeckDesign.Mix(D["background"], D["accent"], 0.07)  # cards carry a breath of the accent, not grey
            for T in (0.05, 0.04, 0.03, 0.02, 0.0):
                Quiet = kDeckDesign.Mix(Tint, D["text"], T)
                if kRules.ContrastRatio(Muted, Quiet) >= 4.5 and kRules.ContrastRatio(D["accent"], Quiet) >= 3:
                    return Quiet, Muted
            Quiet = kDeckDesign.Mix(Tint, D["text"], 0.03)
            for K in range(1, 11):  # darken (or lighten, on dark themes) the muted colour towards the text colour
                Muted = kDeckDesign.Mix(D["muted"], D["text"], K / 10)
                if kRules.ContrastRatio(Muted, Quiet) >= 4.5:
                    break
            return Quiet, Muted
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckDesign.QuietAndMuted")
            return None, None

    @staticmethod
    def ApplyDirection(Prs, D):
        """Write the direction into the theme so everything below uses theme colours and fonts."""
        if kS.ErrorMode:
            return
        try:
            Part = Prs.slide_master.part.part_related_by(RT.THEME)
            Theme = etree.fromstring(Part.blob)
            Scheme = Theme.find(".//a:clrScheme", NS)
            Quiet, Muted = kDeckDesign.QuietAndMuted(D)
            Mix = kDeckDesign.Mix
            Slots = {"dk1": D["text"], "lt1": D["background"], "dk2": Muted,
                     "lt2": Quiet, "accent1": D["accent"],
                     "accent2": Mix(D["accent"], D["background"], 0.45), "accent3": Muted,
                     "accent4": Mix(D["accent"], D["text"], 0.4), "accent5": Mix(D["muted"], D["background"], 0.4),
                     "accent6": Mix(D["accent"], D["background"], 0.7)}
            for Slot, HexV in Slots.items():
                Node = Scheme.find(f"a:{Slot}", NS)
                for Child in list(Node):
                    Node.remove(Child)
                etree.SubElement(Node, f"{{{A}}}srgbClr", val=HexV)
            Scheme.set("name", D["id"])
            Fonts = Theme.find(".//a:fontScheme", NS)
            Fonts.find("a:majorFont/a:latin", NS).set("typeface", D["heading"])
            Fonts.find("a:minorFont/a:latin", NS).set("typeface", D["body"])
            Fonts.set("name", D["id"])
            Part._blob = etree.tostring(Theme, xml_declaration=True, encoding="UTF-8", standalone=True)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckDesign.ApplyDirection")
            return

    @staticmethod
    def OpenTemplate(Path):
        """A template (.pptx, or .potx/.potm with its content type patched) as a Presentation."""
        if kS.ErrorMode:
            return None
        try:
            if Path.lower().endswith((".potx", ".potm")):
                from extract_theme import kThemeReader
                return kThemeReader.OpenAny(Path)
            return Presentation(Path)
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kDeckDesign.OpenTemplate(path={Path})")
            return None

    @staticmethod
    def TitleOnlyLayout(Prs):
        """'Title Only', else the layout with a title and the fewest other text placeholders."""
        if kS.ErrorMode:
            return None
        try:
            for Layout in Prs.slide_layouts:
                if Layout.name.lower().replace("-", " ").startswith("title only"):
                    return Layout
            Best = None
            for Layout in Prs.slide_layouts:
                Types = [Ph.placeholder_format.type for Ph in Layout.placeholders]
                if any(T in TITLE_TYPES for T in Types):
                    Others = sum(1 for T in Types if T not in TITLE_TYPES and T not in FOOTER_TYPES)
                    if Best is None or Others < Best[0]:
                        Best = (Others, Layout)
            if Best is None:
                raise ToolReportableException("the template has no layout with a title placeholder")
            return Best[1]
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckDesign.TitleOnlyLayout")
            return None


class kSlideText:
    """Stateless text helpers: speaker notes, alt text."""

    @staticmethod
    def Alt(Shape, Text):
        """Set a shape's alt text (descr)."""
        if kS.ErrorMode:
            return
        try:
            Shape._element.find(".//p:cNvPr", NS).set("descr", Text)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlideText.Alt")
            return

    @staticmethod
    def QaLine(Q):
        """One Q&A item as notes text."""
        if kS.ErrorMode:
            return ""
        try:
            return f"Q: {Q.get('q', '')}\nA: {Q.get('a', '')}" if isinstance(Q, dict) else f"- {Q}"
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlideText.QaLine")
            return ""

    @staticmethod
    def NotesText(Notes):
        """Speaker notes as text: a string as is, or the structured fields in a fixed order."""
        if kS.ErrorMode:
            return ""
        try:
            if not Notes:
                return ""
            if isinstance(Notes, str):
                return Notes
            Out = []
            if Notes.get("key_fact"):
                Out.append(f"KEY FACT: {Notes['key_fact']}")
            for Label, Key in (("FACTS", "facts"), ("ASSUMPTIONS", "assumptions"), ("PITFALLS", "pitfalls"),
                               ("SOURCES", "sources")):
                Items = Notes.get(Key)
                if Items:
                    Items = [Items] if isinstance(Items, str) else Items
                    Out.append(f"{Label}:\n" + "\n".join(f"- {X}" for X in Items))
            Qa = Notes.get("qa")
            if Qa:
                Qa = [Qa] if isinstance(Qa, dict) else Qa
                Out.append("Q&A:\n" + "\n".join(kSlideText.QaLine(Q) for Q in Qa))
            return "\n\n".join(Out)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlideText.NotesText")
            return ""

    @staticmethod
    def ChartAlt(Sl):
        """Alt text that reads a chart's data out."""
        if kS.ErrorMode:
            return ""
        try:
            Parts = []
            for Ser in Sl["series"]:
                Pairs = ", ".join(f"{C} {V:g}" if isinstance(V, (int, float)) else f"{C} n/a"
                                  for C, V in zip(Sl["categories"], Ser["values"]))
                Parts.append(f"{Ser['name']}: {Pairs}")
            return f"{Sl.get('type', 'column').title()} chart. " + "; ".join(Parts) + "."
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlideText.ChartAlt")
            return ""


class kDeckBuilder:
    """One deck being built: the presentation, its theme fonts, the grid scale and the fit report."""

    def __init__(self, Spec):
        try:
            self.Spec = Spec
            self.Prs = None
            self.Themed = False   # True when a direction was applied (we own fonts/colours)
            self.Layout = None
            self.Major = self.Minor = None
            self.Problems = []    # text that could not be made to fit
            self.Issues = []      # the same, structured: {spec, slide, id, shape, code, severity, message}
            self.LintItems = []   # every lint finding on the built deck (with the spec slide it came from)
            self.SpecOf = {}      # deck slide number -> spec slide number (a quiz answer slide shares its quiz's)
            self.SpecNo = 0       # the spec slide being built
            self.NoNotes = []     # slides the spec gave no speaker notes
            self.SlideNo, self.SlideId = 0, ""
            self.Section = ""     # the current section's eyebrow: the default kicker
            self.TitleInk = None  # top of the title's text (1440 grid) on the slide being built, for the kicker
            self.Kicker = ""      # the kicker of the slide being built, so the title leaves room for it
            self.Extra = []       # notes a pattern adds while it lays the slide out (a trend it had no room for)
            self.Numbers = []     # the deck's slide number of each slide built (--slides builds only some)
            self.Reserve = False  # the slide being built has a footer band (content must end above it)
            self.KickerWidth = 0  # a pattern with a column beside its claim (a statement's points) narrows the kicker
            self._kx = self._ky = 1.0  # template mode: slide size / 1440 x 810, so the grid follows the template
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.__init__")

    def Report(self, Shape, Code, Message):
        """Record text that does not fit on the slide being built: a `fit:` line and a structured issue."""
        if kS.ErrorMode:
            return
        try:
            self.Problems.append(f"slide {self.SlideNo} ({self.SlideId}): {Message}")
            Name = str(Shape or "").strip("'")
            self.Issues.append({"spec": self.SpecNo, "slide": self.SlideNo, "id": self.SlideId,
                                "shape": "" if Name in ("text", "title", "value") else Name, "code": Code,
                                "severity": "error", "message": Message})
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Report")
            return

    def X(self, V):
        """A horizontal 1440-grid length in points on this deck."""
        if kS.ErrorMode:
            return Pt(0)
        try:
            return Pt(V * self._kx)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.X")
            return Pt(0)

    def Y(self, V):
        """A vertical 1440-grid length in points on this deck."""
        if kS.ErrorMode:
            return Pt(0)
        try:
            return Pt(V * self._ky)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Y")
            return Pt(0)

    def TypeScale(self):
        """The factor type sizes are scaled by: never below two-thirds, so body text stays >= 18 pt."""
        if kS.ErrorMode:
            return 1.0
        try:
            return max(min(self._kx, self._ky), 2 / 3)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.TypeScale")
            return 1.0

    def F(self, V):
        """A font size on the 1440 grid, scaled to this deck."""
        if kS.ErrorMode:
            return Pt(0)
        try:
            return Pt(round(V * self.TypeScale(), 1))
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.F")
            return Pt(0)

    def Fit(self, Lines, Size, Wd, Ht, Heading=False, Bold=False, What="text", Spacing=0.5, Floor=None):
        """Largest size <= Size at which Lines fit a Wd x Ht box (1440-grid points), never below the
        floor. Sentences (4+ words) start at the floor at least, so a small role size can't put them below it;
        Floor (a body role: points, support, caption) holds every line, however short, at that floor.
        Records a problem when even the floor doesn't fit."""
        if kS.ErrorMode:
            return Size
        try:
            Sentences = any(len(str(X).split()) >= 4 for X in Lines)
            if (Sentences and not Heading) or Floor:
                Size = max(Size, Floor or FLOOR)
            Smallest = (min(Size, 40) if Heading else (FLOOR if Sentences else LABEL_MIN)) if Size > LABEL_MIN else Size
            Smallest = max(Smallest, Floor) if Floor else Smallest
            Family = self.Major if Heading else self.Minor
            K = self.TypeScale()
            Cur = Size
            while True:
                Paras = [(str(X), Family, Cur * K, Bold or Heading, Cur * K * Spacing) for X in Lines]
                Need, Widest, _ = kMeasure.LooseHeight(Paras, (Wd - 1) * self._kx)  # LibreOffice's face too
                if (Need <= Ht * self._ky * 1.02 and Widest <= Wd * self._kx) or Cur <= Smallest:
                    break
                Cur = max(Smallest, Cur - 2)
            if Need > Ht * self._ky * 1.08 or Widest > Wd * self._kx + 1:
                self.Report(What, "fit", f"{What} does not fit even at {Cur:g} pt "
                                     f"(needs ~{Need / self._ky:.0f} pt of {Ht:.0f}); cut words or split the slide")
            return Cur
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Fit")
            return Size

    def FitAll(self, Groups, Size, Wd, Ht, Heading=False, Bold=False, What="text", Spacing=0.5, Floor=None):
        """One size for sibling boxes (cards in a row): the smallest of each group's own fit, so they match."""
        if kS.ErrorMode:
            return Size
        try:
            Sizes = [self.Fit(G if isinstance(G, list) else [G], Size, Wd, Ht, Heading, Bold, What, Spacing, Floor)
                     for G in Groups if G]
            return min(Sizes) if Sizes else Size
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.FitAll")
            return Size

    def FitLine(self, Texts, Size, Wd, Heading=True, Bold=True, Smallest=40, What="value"):
        """Largest size <= Size at which every text in Texts stays on one line Wd wide (big values, numbers), in
        its own font AND in the face LibreOffice substitutes for it. A text still wider than Wd at Smallest would
        wrap (the unit onto the label below): reported as unfit text."""
        if kS.ErrorMode:
            return Size
        try:
            Cur = Size
            while Cur > Smallest and max(self.LineWidth(T, Cur, Heading, Bold) for T in Texts) \
                    > Wd * 0.9 * self._kx:  # room for renderers' wider fonts
                Cur -= 2
            Widest = max(self.LineWidth(T, Cur, Heading, Bold) for T in Texts) if Texts else 0
            if Widest > Wd * self._kx:
                self.Report(What, "fit_line", f"{What} does not fit on one line even "
                                     f"at {Cur:g} pt (needs ~{Widest / self._kx:.0f} pt of {Wd:.0f}); shorten it "
                                     "(e.g. '19.8 M' with the unit in the label)")
            return Cur
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.FitLine")
            return Size

    def LineWidth(self, Text, Size, Heading=True, Bold=True):
        """Width (points on this deck) of Text set on one line at Size: the wider of its theme font and the face
        LibreOffice substitutes for it (kMeasure.Substitute), as glued by Text() (a unit stays with its number)."""
        if kS.ErrorMode:
            return 0.0
        try:
            K = self.TypeScale()
            Family = self.Major if Heading else self.Minor
            Txt = UNIT_GLUE.sub(NBSP, str(Text))
            Sub = kMeasure.Substitute(Family)
            return max(kMeasure.SafeWidth(Txt, Face, Size * K, Bold or Heading) for Face in (Family, Sub) if Face)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.LineWidth")
            return 0.0

    def Need(self, Lines, Size, Wd, Heading=False, Bold=False, Spacing=0.5):
        """Height (1440-grid points) that Lines take at Size in a box Wd wide, with a little room for renderers
        that set lines looser (LibreOffice)."""
        if kS.ErrorMode:
            return 0.0
        try:
            Lines = Lines if isinstance(Lines, list) else [Lines]
            K = self.TypeScale()
            Family = self.Major if Heading else self.Minor
            Paras = [(str(X), Family, Size * K, Bold or Heading, Size * K * Spacing) for X in Lines]
            Height, _, _ = kMeasure.LooseHeight(Paras, (Wd - 1) * 0.92 * self._kx)  # slack, and LibreOffice's face
            return Height / self._ky + Size * 0.3
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Need")
            return 0.0

    def Wrapped(self, Text, Size, Wd, Heading=False, Bold=False):
        """Height (1440 grid) one paragraph really takes at Size in a box Wd wide: its wrapped lines (in the wider
        of its own font and the face LibreOffice substitutes) times the line height, plus the box's insets - so a
        paragraph cut by a few words gets a shorter box, unlike Need's fixed slack."""
        if kS.ErrorMode:
            return 0.0
        try:
            K = self.TypeScale()
            Family = self.Major if Heading else self.Minor
            Lines = kMeasure.LooseLines(UNIT_GLUE.sub(NBSP, str(Text)), Family, Size * K, (Wd - 1) * self._kx,
                                        Bold or Heading)
            return Lines * Size * K * LINE_HEIGHT / self._ky + 6
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Wrapped")
            return 0.0

    def NewSlide(self, Sl, N):
        """A new title-only slide named after the spec id, with every non-title placeholder removed."""
        if kS.ErrorMode:
            return None
        try:
            S = self.Prs.slides.add_slide(self.Layout)
            S._element.find("p:cSld", NS).set("name", Sl.get("id", f"s{N:02d}"))
            if self.Themed:
                S.background.fill.solid()
                S.background.fill.fore_color.theme_color = BG
            for Ph in list(S.placeholders):
                if Ph.placeholder_format.type not in TITLE_TYPES:
                    Ph._element.getparent().remove(Ph._element)
            if S.shapes.title is None:
                raise ToolReportableException(f"layout '{self.Layout.name}' has no title placeholder; "
                                              "the template needs a layout with one")
            return S
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.NewSlide")
            return None

    def BalancedWidth(self, Text, Size, Wd, Heading=True):
        """The width (1440 grid) at which Text wraps without a lone last word, in its own font AND in the face
        LibreOffice substitutes for it (kMeasure.Substitute): Wd itself when every face fits one line or already
        ends on two words, else the narrowest width that keeps each face's line count and moves a second word down
        (a balanced title instead of a widow)."""
        if kS.ErrorMode:
            return Wd
        try:
            Family = self.Major if Heading else self.Minor
            Sub = kMeasure.Substitute(Family)
            Families = (Family, Sub) if Sub else (Family,)
            Pt = Size * self.TypeScale()
            Counts = [len(kMeasure.LineWords(str(Text), F, Pt, (Wd - 1) * self._kx, True)) for F in Families]
            if not self.AnyWidow(Text, Families, Pt, (Wd - 1, Wd + 2)):  # either side of a renderer's insets
                return Wd
            Try = Wd
            while Try > Wd * 0.55:
                Try -= 8
                Now = [len(kMeasure.LineWords(str(Text), F, Pt, (Try - 1) * self._kx, True)) for F in Families]
                if any(N > C for N, C in zip(Now, Counts)):
                    break
                if not self.AnyWidow(Text, Families, Pt, (Try - 17, Try - 1)):
                    return Try - 16  # a little slack so a wider renderer keeps the same breaks
            return Wd
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.BalancedWidth")
            return Wd

    def AnyWidow(self, Text, Families, SizePt, Widths):
        """True when Text, wrapped in any of Families at any of Widths (1440 grid), leaves one word on its last
        line."""
        if kS.ErrorMode:
            return False
        try:
            for Family in Families:
                for W in Widths:
                    Lines = kMeasure.LineWords(str(Text), Family, SizePt, W * self._kx, True)
                    if len(Lines) >= 2 and len(Lines[-1]) == 1:
                        return True
            return False
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.AnyWidow")
            return False

    def TitleLines(self, Text, Size, Wd, Loose=False):
        """How many lines a title takes at Size in a box Wd wide (1440 grid), from the font's metrics; Loose: in
        the widest renderer (the face LibreOffice substitutes when the theme font is not installed)."""
        if kS.ErrorMode:
            return 1
        try:
            if Loose:
                return kMeasure.LooseLines(str(Text), self.Major, Size * self.TypeScale(), (Wd - 16) * self._kx, True)
            return max(1, len(kMeasure.LineWords(str(Text), self.Major, Size * self.TypeScale(), (Wd - 16) * self._kx,
                                                 True)))
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.TitleLines")
            return 1

    def Widow(self, Text, Size, Wd):
        """True when the title, wrapped in a box Wd wide in its own font or in the face LibreOffice substitutes
        for it (whatever fonts this machine has), leaves one word on its last line."""
        if kS.ErrorMode:
            return False
        try:
            Sub = kMeasure.Substitute(self.Major)
            for Family in (self.Major, Sub) if Sub else (self.Major,):
                for Inner in (Wd - 18, Wd - 12):  # either side of the renderers' insets
                    Lines = kMeasure.LineWords(str(Text), Family, Size * self.TypeScale(), Inner * self._kx, True)
                    if len(Lines) >= 2 and len(Lines[-1]) == 1:
                        return True
            return False
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Widow")
            return False

    def MaxTitleLines(self, Size):
        """The most lines lint_deck.py accepts for a title at Size (1440 grid): 3 for a display title (40 pt or
        more on a 960-pt slide, 60 pt on Full HD), else 2 - so a wide theme font (Verdana) shrinks it to fit."""
        if kS.ErrorMode:
            return 3
        try:
            return 3 if Size * self.TypeScale() >= 40 * 1.5 * self._kx else 2
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.MaxTitleLines")
            return 3

    def TitleFits(self, Room, Size, Lines):
        """True when Lines title lines at Size fit Room (1440 grid; None = any height)."""
        if kS.ErrorMode:
            return True
        try:
            return Room is None or Lines * Size * self.TypeScale() * LINE_HEIGHT / self._ky <= Room * 1.02
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.TitleFits")
            return True

    def AgreeLines(self, Text, Size, Wd, Room=None):
        """(size, width, lines) at which the title wraps to the same number of lines in its own font and in the
        face LibreOffice substitutes for it, so the kicker above it sits right in both: a size up to 4 pt smaller,
        else a narrower box (the own font wraps like the substitute), else a smaller size at which the loose count
        still fits. Every answer is measured to fit Room (the title's height, 1440 grid) and to leave no widow."""
        if kS.ErrorMode:
            return Size, Wd, 1
        try:
            Own, Loose = self.TitleLines(Text, Size, Wd), self.TitleLines(Text, Size, Wd, True)
            if Own == Loose:
                return Size, Wd, Own
            for Try in (Size - 2, Size - 4):
                if Try >= 40 and self.TitleLines(Text, Try, Wd) == self.TitleLines(Text, Try, Wd, True) \
                        and not self.Widow(Text, Try, Wd) \
                        and self.TitleFits(Room, Try, self.TitleLines(Text, Try, Wd)):
                    return Try, Wd, self.TitleLines(Text, Try, Wd)
            Nw = Wd
            while Nw > Wd * 0.6:
                Nw -= 16
                Own, Loose = self.TitleLines(Text, Size, Nw), self.TitleLines(Text, Size, Nw, True)
                if Own == Loose and not self.Widow(Text, Size, Nw) and self.TitleFits(Room, Size, Own):
                    return Size, Nw, Own
                if Own > Loose:
                    break
            Try = Size  # no agreement that fits: the largest size whose loose wrap fits the room without a widow
            while Room is not None and Try > 36:
                Lines = max(self.TitleLines(Text, Try, Wd), self.TitleLines(Text, Try, Wd, True))
                if self.TitleFits(Room, Try, Lines) and not self.Widow(Text, Try, Wd):
                    return Try, Wd, Lines
                Try -= 2
            return Size, Wd, max(self.TitleLines(Text, Size, Wd), self.TitleLines(Text, Size, Wd, True))
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.AgreeLines")
            return Size, Wd, 1

    def Title(self, S, Text, Size=None, Top=TITLE_TOP, Height=TITLE_H, Align=PP_ALIGN.LEFT, Colour=TEXT,
              Width=None):
        """Fill and place the slide's title placeholder: fitted, balanced (no lone last word), bottom-anchored."""
        if kS.ErrorMode:
            return None
        try:
            Text = UNIT_GLUE.sub(NBSP, str(Text))  # '44 %' never breaks between the number and its unit
            Box = Width or W - 2 * M
            Head = Top == TITLE_TOP  # a content slide's title: it may take two lines and has the kicker above it
            Room = Height - 8 + (TITLE_DROP if Head else 0)
            Size = self.Fit([Text], Size or SIZE["title"], Box - 15, Room, Heading=True, What="title")
            Wd = self.BalancedWidth(Text, Size, Box - 15) + 15
            if Wd < Box:  # Fit again at the balanced width: the same size must still fit the height
                Size = self.Fit([Text], Size, Wd - 15, Room, Heading=True, What="title")
            Floor = Size - 8  # no width balances it in every face (the substitute wraps a word sooner): a size
            while self.Widow(Text, Size, Wd) and Size - 2 >= max(Floor, 36):  # up to 8 pt smaller that does
                if self.TitleLines(Text, Size - 2, Wd) > self.MaxTitleLines(Size - 2):
                    break  # below display size a third line is not allowed: a widow beats a claim shrunk to 2 lines
                Size -= 2
                Wd = min(Box, self.BalancedWidth(Text, Size, Box - 15) + 15)
            Lines = self.TitleLines(Text, Size, Wd)
            if Head:  # both renderers must wrap it alike, or a bottom-anchored title grows up into its kicker
                Size, Wd, Lines = self.AgreeLines(Text, Size, Wd, Room)
            Bottom = Top + Height + (TITLE_DROP if Head and Lines > 1 else 0)
            # the kicker and the title are one block: measured with the line height and the line count of the
            # loosest renderer, the title's ink must leave room for the kicker above it, or the title shrinks
            while Head and self.Kicker and Size > 36 and \
                    Bottom - 4 - Lines * Size * LOOSE_LINE - KICKER_GAP - KICKER_H < KICKER_TOP:
                Size -= 2
                if self.Widow(Text, Size, Wd):  # the smaller size wraps differently: balance it again
                    Wd = min(Box, self.BalancedWidth(Text, Size, Box - 15) + 15)
                Lines = self.TitleLines(Text, Size, Wd, True)
                Bottom = Top + Height + (TITLE_DROP if Lines > 1 else 0)
            while Size > 36 and self.TitleLines(Text, Size, Wd) > self.MaxTitleLines(Size):
                Size -= 2  # lint's headline_too_long: a display title takes at most 3 lines, any other 2
                Wd = min(Box, self.BalancedWidth(Text, Size, Box - 15) + 15)
                Lines = self.TitleLines(Text, Size, Wd, Head)
            T = S.shapes.title
            T.left, T.top, T.width, T.height = self.X(M), self.Y(Top), self.X(Wd), self.Y(Bottom - Top)
            self.TitleInk = Bottom - 4 - Lines * Size * LOOSE_LINE - KICKER_GAP if Head else None
            T.text = Text
            Tf = T.text_frame
            Tf.word_wrap = True
            Tf.vertical_anchor = MSO_ANCHOR.BOTTOM
            for Para in Tf.paragraphs:
                Para.alignment = Align
                for R in Para.runs:
                    R.font.size = self.F(Size or SIZE["title"])
                    R.font.bold = True
                    if self.Themed:
                        R.font.color.theme_color = Colour
                        R.font.name = "+mj-lt"
            return T
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Title")
            return None

    def Text(self, S, Name, Txt, Lx, Ty, Wd, Ht, Size, Colour=TEXT, Bold=False, Align=PP_ALIGN.LEFT,
             Anchor=MSO_ANCHOR.TOP, Heading=False, Spacing=0.5, Floor=None):
        """A named text box, its size fitted to the box (Size is the ceiling; Spacing is the gap before each
        paragraph after the first, as a fraction of the size; Floor holds a body role at the floor however short
        its lines). A box that reaches into the footer band of a slide with a footer is reported as unfit."""
        if kS.ErrorMode:
            return None
        try:
            Lines = [UNIT_GLUE.sub(NBSP, str(X)) for X in (Txt if isinstance(Txt, list) else [Txt])]
            Size = self.Fit(Lines, Size, Wd, Ht, Heading=Heading, Bold=Bold, What=f"'{Name}'", Spacing=Spacing,
                            Floor=Floor)
            if self.Reserve and Ty + Ht > FOOTER_TOP - 4:
                self.Report(Name, "footer_band", f"'{Name}' reaches {Ty + Ht:.0f} pt, into "
                                     f"the footer band (content ends at {BODY_BOTTOM}); cut words or split the slide")
            Tb = S.shapes.add_textbox(self.X(Lx), self.Y(Ty), self.X(Wd), self.Y(Ht))
            Tb.name = Name
            Tf = Tb.text_frame
            Tf.word_wrap = True
            Tf.vertical_anchor = Anchor
            Tf.margin_left = Pt(0)
            Tf.margin_right = Pt(0)
            Tf.margin_top = Tf.margin_bottom = Pt(2)
            for I, Line in enumerate(Lines):
                Para = Tf.paragraphs[0] if I == 0 else Tf.add_paragraph()
                Marker = Line[:2] if Line.startswith(("\u2022 ", "\u25A0 ")) else ""
                if Marker:  # a bullet: the glyph in the accent, the words in the text colour
                    Lead = Para.add_run()
                    Lead.text = Marker
                    Lead.font.color.theme_color = ACCENT if Colour == TEXT else Colour
                    Body = Para.add_run()
                    Body.text = Line[2:]
                else:
                    Para.text = Line
                Para.alignment = Align
                if I:
                    Para.space_before = self.F(Size * Spacing)
                for Ri, R in enumerate(Para.runs):
                    R.font.size, R.font.bold = self.F(Size), Bold
                    if not (Marker and Ri == 0):
                        R.font.color.theme_color = Colour
                    if self.Themed:
                        R.font.name = "+mj-lt" if Heading else "+mn-lt"
            return Tb
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Text")
            return None

    def ShapeText(self, Shape, Text, Size, Colour=BG, Bold=True, Heading=True, Inset=0):
        """Text written into a shape itself (a badge's number, a chip's level), centred both ways."""
        if kS.ErrorMode:
            return None
        try:
            Tf = Shape.text_frame
            Tf.word_wrap = True
            Tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            Tf.margin_top = Tf.margin_bottom = Pt(0)
            Tf.margin_left = Tf.margin_right = self.X(Inset)
            Tf.text = Text
            Para = Tf.paragraphs[0]
            Para.alignment = PP_ALIGN.CENTER
            for R in Para.runs:
                R.font.size, R.font.bold = self.F(Size), Bold
                R.font.color.theme_color = Colour
                if self.Themed:
                    R.font.name = "+mj-lt" if Heading else "+mn-lt"
            return Shape
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.ShapeText")
            return None

    def Backdrop(self, S, Name, Colour):
        """A full-slide panel behind everything else on the slide (the title placeholder included)."""
        if kS.ErrorMode:
            return None
        try:
            R = self.Rect(S, Name, 0, 0, W, H, Colour)
            Tree = S.shapes._spTree
            Tree.remove(R._element)
            Tree.insert(2, R._element)  # after nvGrpSpPr and grpSpPr: the first drawn, so the back of the stack
            return R
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Backdrop")
            return None

    def Card(self, S, Name, Lx, Ty, Wd, Ht, On=False, Edge="top", EdgeOn=None):
        """The card language: a tinted card with an accent edge (top or left); the highlighted one filled with
        the accent (its text then goes in the background colour). EdgeOn marks the lead card by its edge only."""
        if kS.ErrorMode:
            return None
        try:
            Card = self.Rect(S, Name, Lx, Ty, Wd, Ht, ACCENT if On else QUIET)
            if not On and Edge:
                Lead = ACCENT if EdgeOn or EdgeOn is None else SOFT
                if Edge == "left":
                    self.Rect(S, f"{Name}Edge", Lx, Ty, EDGE, Ht, Lead)
                else:
                    self.Rect(S, f"{Name}Edge", Lx, Ty, Wd, EDGE, Lead)
            return Card
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Card")
            return None

    def KickerFor(self, Sl):
        """The kicker this slide shows ("" for none): its own, else the current section's eyebrow."""
        if kS.ErrorMode:
            return ""
        try:
            if not self.Spec.get("kickers", True) or Sl.get("pattern") in NO_KICKER:
                return ""
            return str(Sl.get("kicker", self.Section) or "")
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.KickerFor")
            return ""

    def Chrome(self, S, Sl, N):
        """The deck's furniture on a content slide: the kicker above the title, the footer and the page number."""
        if kS.ErrorMode:
            return
        try:
            Pat = Sl.get("pattern")
            if Pat == "section":
                self.Section = str(Sl.get("kicker") or Sl.get("eyebrow") or "")
            if Pat in NO_CHROME:
                return
            Kick = self.KickerFor(Sl)
            if Kick and self.TitleInk is not None:
                Tb = self.Text(S, "Kicker", Kick.upper(), M, max(KICKER_TOP, self.TitleInk - KICKER_H),
                               self.KickerWidth or W - 2 * M, KICKER_H,
                               KICKER_SIZE, ACCENT, Bold=True, Anchor=MSO_ANCHOR.BOTTOM)
                for R in Tb.text_frame.paragraphs[0].runs:
                    R.font._rPr.set("spc", "200")  # tracked caps read as a label, not a sentence
            Foot = str(self.Spec.get("footer") or "")
            if Foot:
                self.Text(S, "Footer", Foot, M, FOOTER_TOP, W - 2 * M - 160, FOOTER_H, FOOTER_SIZE, MUTED,
                          Anchor=MSO_ANCHOR.MIDDLE)
            if self.Spec.get("page_numbers", True):
                self.Text(S, "PageNumber", str(N), W - M - 120, FOOTER_TOP, 120, FOOTER_H, FOOTER_SIZE, MUTED,
                          Bold=True, Align=PP_ALIGN.RIGHT, Anchor=MSO_ANCHOR.MIDDLE)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Chrome")
            return

    def HasFooter(self, Sl):
        """True when Chrome will draw a footer or page number on this slide: its band is reserved."""
        if kS.ErrorMode:
            return False
        try:
            return Sl.get("pattern") not in NO_CHROME and bool(self.Spec.get("footer")
                                                                or self.Spec.get("page_numbers", True))
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.HasFooter")
            return False

    def Rect(self, S, Name, Lx, Ty, Wd, Ht, Colour=QUIET, Shape=MSO_SHAPE.RECTANGLE, Line=None):
        """A named, theme-filled shape with no shadow; no outline unless Line (a theme colour) is given."""
        if kS.ErrorMode:
            return None
        try:
            R = S.shapes.add_shape(Shape, self.X(Lx), self.Y(Ty), self.X(Wd), self.Y(Ht))
            R.name = Name
            R.fill.solid()
            R.fill.fore_color.theme_color = Colour
            if Line is None:
                R.line.fill.background()
            else:
                R.line.color.theme_color = Line
                R.line.width = Pt(1.5)
            R.shadow.inherit = False
            Ref = R._element.find(".//a:effectRef", NS)
            if Ref is not None:  # the shape style's theme effect is a shadow in some renderers: none, flat cards
                Ref.set("idx", "0")
            return R
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Rect")
            return None

    def OpenDeck(self):
        """The template with its slides dropped, or a blank deck themed with the direction; sets the scale."""
        if kS.ErrorMode:
            return
        try:
            if self.Spec.get("template"):
                self.Prs = kDeckDesign.OpenTemplate(self.Spec["template"])
                for Sld in list(self.Prs.slides._sldIdLst):  # start from the template's masters, not its slides
                    self.Prs.part.drop_rel(Sld.rId)
                    self.Prs.slides._sldIdLst.remove(Sld)
                self.Themed = False
            else:
                self.Prs = Presentation()
                kDeckDesign.ApplyDirection(self.Prs, kDeckDesign.LoadDirection(self.Spec.get("direction",
                                                                                               "clean-corporate")))
                self.Themed = True
            if self.Themed:
                self.Prs.slide_width, self.Prs.slide_height = Pt(W), Pt(H)
                self._kx = self._ky = 1.0
            else:  # keep the template's size; scale the 1440 x 810 grid onto it
                self._kx, self._ky = self.Prs.slide_width / Pt(W), self.Prs.slide_height / Pt(H)
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.OpenDeck")
            return

    def Notes(self, S, Sl, Extra=""):
        """Write the slide's speaker notes (the spec's, then what the pattern adds); remember slides left without."""
        if kS.ErrorMode:
            return
        try:
            Text = "\n\n".join(X for X in (kSlideText.NotesText(Sl.get("notes")), Extra, "\n".join(self.Extra),
                                            "\n".join(Sl.get("_moved") or []), self.DeckSources(Sl)) if X)
            self.Extra = []
            if Text.strip():
                S.notes_slide.notes_text_frame.text = Text
            if not kSlideText.NotesText(Sl.get("notes")).strip():
                self.NoNotes.append(f"slide {self.SlideNo} ({self.SlideId})")
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.Notes")
            return

    def LayoutCheck(self, Out):
        """Lint the saved deck for the layout findings the builder answers for - text running out of its card,
        a kicker on the title, tile text below the label floor, a value that wraps, body text under the floor - and
        every lint error (shapes that overlap, content on the footer) - and report each one as unfit text (exit 3),
        so the builder and lint_deck.py never disagree about a box."""
        if kS.ErrorMode:
            return
        try:
            from lint_deck import kLintDeck  # the same checks lint_deck.py runs, in this process
            _, Found = kLintDeck.Lint(Out, 18.0, 12)
            Counted = set()
            for I in (Found.Items if Found else []):
                No = self.Numbers[I["slide"] - 1] if 0 < I["slide"] <= len(self.Numbers) else I["slide"]
                Item = dict(I, slide=No, spec=self.SpecOf.get(No, 0), id="")
                if I["code"] == "word_budget" and Item["spec"]:
                    continue  # a built slide's words are counted from its spec below, as --plan counts them
                self.LintItems.append(Item)
                if I["code"] in LAYOUT_CODES or I["severity"] == "error":  # an overlap, a clash: never ship it
                    Shape = f" [{I['shape']}]" if I.get("shape") else ""
                    self.Problems.append(f"slide {No}: {I['code']}{Shape} {I['message']}")
                    self.Issues.append(dict(Item, severity="error"))
            for No in self.Numbers:
                SpecNo = self.SpecOf.get(No, 0)
                if not SpecNo or SpecNo in Counted or SpecNo > len(self.Spec["slides"]):
                    continue
                Counted.add(SpecNo)
                Sl = self.Spec["slides"][SpecNo - 1]
                Words, Budget = kPlanCheck.Words(Sl), kPlanCheck.Budget(Sl)
                if Words > Budget and Sl.get("pattern") not in NO_CHROME:
                    self.LintItems.append({"slide": No, "spec": SpecNo, "id": "", "severity": "info",
                                           "code": "word_budget", "shape": "", "message":
                                           f"about {Words} visible words (budget {Budget} for a {Sl.get('pattern')} "
                                           "slide; counted from the spec as --plan counts them). Fine for chart, "
                                           "quote and reference slides; otherwise cut or move to the notes."})
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.LayoutCheck")
            return

    def DeckSources(self, Sl):
        """The deck-level 'sources' as a SOURCES: section, for a content slide whose own notes name none."""
        if kS.ErrorMode:
            return ""
        try:
            Src = self.Spec.get("sources")
            if not Src or Sl.get("pattern") in NO_CHROME:
                return ""
            Notes = Sl.get("notes")
            if isinstance(Notes, dict) and Notes.get("sources"):
                return ""
            if isinstance(Notes, str) and re.search(r"source", Notes, re.I):
                return ""
            Src = [Src] if isinstance(Src, str) else Src
            return "SOURCES:\n" + "\n".join(f"- {X}" for X in Src)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kDeckBuilder.DeckSources")
            return ""

    def Build(self, Out, Only=None):
        """Build every slide (or only the spec slides numbered in Only, with the page numbers and section kickers
        they have in the whole deck) and save to Out. Returns the slide count (fit problems are in self.Problems)."""
        if kS.ErrorMode:
            return 0
        try:
            self.OpenDeck()
            self.Layout = kDeckDesign.TitleOnlyLayout(self.Prs)
            Th = kTheme(self.Prs.slide_master)
            self.Major, self.Minor = Th.Font("+mj-lt"), Th.Font("+mn-lt")
            Patterns = kSlidePatterns(self)
            N = 0
            for SpecNo, Sl in enumerate(self.Spec["slides"], 1):
                N += 1
                if Only and SpecNo not in Only:  # skipped, but it still numbers the pages and sets the section
                    if Sl["pattern"] == "section":
                        self.Section = str(Sl.get("kicker") or Sl.get("eyebrow") or "")
                    if Sl["pattern"] == "quiz" and Sl.get("reveal") == "slide":
                        N += 1
                    continue
                self.Numbers.append(N)
                self.SpecNo, self.SpecOf[N] = SpecNo, SpecNo
                self.SlideNo, self.SlideId = N, Sl.get("id", f"s{N:02d}")
                S = self.NewSlide(Sl, N)
                self.TitleInk, self.KickerWidth = None, 0
                self.Kicker = self.KickerFor(Sl)
                self.Reserve = self.HasFooter(Sl)
                getattr(Patterns, PATTERNS[Sl["pattern"]])(S, Sl)
                self.Reserve = False  # the footer and page number themselves sit in the band
                self.Chrome(S, Sl, N)
                self.Notes(S, Sl, kSlidePatterns.ExtraNotes(Sl))
                if Sl["pattern"] == "quiz" and Sl.get("reveal") == "slide":  # the answer on a slide of its own
                    N += 1
                    self.Numbers.append(N)
                    self.SpecOf[N] = SpecNo
                    Answer = dict(Sl, id=f"{self.SlideId}-answer")
                    self.SlideNo, self.SlideId = N, Answer["id"]
                    S = self.NewSlide(Answer, N)
                    self.TitleInk, self.KickerWidth = None, 0
                    self.Kicker = self.KickerFor(Answer)
                    self.Reserve = self.HasFooter(Answer)
                    Patterns.QuizAnswer(S, Answer)
                    self.Reserve = False
                    self.Chrome(S, Answer, N)
                    self.Notes(S, Answer, kSlidePatterns.ExtraNotes(Sl))
            if kS.ErrorMode:
                return 0  # a step was reported and halted: do not write a half-built deck
            self.Prs.save(Out)
            self.LayoutCheck(Out)
            return len(self.Numbers)
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kDeckBuilder.Build(out={Out})")
            return 0


class kSlidePatterns:
    """The slide patterns: one method per pattern, drawing onto a slide through a kDeckBuilder. Each one lays its
    content out over the whole body area (BODY_TOP..BODY_BOTTOM, margin to margin) and grows text up to the
    ceilings in GROW, so a slide with little to say still reads from the back of the room."""

    def __init__(self, Builder):
        try:
            self._b = Builder
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.__init__")

    @staticmethod
    def ReadingOrder(Calls):
        """An email's callouts in the order the eye meets their targets (from, to, subject, attachment, body
        lines top to bottom), so marker 1 is the first red flag read and the explanation list follows suit."""
        if kS.ErrorMode:
            return []
        try:
            return sorted([C for C in Calls if isinstance(C, dict)], key=kSlidePatterns.ReadingRank)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.ReadingOrder")
            return []

    @staticmethod
    def ReadingRank(Call):
        """Where a callout's target sits on the drawn email: (field rank, body line)."""
        if kS.ErrorMode:
            return (0, 0)
        try:
            Rank = {"from": 0, "to": 1, "subject": 2, "attachment": 3, "body": 4}
            Target = Call.get("target", "body")
            Line = Call.get("line", 0) if Target == "body" else 0
            return (Rank.get(Target, 4), Line if isinstance(Line, int) else 0)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.ReadingRank")
            return (0, 0)

    @staticmethod
    def ExtraNotes(Sl):
        """Notes the pattern adds after the spec's own: an email's callouts, a quiz's answer."""
        if kS.ErrorMode:
            return ""
        try:
            Pat = Sl.get("pattern")
            if Pat == "email" and Sl.get("callouts"):
                return "CALLOUTS:\n" + "\n".join(f"{I}. {C.get('note', '')} ({C.get('target', 'body')})"
                                                 for I, C in enumerate(kSlidePatterns.ReadingOrder(Sl["callouts"]), 1))
            if Pat == "quiz":
                Right = [f"{'ABCD'[I]}: {O.get('text', '')}" for I, O in enumerate(Sl["options"]) if O.get("correct")]
                return "ANSWER: " + "; ".join(Right) + (f"\nWHY: {Sl['explain']}" if Sl.get("explain") else "")
            return ""
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.ExtraNotes")
            return ""

    def Badge(self, S, Name, Text, Lx, Ty, D, Fill=ACCENT, Fg=BG, Size=28):
        """A numbered or lettered circle: one shape with its text inside, so nothing can drift apart."""
        if kS.ErrorMode:
            return None
        try:
            B = self._b
            Dot = B.Rect(S, Name, Lx, Ty, D, D, Fill, MSO_SHAPE.OVAL)
            B.ShapeText(Dot, Text, Size, Fg)
            return Dot
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Badge")
            return None

    def Title(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            Fg, Sub = TEXT, MUTED
            if B.Themed:  # the cover is a full panel in the text colour: the deck opens with weight
                B.Backdrop(S, "CoverPanel", TEXT)
                B.Rect(S, "CoverBand", 0, H - 24, W, 24, ACCENT)
                Fg = Sub = BG
            B.Rect(S, "AccentRule", M, 236, 160, 10, ACCENT)
            B.Title(S, Sl.get("title", ""), Size=SIZE["section"], Top=262, Height=230, Colour=Fg)
            if Sl.get("subtitle"):
                B.Text(S, "Subtitle", Sl["subtitle"], M, 516, W - 2 * M - 160, 130, SIZE["h2"], Sub)
            if B.Spec.get("footer"):
                B.Text(S, "CoverFooter", str(B.Spec["footer"]), M, FOOTER_TOP - 40, W - 2 * M, FOOTER_H,
                       FOOTER_SIZE, Sub, Anchor=MSO_ANCHOR.MIDDLE)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Title")
            return

    def Section(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            Fg = TEXT
            if B.Themed:  # dividers echo the cover, so the deck has a rhythm
                B.Backdrop(S, "SectionPanel", TEXT)
                Fg = BG
            if Sl.get("eyebrow"):
                Tb = B.Text(S, "Eyebrow", Sl["eyebrow"].upper(), M, 250, W - 2 * M, 50, SIZE["caption"],
                            BG if B.Themed else ACCENT, Bold=True)
                for R in Tb.text_frame.paragraphs[0].runs:
                    R.font._rPr.set("spc", "300")
            B.Title(S, Sl["title"], Size=SIZE["section"], Top=300, Height=220, Colour=Fg)
            B.Rect(S, "AccentRule", M, 540, 160, 10, ACCENT)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Section")
            return

    def Statement(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            if Sl.get("decision"):
                self.Decision(S, Sl)
                return
            Points = [str(X) for X in Sl.get("points") or []]
            Pw = POINTS_W if Points else 0  # the points' column on the right: numbered cards, the full body high
            Wd = W - 2 * M - (Pw + 2 * GAP if Pw else 80)
            Ts = B.Fit([Sl["title"]], SIZE["statement"] + (0 if Pw else 16), Wd - 15, 300, Heading=True,
                       What="statement")
            Tw = B.BalancedWidth(Sl["title"], Ts, Wd - 15) + 15
            Display = 40 * max(1.0, B.Prs.slide_width / 12700 / 960) / B.TypeScale()  # lint's display size
            Cap = 380 if Pw else 300  # beside points the claim may take a third display line, not shrink to two
            while Ts > 48 and (B.TitleLines(Sl["title"], Ts, Tw + 1, True) * Ts * LOOSE_LINE + 8 > Cap
                               or B.TitleLines(Sl["title"], Ts, Tw) > B.MaxTitleLines(Ts)):  # as Title() checks
                Next = Ts - 4  # the claim must also fit where LibreOffice wraps it wider (a substituted font)
                if Next < Display and B.TitleLines(Sl["title"], Next, B.BalancedWidth(Sl["title"], Next, Wd - 15)
                                                   + 15) > 2:
                    break  # below display size a claim may take two lines only: keep the measured three
                Ts = Next
                Tw = B.BalancedWidth(Sl["title"], Ts, Wd - 15) + 15
            while Ts > 36 and B.TitleLines(Sl["title"], Ts, Tw) > B.MaxTitleLines(Ts):
                Ts -= 2  # Title() would shrink it to this anyway: size the box, the support and the block from it
                Tw = B.BalancedWidth(Sl["title"], Ts, Wd - 15) + 15
            LooseH = B.TitleLines(Sl["title"], Ts, Tw + 1, True) * Ts * LOOSE_LINE + 8  # the loosest renderer's wrap
            Th = min(max(300, LooseH), max(B.Need([Sl["title"]], Ts, Tw - 15, Heading=True) + 10, LooseH))
            Support = str(Sl.get("support") or "")
            # the support line never outranks the claim it supports: at most three quarters of the claim's size
            Ss = B.Fit([Support], max(FLOOR, min(44 if Pw else 48, int(Ts * 0.75))), Wd, 170, Floor=FLOOR) \
                if Support else 0
            Sh = B.Need(Support, Ss, Wd) if Support else 0
            Kh = KICKER_H + 12 if B.Kicker else 0  # the kicker sits between the rule and the claim, as on every slide
            Lead = 46 + Kh
            Block = 10 + 36 + Kh + Th + (34 + Sh if Support else 0)
            Top = max(120 - Kh, (120 + BODY_BOTTOM - Block) / 2)  # the block sits in the optical middle of the slide
            if Pw:
                B.Rect(S, "AccentRule", M, Top, 160, 10, ACCENT)
            else:  # the anchor: an accent bar the height of the claim and its support, left of the text
                B.Rect(S, "AnchorBar", M - 48, Top + Lead, 16, Block - Lead, ACCENT)
            B.Title(S, Sl["title"], Size=Ts, Top=Top + Lead, Height=Th, Width=Tw)
            if B.Kicker:  # Chrome draws it ending here: above the claim's box, whatever the renderer's line height
                B.TitleInk, B.KickerWidth = Top + Lead - 6, Wd
            if Support:
                B.Text(S, "Support", Support, M, Top + Lead + Th + 34, Wd, Sh, Ss, MUTED, Floor=FLOOR)
            if Pw:
                self.PointCards(S, Points, W - M - Pw, 120, Pw, BODY_BOTTOM - 120)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Statement")
            return

    @staticmethod
    def Share(Sl):
        """(part, whole) that a big number states - '41' with unit '%' is (41, 100), '14/20' is (14, 20), '70 %'
        is (70, 100) - or None when it is not a share of a whole of at most 100."""
        if kS.ErrorMode:
            return None
        try:
            Num = str(Sl.get("number", "")).replace(NBSP, " ").replace(",", ".").strip()
            Unit = str(Sl.get("unit") or "").strip()
            Frac = re.fullmatch(r"(\d+)\s*(?:/|of)\s*(\d+)", Num)
            if Frac and 0 < int(Frac.group(2)) <= 100 and int(Frac.group(1)) <= int(Frac.group(2)):
                return int(Frac.group(1)), int(Frac.group(2))
            Pct = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(%?)", Num)
            if Pct and (Pct.group(2) or Unit == "%") and float(Pct.group(1)) <= 100:
                return float(Pct.group(1)), 100
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Share")
            return None

    def ShareVisual(self, S, Part, Whole, Kind, Lx, Ty, Wd, Ht):
        """The share as a picture: a dot grid (one dot per unit, the part filled with the accent) or a filled bar,
        with a 'part of whole' label under it. Returns the height used."""
        if kS.ErrorMode:
            return 0
        try:
            B = self._b
            Label = (f"{Part:g} in every 100" if Whole == 100 else f"{Part:g} of {Whole}")
            Lh = 44
            if Kind == "bar":
                Bh = min(90, Ht - Lh - 16)
                Y = Ty + (Ht - Bh - Lh - 16) / 2
                B.Rect(S, "ShareTrack", Lx, Y, Wd, Bh, FAINT)
                B.Rect(S, "ShareFill", Lx, Y, max(4, Wd * Part / Whole), Bh, ACCENT)
                B.Text(S, "ShareLabel", Label, Lx, Y + Bh + 16, Wd, Lh, 32, MUTED, Bold=True)
                return Ht
            Cols = 10 if Whole > 25 else (5 if Whole > 6 else Whole)  # 20 is 5 x 4, 100 is 10 x 10
            Rows = -(-Whole // Cols)
            Pitch = min(Wd / Cols, (Ht - Lh - 16) / Rows)
            D = Pitch * 0.72
            Gw = Pitch * Cols
            X0 = Lx + (Wd - Gw) / 2
            Y0 = Ty + (Ht - Pitch * Rows - Lh - 16) / 2
            Filled = round(Part)
            for K in range(Whole):
                B.Rect(S, f"Share{K + 1}", X0 + (K % Cols) * Pitch, Y0 + (K // Cols) * Pitch, D, D,
                       ACCENT if K < Filled else FAINT, MSO_SHAPE.OVAL)
            B.Text(S, "ShareLabel", Label, X0, Y0 + Pitch * Rows + 16 - Pitch + D, Gw, Lh, 32, MUTED, Bold=True)
            return Ht
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.ShareVisual")
            return 0

    def PointCards(self, S, Points, Lx, Ty, Wd, Ht):
        """Two or three short points as numbered cards stacked in a column (tint, accent edge, a badge), sharing
        one text size - the evidence beside a statement."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            N = len(Points)
            Ch = min(200, (Ht - GAP * (N - 1)) / N)
            Y0 = Ty + (Ht - Ch * N - GAP * (N - 1)) / 2
            D, Pad = 56, 28
            Tw = Wd - EDGE - Pad * 3 - D
            Size = B.FitAll(Points, 36, Tw, Ch - 2 * 20, Floor=FLOOR)
            for I, Pt_ in enumerate(Points):
                Y = Y0 + I * (Ch + GAP)
                B.Card(S, f"PointCard{I + 1}", Lx, Y, Wd, Ch, False, "left")
                self.Badge(S, f"PointNo{I + 1}", str(I + 1), Lx + EDGE + Pad, Y + (Ch - D) / 2, D, ACCENT, BG, 26)
                B.Text(S, f"StatementPoint{I + 1}", Pt_, Lx + EDGE + Pad * 2 + D, Y + 20, Tw, Ch - 40, Size, TEXT,
                       Anchor=MSO_ANCHOR.MIDDLE, Floor=FLOOR)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.PointCards")
            return

    def ReasonCards(self, S, Points, Lx, Ty, Wd, Ht):
        """The reasons under a decision: two or three numbered cards side by side (tint, accent edge on top, a
        badge beside the text), sharing one text size that never goes below the floor."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            N = len(Points)
            Cw = (Wd - GAP * (N - 1)) / N
            Ch = min(220, Ht)
            if Ch < 100:
                B.Report("DecisionBox", "fit", f"no room for the reasons under the decision "
                                  f"({Ht:.0f} pt left); shorten the decision or support, or drop the figure")
                return
            D, Pad = 48, 24
            Tw = Cw - Pad * 3 - D
            Size = B.FitAll(Points, 34, Tw, Ch - EDGE - 2 * Pad, What="the reasons", Floor=FLOOR)
            for I, Pt_ in enumerate(Points):
                X = Lx + I * (Cw + GAP)
                B.Card(S, f"ReasonCard{I + 1}", X, Ty, Cw, Ch, False, "top")
                self.Badge(S, f"ReasonNo{I + 1}", str(I + 1), X + Pad, Ty + EDGE + (Ch - EDGE - D) / 2, D, ACCENT,
                           BG, 24)
                B.Text(S, f"Reason{I + 1}", Pt_, X + Pad * 2 + D, Ty + EDGE + Pad, Tw, Ch - EDGE - 2 * Pad, Size,
                       TEXT, Anchor=MSO_ANCHOR.MIDDLE, Floor=FLOOR)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.ReasonCards")
            return

    def BigNumber(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Num = str(Sl["number"]).replace(" ", NBSP) + (f"{NBSP}{Sl['unit']}" if Sl.get("unit") else "")
            Share = self.Share(Sl)
            Kind = Sl.get("visual", "auto")
            Kind = ("dots" if Share else "none") if Kind == "auto" else Kind
            if Kind == "dots" and Share and Share[1] > 100:
                Kind = "bar"
            Points = [str(X) for X in Sl.get("points") or []]
            if Kind == "none" or not Share:
                if not Points:  # the classic layout: the number, then its caption
                    Ns = B.FitLine([Num], SIZE["hero"] + 20, W - 2 * M, Smallest=96)
                    B.Text(S, "HeroNumber", Num, M, 220, W - 2 * M, 300, Ns, ACCENT, Bold=True, Heading=True)
                    if Sl.get("caption"):
                        B.Rect(S, "CaptionEdge", M, 550, EDGE, 120, ACCENT)
                        B.Text(S, "Caption", Sl["caption"], M + EDGE + 32, 550, W - 2 * M - 160, 150, SIZE["h2"], TEXT)
                    return
            Vw = 0 if Kind == "none" or not Share else SHARE_W  # the visual's column on the right
            Lw = W - 2 * M - (Vw + 2 * GAP if Vw else 0)
            Ns = B.FitLine([Num], SIZE["hero"], Lw, Smallest=96, What="the number")
            Cap = str(Sl.get("caption") or "")
            Cs = B.Fit([Cap], 40, Lw - EDGE - 32, 130, Floor=FLOOR) if Cap else 0
            Ch = B.Need(Cap, Cs, Lw - EDGE - 32) if Cap else 0
            Bul = ["\u2022 " + X for X in Points]
            Avail = BODY_BOTTOM - BODY_TOP  # the column ends above the footer band: the number gives way first
            Rest = (24 + Ch if Cap else 0) + 36
            Ps = B.FitAll(Points, 34, Lw - 40, 200, Floor=FLOOR) if Points else 0
            Ph = B.Need(Bul, Ps, Lw - 40) if Points else 0
            while Points and Ns * 1.15 + Rest + Ph > Avail and Ns > 160:
                Ns -= 8
            while Points and Ns * 1.15 + Rest + Ph > Avail and Ps > FLOOR:
                Ps = max(FLOOR, Ps - 2)
                Ph = B.Need(Bul, Ps, Lw - 40)
            while Points and Ns * 1.15 + Rest + Ph > Avail and Ns > 96:
                Ns -= 8
            Hh = Ns * 1.15
            Block = Hh + (24 + Ch if Cap else 0) + (36 + Ph if Points else 0)
            Y = BODY_TOP + max(0, (BODY_BOTTOM - BODY_TOP - Block) / 2)  # the column sits in the body's middle
            B.Text(S, "HeroNumber", Num, M, Y, Lw, Hh, Ns, ACCENT, Bold=True, Heading=True)
            Y += Hh + 24
            if Cap:
                B.Rect(S, "CaptionEdge", M, Y, EDGE, Ch, ACCENT)
                B.Text(S, "Caption", Cap, M + EDGE + 32, Y, Lw - EDGE - 32, Ch, Cs, TEXT, Floor=FLOOR)
                Y += Ch + 36
            if Points:
                B.Text(S, "Points", Bul, M, Y, Lw - 40, Ph, Ps, TEXT, Floor=FLOOR)
            if Vw:
                self.ShareVisual(S, Share[0], Share[1], Kind, W - M - Vw, BODY_TOP + 10, Vw,
                                 BODY_BOTTOM - BODY_TOP - 10)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.BigNumber")
            return

    def Tiles(self, S, Ms, Hi, Lx, Ty, Wd, Ht, Across, ValueMax, LabelMax):
        """Metric tiles (value, label, optional note and trend) in a row (Across) or a column, filling Wd x Ht;
        the highlighted one filled with the accent, the others tinted with an accent edge and the value in the
        accent. Values, labels and notes share one size each and line up. Labels, notes and the trend's figures
        never go below the label floor (LABEL_MIN): the trend and the value give way first, and what still does
        not fit is reported."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            N = len(Ms)
            Tw = (Wd - GAP * (N - 1)) / N if Across else Wd
            Th = Ht if Across else (Ht - GAP * (N - 1)) / N
            Pad = 32 if Across else 28
            Iw = Tw - 2 * Pad - (0 if Across else EDGE)
            Ix = Pad + (0 if Across else EDGE)
            Values = [self.Value(X["value"]) for X in Ms]
            Notes = [str(X.get("note", "")) for X in Ms]
            Trends = any(X.get("trend") for X in Ms)
            Periods = any(X.get("trend_labels") for X in Ms if X.get("trend"))
            TrMin = TREND_LABEL * 1.3 + TREND_BARS + (TREND_LABEL * 1.3 if Periods else 0)  # figures, bars, periods
            Tr = max(TrMin, min(TREND_MAX, Th * 0.42)) if Trends else 0
            Vs = B.FitLine(Values, ValueMax, Iw)
            # first guesses (the loop below has the last word, and reports what cannot fit)
            Ls = B.FitAll([str(X["label"]) for X in Ms], LabelMax, Iw, Th * (0.28 if Across else 0.3) + 30)
            Ns = B.FitAll([X for X in Notes if X], GROW["note"] - 2, Iw, Th * 0.22 + 30) if any(Notes) else 0
            Room = Th - Pad * 0.8 - (EDGE if Across else 0)
            VFloor = max(64, ValueMax * 0.6)
            while Trends:  # legible bars or none: the value and the labels give a little first
                Hl = max(B.Need(str(X["label"]), Ls, Iw) for X in Ms)
                Hn = max(B.Need(X, Ns, Iw) for X in Notes if X) if Ns else 0
                if Room - (Vs * 1.28 + 10 + Hl + (14 + Hn if Ns else 0)) >= TrMin + 18:
                    break
                if Vs > VFloor:
                    Vs = max(VFloor, Vs - 4)
                elif Ls > 28 or Ns > 28:
                    Ls, Ns = max(28, Ls - 2), (max(28, Ns - 2) if Ns else 0)
                else:
                    Trends = False
            if any(X.get("trend") for X in Ms) and not Trends:
                # no room for legible bars: the trend moves to the note (when there is none) and the speaker notes
                Vs = B.FitLine(Values, ValueMax, Iw)
                Tr = 0
                for I, Mt in enumerate(Ms):
                    if Mt.get("trend"):
                        Text = self.TrendText(Mt)
                        Notes[I] = Notes[I] or Text
                        B.Extra.append(f"TREND ({Mt['label']}, not drawn - too little room on the slide): {Text}")
                Ns = B.FitAll([X for X in Notes if X], GROW["note"] - 2, Iw, Th * 0.22 + 30)
            while True:  # value, label, note and trend must fit the tile together: the trend gives way first
                Hv = Vs * 1.28  # FitLine keeps every value on one line
                Hl = max(B.Need(str(X["label"]), Ls, Iw) for X in Ms)
                Hn = max(B.Need(X, Ns, Iw) for X in Notes if X) if Ns else 0
                Block = Hv + 10 + Hl + (14 + Hn if Ns else 0) + (18 + Tr if Tr else 0)
                if Block <= Room:
                    break
                if Tr > TrMin:
                    Tr = max(TrMin, Tr - 10)
                elif Ls > 28 or Ns > 28:  # the labels give way before the value: the value is the point
                    Ls, Ns = max(28, Ls - 2), (max(28, Ns - 2) if Ns else 0)
                elif Vs > 56:
                    Vs -= 4
                elif Ls > LABEL_MIN or Ns > LABEL_MIN:
                    Ls, Ns = max(LABEL_MIN, Ls - 2), (max(LABEL_MIN, Ns - 2) if Ns else 0)
                elif Vs > 40:
                    Vs -= 4
                else:
                    B.Report("Note", "fit", f"the metric tiles do not fit at the label "
                                      f"floor ({LABEL_MIN} pt); shorten labels or notes, or drop a metric")
                    break
            if Tr and Block < Room:  # room to spare: the trend takes it, so its bars read from the back
                Grow = min(Room - Block, TREND_MAX - Tr) * 0.8
                Tr, Block = Tr + max(0, Grow), Block + max(0, Grow)
            Figure = ACCENT if B.Themed else TEXT
            for I, Mt in enumerate(Ms):
                X0 = Lx + (I * (Tw + GAP) if Across else 0)
                Y0 = Ty + (0 if Across else I * (Th + GAP))
                On = I == Hi
                B.Card(S, f"Card{I + 1}", X0, Y0, Tw, Th, On, "top" if Across else "left", EdgeOn=False)
                Fg = BG if On else TEXT
                Top = Y0 + max(Pad * 0.6 + (EDGE if Across else 0), (Th - Block) / 2)
                B.Text(S, f"Value{I + 1}", Values[I], X0 + Ix, Top, Iw, Hv, Vs, BG if On else Figure, Bold=True,
                       Heading=True)
                B.Text(S, f"Label{I + 1}", str(Mt["label"]), X0 + Ix, Top + Hv + 10, Iw, Hl, Ls, Fg)
                Y = Top + Hv + 10 + Hl
                if Notes[I]:
                    B.Text(S, f"Note{I + 1}", Notes[I], X0 + Ix, Y + 14, Iw, Hn, Ns, BG if On else MUTED)
                if Tr and Mt.get("trend"):
                    self.Trend(S, I, Mt["trend"], Mt.get("trend_labels"), X0 + Ix, Top + Block - Tr, Iw, Tr, On)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Tiles")
            return

    def Kpi(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Top = BODY_TOP + 24
            Bottom = BODY_BOTTOM - 6
            Hi = Sl.get("highlight", 0)
            Lead = self.ChartMetric(Sl["metrics"], Hi)
            if Lead is not None and not Sl.get("decision") and Sl.get("trend_chart", True) \
                    and len(Sl["metrics"]) <= 4:  # a metric with a series: its chart beside the tiles
                self.TilesAndChart(S, Sl, Hi, Lead)
                return
            if Sl.get("decision"):  # an exec summary: the numbers, then the decision they lead to
                Bh = self.DecisionSize(Sl, W - 2 * M, 36)[3]
                Bottom = BODY_BOTTOM - Bh - GAP
                self.DecisionBox(S, Sl, M, Bottom + GAP, W - 2 * M, 36)
            self.Tiles(S, Sl["metrics"], Hi, M, Top, W - 2 * M, Bottom - Top, True, GROW["value"], 40)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Kpi")
            return

    def Bullets(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            if all(len(str(X)) <= 80 for X in Sl["items"]):
                if len(Sl["items"]) <= 5:
                    self.BulletRows(S, Sl)  # a few short points: one numbered band each, filling the slide
                else:
                    self.BulletGrid(S, Sl)  # six or seven: numbered cards in two columns, never a bare list
                return
            Items = ["\u2022 " + str(X) for X in Sl["items"]]
            Inset = EDGE + 32  # long points keep a plain list, but on an accent rule, not floating on the slide
            Wd, Ht = W - 2 * M - 120 - Inset, BODY_BOTTOM - BODY_TOP - 20
            Size = B.Fit(Items, GROW["body"], Wd, Ht)
            Spare = max(0.0, Ht - B.Need(Items, Size, Wd))
            Spacing = 0.5 + (min(1.1, Spare / (len(Items) - 1) / Size * 0.7) if len(Items) > 1 else 0)
            Used = min(Ht, B.Need(Items, Size, Wd, Spacing=Spacing))
            B.Rect(S, "PointsRule", M, BODY_TOP + 20, EDGE, Used, ACCENT)
            B.Text(S, "Points", Items, M + Inset, BODY_TOP + 20, Wd, Ht, Size, Spacing=Spacing)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Bullets")
            return

    def BulletRows(self, S, Sl):
        """Up to five short points as full-width tinted bands with an accent edge and a number badge, sized to fill
        the body: a list the eye can count, never bare bullets on an empty slide."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            Items = [str(X) for X in Sl["items"]]
            N = len(Items)
            Ht = min(150, (BODY_BOTTOM - BODY_TOP - 10 - GAP * (N - 1)) / N)
            Pad = 32
            D = min(64, Ht - 24)
            Tw = W - 2 * M - EDGE - 3 * Pad - D
            Size = B.FitAll(Items, GROW["body"], Tw, Ht - 20)
            for I, Item in enumerate(Items):
                Ty = BODY_TOP + 10 + I * (Ht + GAP)
                B.Card(S, f"PointBand{I + 1}", M, Ty, W - 2 * M, Ht, False, "left")
                self.Badge(S, f"BulletNo{I + 1}", str(I + 1), M + EDGE + Pad, Ty + (Ht - D) / 2, D, ACCENT, BG,
                           min(32, D * 0.5))
                B.Text(S, f"Point{I + 1}", Item, M + EDGE + 2 * Pad + D, Ty + 10, Tw, Ht - 20, Size, TEXT,
                       Anchor=MSO_ANCHOR.MIDDLE)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.BulletRows")
            return

    def BulletGrid(self, S, Sl):
        """Six or seven short points as numbered cards in two columns (read down the left column, then the right):
        each card a tinted band with an accent edge and a number badge."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            Items = [str(X) for X in Sl["items"]]
            N = len(Items)
            Rows = -(-N // 2)
            Wd = (W - 2 * M - GAP) / 2
            Ht = (BODY_BOTTOM - BODY_TOP - 10 - GAP * (Rows - 1)) / Rows
            Pad, D = 24, min(56, Ht - 24)
            Tw = Wd - EDGE - 3 * Pad - D
            Size = B.FitAll(Items, GROW["point"], Tw, Ht - 16)
            for I, Item in enumerate(Items):
                Lx = M + (I // Rows) * (Wd + GAP)
                Ty = BODY_TOP + 10 + (I % Rows) * (Ht + GAP)
                B.Card(S, f"PointBand{I + 1}", Lx, Ty, Wd, Ht, False, "left")
                self.Badge(S, f"BulletNo{I + 1}", str(I + 1), Lx + EDGE + Pad, Ty + (Ht - D) / 2, D, ACCENT, BG,
                           min(28, D * 0.5))
                B.Text(S, f"Point{I + 1}", Item, Lx + EDGE + 2 * Pad + D, Ty + 8, Tw, Ht - 16, Size, TEXT,
                       Anchor=MSO_ANCHOR.MIDDLE)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.BulletGrid")
            return

    def Compare(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Cols = Sl["columns"]
            Hi = Sl.get("highlight")
            N = len(Cols)
            Gap = GAP * 1.5
            Wd = (W - 2 * M - Gap * (N - 1)) / N
            Pad = 32
            Iw = Wd - 2 * Pad
            BodyH = BODY_BOTTOM - BODY_TOP - 10
            Hs = B.FitAll([str(C["heading"]) for C in Cols], GROW["heading"] - 4, Iw, 120, Heading=True)
            Hh = max(B.Need(str(C["heading"]), Hs, Iw, Heading=True) for C in Cols)
            Room = BodyH - EDGE - 2 * Pad - Hh - 20
            Points = [["• " + str(P) for P in C["points"]] for C in Cols]
            Ps = B.FitAll(Points, GROW["point"], Iw, Room)
            Ph = max(B.Need(P, Ps, Iw) for P in Points)
            Most = max(len(P) for P in Points)
            Spacing = 0.5 + (min(0.9, max(0.0, Room - Ph) / (Most - 1) / Ps * 0.6) if Most > 1 else 0)
            Ph = min(max(B.Need(P, Ps, Iw, Spacing=Spacing) for P in Points), Room)  # overflow stays visible to lint
            CardH = min(BodyH, max(EDGE + 2 * Pad + Hh + 20 + Ph, BodyH * 0.78))
            Ty = BODY_TOP + 10 + (BodyH - CardH) / 2
            for I, C in enumerate(Cols):
                Lx = M + I * (Wd + Gap)
                B.Card(S, f"Card{I + 1}", Lx, Ty, Wd, CardH, False, "top", EdgeOn=Hi is None or I == Hi)
                B.Text(S, f"Heading{I + 1}", str(C["heading"]), Lx + Pad, Ty + EDGE + Pad, Iw, Hh, Hs,
                       ACCENT if I == Hi and B.Themed else TEXT, Bold=True, Heading=True)
                B.Text(S, f"Points{I + 1}", Points[I], Lx + Pad, Ty + EDGE + Pad + Hh + 20, Iw, Ph, Ps, TEXT,
                       Spacing=Spacing)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Compare")
            return

    def Process(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Steps = Sl["steps"]
            if len(Steps) > 5:
                self.ProcessGrid(S, Sl)
                return
            if len(Steps) == 5:
                self.ProcessRows(S, Sl)
                return
            N = len(Steps)
            Hi = Sl.get("highlight")
            Wd = (W - 2 * M - GAP * (N - 1)) / N
            Pad = 26
            Iw = Wd - 2 * Pad
            ChevH, BodyH = 96, BODY_BOTTOM - BODY_TOP - 10
            Ls = B.FitAll([str(X["label"]) for X in Steps], 44, Iw, 150, Bold=True)
            Lh = max(B.Need(str(X["label"]), Ls, Iw, Bold=True) for X in Steps)
            Details = [str(X.get("detail", "")) for X in Steps]
            Room = BodyH - ChevH - 20 - 2 * Pad - EDGE - Lh - 14
            Ds = B.FitAll([X for X in Details if X], min(GROW["detail"] + 4, max(Ls, FLOOR)), Iw, Room) \
                if any(Details) else 0  # never louder than the label above it
            Dh = max(B.Need(X, Ds, Iw) for X in Details if X) if Ds else 0
            Content = 2 * Pad + EDGE + Lh + (14 + Dh if Ds else 0)
            CardH = min(BodyH - ChevH - 20, max(Content, BodyH * 0.6 - ChevH))  # no tall empty cards
            Top = BODY_TOP + 10 + (BodyH - ChevH - 20 - CardH) / 2
            Depth = 0.32
            for I, St in enumerate(Steps):
                Lx = M + I * (Wd + GAP)
                On = I == Hi
                Arrow = B.Rect(S, f"Step{I + 1}", Lx, Top, Wd + (GAP * 0.6 if I < N - 1 else 0), ChevH,
                               ACCENT if On else QUIET, MSO_SHAPE.PENTAGON if I == 0 else MSO_SHAPE.CHEVRON)
                Arrow.adjustments[0] = Depth
                # the number sits in the arrow's body, clear of the notch and the point
                B.ShapeText(Arrow, str(I + 1), SIZE["h2"], BG if On else ACCENT, Inset=ChevH * Depth + 6)
                Card = Top + ChevH + 20
                B.Card(S, f"StepCard{I + 1}", Lx, Card, Wd, CardH, False, "top", EdgeOn=Hi is None or On)
                Block = Lh + (14 + Dh if Ds else 0)
                Y = Card + EDGE + max(Pad, (CardH - EDGE - Block) / 2)  # the text block in the card's middle
                B.Text(S, f"StepLabel{I + 1}", str(St["label"]), Lx + Pad, Y, Iw, Lh, Ls, TEXT, Bold=True)
                if Details[I]:
                    B.Text(S, f"StepDetail{I + 1}", Details[I], Lx + Pad, Y + Lh + 14, Iw, Dh, Ds, TEXT)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Process")
            return

    def ProcessGrid(self, S, Sl):
        """Six to eight steps: numbered cards in two rows, so labels keep a readable size."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            Steps = Sl["steps"]
            Hi = Sl.get("highlight")
            Cols = -(-len(Steps) // 2)
            Wd = (W - 2 * M - GAP * (Cols - 1)) / Cols
            Ht = (BODY_BOTTOM - BODY_TOP - 10 - GAP) / 2
            Pad, D = 24, 56
            Iw = Wd - 2 * Pad
            Ls = B.FitAll([str(X["label"]) for X in Steps], 34, Iw, 90, Bold=True)
            Lh = max(B.Need(str(X["label"]), Ls, Iw, Bold=True) for X in Steps)
            Details = [str(X.get("detail", "")) for X in Steps]
            Room = Ht - 2 * Pad - EDGE - D - 14 - Lh - 10
            Ds = B.FitAll([X for X in Details if X], 30, Iw, Room) if any(Details) else 0
            for I, St in enumerate(Steps):
                Lx = M + (I % Cols) * (Wd + GAP)
                Ty = BODY_TOP + 10 + (I // Cols) * (Ht + GAP)
                On = I == Hi
                B.Card(S, f"Step{I + 1}", Lx, Ty, Wd, Ht, On, "top")
                self.Badge(S, f"StepNo{I + 1}", str(I + 1), Lx + Pad, Ty + EDGE + Pad, D, BG if On else ACCENT,
                           ACCENT if On else BG)
                Fg = BG if On else TEXT
                Y = Ty + EDGE + Pad + D + 14
                B.Text(S, f"StepLabel{I + 1}", str(St["label"]), Lx + Pad, Y, Iw, Lh, Ls, Fg, Bold=True)
                if Details[I]:
                    B.Text(S, f"StepDetail{I + 1}", Details[I], Lx + Pad, Y + Lh + 10, Iw, Room, Ds, Fg)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.ProcessGrid")
            return

    def Timeline(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Ev = Sl["events"]
            N = len(Ev)
            if not any(E.get("date") for E in Ev):  # undated sequential stages: numbered badges on the rail
                self.Stages(S, Sl)
                return
            if self.TimelineInRow(S, Sl):  # dates above the rail, every label below it, all on one line
                return
            Ry = (BODY_TOP + BODY_BOTTOM) / 2 + 10
            Step = (W - 2 * M) / N
            Lw = min(2 * Step - 32, 440) if N > 2 else Step - 32  # neighbours sit on the other side of the rail
            Hi = Sl.get("highlight")
            Centres = [M + Step * I + Step / 2 for I in range(N)]
            Widths = [min(Lw, 2 * (Cx - M), 2 * (W - M - Cx)) for Cx in Centres]  # centred on the dot, on the slide
            Ds = min(B.Fit([str(E["date"])], 54, Widths[I], 70, Heading=True, Bold=True) for I, E in enumerate(Ev))
            Dh = max(B.Need(str(E["date"]), Ds, Widths[I], Heading=True) for I, E in enumerate(Ev))
            Room = min(Ry - 30 - Dh - 8 - BODY_TOP, BODY_BOTTOM - (Ry + 30 + Dh + 8))  # above AND below the rail
            Ls = min(B.Fit([str(E["label"])], 40, Widths[I], Room) for I, E in enumerate(Ev))
            Lh = min(Room, max(B.Need(str(E["label"]), Ls, Widths[I]) for I, E in enumerate(Ev)))
            B.Rect(S, "Rail", M, Ry - 4, W - 2 * M, 8, QUIET)
            for I, E in enumerate(Ev):
                Cx, Wd = Centres[I], Widths[I]
                On = I == Hi or Hi is None
                D = 48 if I == Hi else 32
                B.Rect(S, f"Dot{I + 1}", Cx - D / 2, Ry - D / 2, D, D, ACCENT if On else MUTED, MSO_SHAPE.OVAL)
                Lx = Cx - Wd / 2
                if I % 2 == 0:
                    B.Text(S, f"Date{I + 1}", str(E["date"]), Lx, Ry - 30 - Dh, Wd, Dh, Ds, ACCENT, Bold=True,
                           Align=PP_ALIGN.CENTER, Anchor=MSO_ANCHOR.BOTTOM, Heading=True)
                    B.Text(S, f"Event{I + 1}", str(E["label"]), Lx, Ry - 30 - Dh - 8 - Lh, Wd, Lh, Ls, TEXT,
                           Align=PP_ALIGN.CENTER, Anchor=MSO_ANCHOR.BOTTOM)
                else:
                    B.Text(S, f"Date{I + 1}", str(E["date"]), Lx, Ry + 30, Wd, Dh, Ds, ACCENT, Bold=True,
                           Align=PP_ALIGN.CENTER, Heading=True)
                    B.Text(S, f"Event{I + 1}", str(E["label"]), Lx, Ry + 30 + Dh + 8, Wd, Lh, Ls, TEXT,
                           Align=PP_ALIGN.CENTER)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Timeline")
            return

    def Stages(self, S, Sl):
        """A timeline without dates: undated stages in order - a numbered badge on the rail for each, its label
        above or below (alternating, so neighbours get twice the width), the highlighted stage in the accent."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            if self.StagesInRow(S, Sl):  # every label below its badge, aligned, when they fit one column each
                return
            Ev = Sl["events"]
            N = len(Ev)
            Ry = (BODY_TOP + BODY_BOTTOM) / 2 + 10
            Step = (W - 2 * M) / N
            Lw = min(2 * Step - 32, 440) if N > 2 else Step - 32
            Hi = Sl.get("highlight")
            Centres = [M + Step * I + Step / 2 for I in range(N)]
            Widths = [min(Lw, 2 * (Cx - M), 2 * (W - M - Cx)) for Cx in Centres]
            D = 88
            Room = min(Ry - D / 2 - 24 - BODY_TOP, BODY_BOTTOM - (Ry + D / 2 + 24))  # above AND below the rail
            Ls = min(B.Fit([str(E["label"])], 48, Widths[I], Room) for I, E in enumerate(Ev))
            Lh = min(Room, max(B.Need(str(E["label"]), Ls, Widths[I]) for I, E in enumerate(Ev)))
            B.Rect(S, "Rail", M, Ry - 4, W - 2 * M, 8, QUIET)
            for I, E in enumerate(Ev):
                Cx, Wd = Centres[I], Widths[I]
                On = I == Hi or Hi is None
                self.Badge(S, f"Stage{I + 1}", str(I + 1), Cx - D / 2, Ry - D / 2, D, ACCENT if On else MUTED, BG, 40)
                Lx = Cx - Wd / 2
                if I % 2 == 0:
                    B.Text(S, f"Event{I + 1}", str(E["label"]), Lx, Ry - D / 2 - 24 - Lh, Wd, Lh, Ls, TEXT,
                           Bold=I == Hi, Align=PP_ALIGN.CENTER, Anchor=MSO_ANCHOR.BOTTOM)
                else:
                    B.Text(S, f"Event{I + 1}", str(E["label"]), Lx, Ry + D / 2 + 24, Wd, Lh, Ls, TEXT, Bold=I == Hi,
                           Align=PP_ALIGN.CENTER)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Stages")
            return

    def RowLabels(self, Labels, Wd, Room, Size=40, Hi=None):
        """The size every label fits in one column of width Wd and height Room (label Hi measured bold), and the
        tallest label's height at it; (0, 0) when a label would need more than three lines or drop below 32 pt -
        then the row alternates."""
        if kS.ErrorMode:
            return 0, 0
        try:
            B = self._b
            Three = (0, 0)
            if Wd < 160:  # a column this narrow breaks words: alternate above and below instead
                return Three
            # measured with Need, not Fit: a failed try must not report unfit text - the alternating row follows
            Words = [(Word, I == Hi) for I, X in enumerate(Labels) for Word in X.split()]
            for Try in range(int(Size), 31, -2):  # two even lines at a smaller size read better than three
                if max(B.LineWidth(Word, Try, False, Bold) for Word, Bold in Words) > (Wd - 4) * B._kx:
                    continue  # a word wider than the column would break mid-word
                Th = max(B.Need(X, Try, Wd, Bold=I == Hi) for I, X in enumerate(Labels))
                if Th <= 2.5 * Try * LINE_HEIGHT:
                    return Try, Th
                if not Three[0] and Th <= min(Room, 3.5 * Try * LINE_HEIGHT):  # 3 lines + box insets
                    Three = (Try, Th)
            return Three
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.RowLabels")
            return 0, 0

    def TimelineInRow(self, S, Sl):
        """A dated timeline with every date above the rail and every label below it, one column per event, the
        block centred in the body. False (nothing drawn) when a label needs more than three lines at 32 pt."""
        if kS.ErrorMode:
            return False
        try:
            B = self._b
            Ev = Sl["events"]
            N, Hi = len(Ev), Sl.get("highlight")
            Step = (W - 2 * M) / N
            Wd, BodyH = Step - 24, BODY_BOTTOM - BODY_TOP - 10
            if Wd - 2 * ROW_PAD < 160:  # six or seven events: too narrow for a column each - alternate instead
                return False
            Ds = B.FitAll([str(E["date"]) for E in Ev], 54, Wd, 70, Heading=True, Bold=True)
            Dh = max(B.Need(str(E["date"]), Ds, Wd, Heading=True) for E in Ev)
            Ls, Lh = self.RowLabels([str(E["label"]) for E in Ev], Wd - 2 * ROW_PAD, BodyH - Dh - 60 - 20 - 2 *
                                    ROW_PAD - EDGE, 40, Hi)
            if not Ls:
                return False
            Ch = max(Lh + 2 * ROW_PAD + EDGE, ROW_CARD)
            Ry = BODY_TOP + 10 + (BodyH - (Dh + 60 + Ch)) / 2 + Dh + 30
            B.Rect(S, "Rail", M, Ry - 4, W - 2 * M, 8, QUIET)
            for I, E in enumerate(Ev):
                Cx = M + Step * I + Step / 2
                On = I == Hi or Hi is None
                D = 48 if I == Hi else 32
                B.Rect(S, f"Dot{I + 1}", Cx - D / 2, Ry - D / 2, D, D, ACCENT if On else MUTED, MSO_SHAPE.OVAL)
                B.Text(S, f"Date{I + 1}", str(E["date"]), Cx - Wd / 2, Ry - 30 - Dh, Wd, Dh, Ds, ACCENT, Bold=True,
                       Align=PP_ALIGN.CENTER, Anchor=MSO_ANCHOR.BOTTOM, Heading=True)
            self.LabelCards(S, Ev, Hi, Step, Ry + 30, Ch, Ls)
            return True
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.TimelineInRow")
            return False

    def StagesInRow(self, S, Sl):
        """Undated stages with every label below its numbered badge, one column each, top-aligned so the row
        reads level; False (nothing drawn) when a label needs more than three lines at 32 pt."""
        if kS.ErrorMode:
            return False
        try:
            B = self._b
            Ev = Sl["events"]
            N, Hi = len(Ev), Sl.get("highlight")
            Step = (W - 2 * M) / N
            Wd, BodyH, D = Step - 24, BODY_BOTTOM - BODY_TOP - 10, 88
            Ls, Lh = self.RowLabels([str(E["label"]) for E in Ev], Wd - 2 * ROW_PAD, BodyH - D - 24 - 20 - 2 *
                                    ROW_PAD - EDGE, 44, Hi)
            if not Ls:
                return False
            Ch = max(Lh + 2 * ROW_PAD + EDGE, ROW_CARD)
            Ry = BODY_TOP + 10 + (BodyH - (D + 24 + Ch)) / 2 + D / 2
            B.Rect(S, "Rail", M, Ry - 4, W - 2 * M, 8, QUIET)
            for I, E in enumerate(Ev):
                Cx = M + Step * I + Step / 2
                On = I == Hi or Hi is None
                self.Badge(S, f"Stage{I + 1}", str(I + 1), Cx - D / 2, Ry - D / 2, D, ACCENT if On else MUTED, BG, 40)
            self.LabelCards(S, Ev, Hi, Step, Ry + D / 2 + 24, Ch, Ls)
            return True
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.StagesInRow")
            return False

    def LabelCards(self, S, Ev, Hi, Step, Top, Ch, Ls):
        """One tinted card per event under the rail, all the same height and top so the row reads level; the
        highlighted event's card keeps the accent edge (the others a soft one) and its label goes bold."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            Wd = Step - 24
            for I, E in enumerate(Ev):
                Lx = M + Step * I + 12
                B.Card(S, f"EventCard{I + 1}", Lx, Top, Wd, Ch, False, "top", EdgeOn=Hi is None or I == Hi)
                B.Text(S, f"Event{I + 1}", str(E["label"]), Lx + ROW_PAD, Top + EDGE + ROW_PAD, Wd - 2 * ROW_PAD,
                       Ch - EDGE - 2 * ROW_PAD, Ls, TEXT, Bold=I == Hi, Align=PP_ALIGN.CENTER,
                       Anchor=MSO_ANCHOR.MIDDLE)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.LabelCards")
            return

    def Quote(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            # the title stays a real placeholder (outline, screen readers) but reads as a small label
            T = B.Title(S, Sl.get("title") or "In their words", Size=SIZE["caption"], Top=56, Height=50, Colour=MUTED)
            T.text_frame.vertical_anchor = MSO_ANCHOR.TOP
            B.Text(S, "QuoteMark", "“", M - 10, 120, 120, 170, 180, ACCENT, Bold=True, Heading=True)
            B.Text(S, "Quote", Sl["quote"], M + 120, 220, W - 2 * M - 240, 360, SIZE["statement"] - 12, TEXT,
                   Heading=True)
            Who = Sl.get("attribution", "") + (f", {Sl['role']}" if Sl.get("role") else "")
            if Who:
                B.Text(S, "Attribution", "— " + Who, M + 120, 620, W - 2 * M - 240, 60, SIZE["small"], MUTED)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Quote")
            return

    def DrawChart(self, S, Sl, Lx, Ty, Wd, Ht, Rest=QUIET):
        """A native chart of Sl's categories and series in the box: one highlighted finding, the rest quiet (Rest:
        the colour of the other bars - SOFT on a tinted card, where QUIET would vanish)."""
        if kS.ErrorMode:
            return None
        try:
            B = self._b
            Kind = Sl.get("type", "column")
            Data = CategoryChartData()
            Data.categories = Sl["categories"]
            for Ser in Sl["series"]:
                Data.add_series(Ser["name"], Ser["values"])
            Gf = S.shapes.add_chart(CHART_TYPES[Kind], B.X(Lx), B.Y(Ty), B.X(Wd), B.Y(Ht), Data)
            Gf.name = "Chart"
            Ch = Gf.chart
            Single = len(Sl["series"]) == 1
            Hi = Sl.get("highlight")
            if isinstance(Hi, str):
                Hi = Sl["categories"].index(Hi)  # CheckSpec guarantees it exists
            Ch.has_title = False
            Ch.font.size = B.F(26)
            Ch.font.color.theme_color = MUTED
            # legend: none for a single series; on the right otherwise (never top/bottom - it steals plot height)
            Ch.has_legend = not Single and Kind != "pie"
            if Ch.has_legend:
                Ch.legend.position = XL_LEGEND_POSITION.RIGHT
                Ch.legend.include_in_layout = False
            Plot = Ch.plots[0]
            # data labels: pie, or a single series of bars/columns; never on line charts
            Labels = Kind == "pie" or (Single and Kind in ("column", "bar"))
            Plot.has_data_labels = Labels
            if Labels:
                Dl = Plot.data_labels
                Dl.font.size, Dl.font.bold = B.F(30), True
                Dl.font.color.theme_color = TEXT
                if Sl.get("number_format"):
                    Dl.number_format, Dl.number_format_is_linked = Sl["number_format"], False
                if Kind != "pie":
                    Dl.position = XL_LABEL_POSITION.OUTSIDE_END
                else:  # a pie has no axis or legend: name every slice on the slice
                    Dl.show_category_name, Dl.show_value = True, True
                    Dl.position = XL_LABEL_POSITION.OUTSIDE_END
            if Kind in ("column", "bar"):
                Plot.gap_width = 60
            if Kind != "pie":
                Va = Ch.value_axis
                Va.has_major_gridlines = not Labels
                if Va.has_major_gridlines:
                    Va.major_gridlines.format.line.width = Pt(0.5)
                    Va.major_gridlines.format.line.color.theme_color = QUIET
                Va.visible = not Labels
                Va.has_minor_gridlines = False
                if Sl.get("number_format"):
                    Va.tick_labels.number_format, Va.tick_labels.number_format_is_linked = Sl["number_format"], False
                Ca = Ch.category_axis
                Ca.tick_labels.font.size = B.F(28)
                if Kind == "bar":  # bars list top-down in the order given, like a table
                    Ca.reverse_order = True
                Ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
                Ca.has_major_gridlines = False
            # colour: one series -> highlight one point in the accent, the rest quiet; several -> accent ramp
            for Si, Ser in enumerate(Plot.series):
                if Kind == "line":
                    Ser.format.line.width = Pt(3)
                    Ser.format.line.color.theme_color = ACCENT
                    Ser.format.line.color.brightness = RAMP[Si % len(RAMP)]
                    Ser.smooth = False
                    continue
                if Single:
                    NPts = len(Sl["categories"])
                    for Pi, Point in enumerate(Ser.points):
                        Point.format.fill.solid()
                        Point.format.fill.fore_color.theme_color = ACCENT if (Hi is None or Pi == Hi) else Rest
                        if Kind == "pie" and Hi is None:  # slices must differ: one hue, light to dark
                            Point.format.fill.fore_color.brightness = round(-0.4 + 0.8 * Pi / max(1, NPts - 1), 2)
                else:
                    Ser.format.fill.solid()
                    Ser.format.fill.fore_color.theme_color = ACCENT
                    Ser.format.fill.fore_color.brightness = RAMP[Si % len(RAMP)]
            kSlideText.Alt(Gf, Sl.get("alt") or kSlideText.ChartAlt(Sl))
            return Gf
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.DrawChart")
            return None

    def Chart(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            if self.ChartWithLead(S, Sl):  # a short caption leads above a full-width chart
                return
            HasSide = bool(Sl.get("caption"))
            Cw = W - 2 * M - (420 if HasSide else 0)
            self.DrawChart(S, Sl, M, BODY_TOP, Cw, BODY_BOTTOM - BODY_TOP)
            if HasSide:
                Px, Pw = W - M - 380, 380
                Size = B.Fit([Sl["caption"]], 38, Pw, 380)
                Ht = B.Need(Sl["caption"], Size, Pw)
                Top = BODY_TOP + max(30, (BODY_BOTTOM - BODY_TOP - Ht - 30) / 2)
                B.Rect(S, "NoteRule", Px, Top, 120, 8, ACCENT)
                B.Text(S, "ChartNote", Sl["caption"], Px, Top + 30, Pw, Ht, Size, TEXT)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Chart")
            return

    def ChartWithLead(self, S, Sl):
        """A caption that fits two lines across the slide sits above the chart on an accent rule, and the chart
        takes the full width below it - no side column leaving the chart small in white space. False (nothing
        drawn) for a longer caption, which keeps the side column."""
        if kS.ErrorMode:
            return False
        try:
            B = self._b
            Cap = str(Sl.get("caption") or "")
            if not Cap:
                return False
            Inset = EDGE + 24
            Tw = W - 2 * M - Inset
            Size = B.Fit([Cap], 34, Tw, 200)
            Ht = B.Need(Cap, Size, Tw)
            if Size < 30 or Ht > 2.2 * Size * LINE_HEIGHT:
                return False
            B.Rect(S, "NoteRule", M, BODY_TOP, EDGE, Ht, ACCENT)
            B.Text(S, "ChartNote", Cap, M + Inset, BODY_TOP, Tw, Ht, Size, TEXT, Anchor=MSO_ANCHOR.MIDDLE)
            Top = BODY_TOP + Ht + GAP
            self.DrawChart(S, Sl, M, Top, W - 2 * M, BODY_BOTTOM - Top)
            return True
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.ChartWithLead")
            return False

    def KpiChart(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Hi = Sl.get("highlight_metric")
            Lead = None if "series" in Sl else self.ChartMetric(Sl["metrics"], Hi)  # no series: chart a trend
            self.TilesAndChart(S, Sl, Hi, Lead)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.KpiChart")
            return

    @staticmethod
    def ChartMetric(Ms, Hi):
        """The index of the metric whose trend becomes the chart: the highlighted one if it has a trend, else the
        first with one; None when no metric has a trend."""
        if kS.ErrorMode:
            return None
        try:
            With = [I for I, X in enumerate(Ms) if isinstance(X, dict) and len(X.get("trend") or []) >= 2]
            if not With:
                return None
            return Hi if Hi in With else With[0]
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.ChartMetric")
            return None

    @staticmethod
    def TrendChart(Mt):
        """Chart fields (categories, one series, the last point highlighted, a number format) from a metric's
        trend: its trend_labels as categories, else 'Start' / 'Now' for two values and 1..n for more."""
        if kS.ErrorMode:
            return {}
        try:
            Vals = list(Mt["trend"])
            Names = [str(X) for X in (Mt.get("trend_labels") or [])]
            if len(Names) == len(Vals):
                Cats = Names
            elif len(Names) == 2:
                Cats = [Names[0]] + [NBSP * (J + 1) for J in range(len(Vals) - 2)] + [Names[1]]
            else:
                Cats = ["Start", "Now"] if len(Vals) == 2 else [str(J + 1) for J in range(len(Vals))]
            Places = max((len(f"{X:g}".split(".")[1]) if "." in f"{X:g}" else 0) for X in Vals
                         if isinstance(X, (int, float)))
            return {"type": "column", "categories": Cats, "series": [{"name": str(Mt["label"]), "values": Vals}],
                    "highlight": len(Vals) - 1, "number_format": "0" + ("." + "0" * min(Places, 2) if Places else ""),
                    "alt": f"Column chart of {Mt['label']}: " + ", ".join(
                        f"{C.strip() or J + 1} {V:g}" for J, (C, V) in enumerate(zip(Cats, Vals))
                        if isinstance(V, (int, float))) + "."}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.TrendChart")
            return {}

    def TilesAndChart(self, S, Sl, Hi, Lead):
        """Tiles in a column on the left, a native chart on the right. With Lead (a metric whose trend becomes the
        chart) the chart sits in that metric's own card - its value, label and note as the card's header - and
        the other metrics are the tiles."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            Ms = Sl["metrics"]
            Tw = 420
            Cx = M + Tw + 48
            Ty, Th = BODY_TOP + 10, BODY_BOTTOM - BODY_TOP - 10
            if Lead is None:
                Bottom = BODY_BOTTOM - (70 if Sl.get("caption") else 0)
                self.Tiles(S, Ms, Hi, M, Ty, Tw, Th, False, 80, 32)
                self.DrawChart(S, Sl, Cx, BODY_TOP, W - M - Cx, Bottom - BODY_TOP)
                if Sl.get("caption"):
                    B.Text(S, "ChartNote", Sl["caption"], Cx, Bottom + 10, W - M - Cx, 60, 30, MUTED)
                return
            Mt = Ms[Lead]
            Others = [X for J, X in enumerate(Ms) if J != Lead]
            HiO = None if Hi is None or Hi == Lead else (Hi if Hi < Lead else Hi - 1)
            if Others:
                self.Tiles(S, Others, HiO, M, Ty, Tw, Th, False, 80, 32)
            else:
                Cx = M
            Cw = W - M - Cx
            Pad = 32
            B.Card(S, "ChartCard", Cx, Ty, Cw, Th, False, "top", EdgeOn=Hi in (None, Lead))
            Value = self.Value(Mt["value"])
            Vs = B.FitLine([Value], 88, Cw * 0.42)
            Vw = min(Cw * 0.45, kMeasure.SafeWidth(Value, B.Major, Vs * B.TypeScale(), True) * 1.12 / B._kx + 24)
            Hh = Vs * 1.28
            Y = Ty + EDGE + Pad - 8
            B.Text(S, "LeadValue", Value, Cx + Pad, Y, Vw, Hh, Vs, ACCENT if B.Themed else TEXT, Bold=True,
                   Heading=True)
            Lx = Cx + Pad + Vw + 32
            Lw = Cx + Cw - Pad - Lx
            Note = str(Mt.get("note") or "")
            Ls = B.Fit([str(Mt["label"])], 34, Lw, Hh * (0.55 if Note else 1))
            Lh = B.Need(str(Mt["label"]), Ls, Lw)
            Ns = B.Fit([Note], 28, Lw, max(Hh - Lh, 80)) if Note else 0  # two lines of note may push the chart down
            Nh = B.Need(Note, Ns, Lw) if Note else 0
            Ly = Y + max(0, (Hh - Lh - Nh) / 2)
            B.Text(S, "LeadLabel", str(Mt["label"]), Lx, Ly, Lw, Lh, Ls, TEXT, Bold=True)
            if Note:
                B.Text(S, "LeadNote", Note, Lx, Ly + Lh, Lw, Nh, Ns, MUTED)
            Top = Y + max(Hh, Lh + Nh) + 16
            self.DrawChart(S, dict(Sl, **self.TrendChart(Mt)), Cx + Pad / 2, Top, Cw - Pad, Ty + Th - Pad / 2 - Top,
                           SOFT)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.TilesAndChart")
            return

    @staticmethod
    def IsNumber(Val):
        """True for a table cell that reads as a number (right-aligned)."""
        if kS.ErrorMode:
            return False
        try:
            return str(Val).replace(",", "").replace(".", "").replace("%", "").replace("-", "").replace(
                " ", "").replace(" ", "").strip().isdigit()
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.IsNumber")
            return False

    def Cell(self, Tbl, R, C, Val, Size, Fill, Fg, Bold=False, Right=False):
        """Fill one table cell: text, size, colours, generous margins, centred vertically."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            Cell = Tbl.cell(R, C)
            Cell.text = str(Val)
            Cell.fill.solid()
            Cell.fill.fore_color.theme_color = Fill
            Cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            Cell.margin_left = Cell.margin_right = B.X(16)
            Cell.margin_top = Cell.margin_bottom = B.Y(4)
            Para = Cell.text_frame.paragraphs[0]
            if Right:
                Para.alignment = PP_ALIGN.RIGHT
            for Run in Para.runs:
                Run.font.size, Run.font.bold = B.F(Size), Bold
                Run.font.color.theme_color = Fg
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Cell")
            return

    def Table(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Rows, Cols = len(Sl["rows"]) + 1, len(Sl["header"])
            RowH = min(84, (BODY_BOTTOM - BODY_TOP - 20) / Rows)
            Size = 32 if RowH >= 72 else (30 if RowH >= 62 else 27)
            Gf = S.shapes.add_table(Rows, Cols, B.X(M), B.Y(BODY_TOP + 10), B.X(W - 2 * M), B.Y(RowH * Rows))
            Gf.name = "Table"
            Tbl = Gf.table
            Longest = [max([len(str(Sl["header"][C]))] + [len(str(R[C])) for R in Sl["rows"]]) for C in range(Cols)]
            Share = [max(L, 6) for L in Longest]
            for C in range(Cols):
                Tbl.columns[C].width = B.X((W - 2 * M) * Share[C] / sum(Share))
            for R in range(Rows):
                Tbl.rows[R].height = B.Y(RowH)
            Hi = Sl.get("highlight_row")
            for C, Head in enumerate(Sl["header"]):
                self.Cell(Tbl, 0, C, Head, Size, TEXT, BG, Bold=True, Right=C > 0 and self.IsNumber(Sl["rows"][0][C]))
            for Ri, Row in enumerate(Sl["rows"], 1):
                for C, Val in enumerate(Row):
                    On = Hi == Ri - 1
                    self.Cell(Tbl, Ri, C, Val, Size, ACCENT if On else (QUIET if Ri % 2 == 0 else BG),
                              BG if On else TEXT, Bold=On, Right=C > 0 and self.IsNumber(Val))
            kSlideText.Alt(Gf, Sl.get("alt") or f"Table: {', '.join(str(X) for X in Sl['header'])}; "
                                                 f"{len(Sl['rows'])} rows.")
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Table")
            return

    @staticmethod
    def Amount(Val, Decimals):
        """A cost as text: thousands grouped with a no-break space, a fixed number of decimals."""
        if kS.ErrorMode:
            return ""
        try:
            return f"{Val:,.{Decimals}f}".replace(",", " ")
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Amount")
            return ""

    def CostTable(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Items = Sl["rows"]
            Unit = Sl.get("unit", "")
            Detail = any(R.get("detail") for R in Items)
            Decimals = max(len(f"{R['amount']:g}".split(".")[1]) if "." in f"{R['amount']:g}" else 0 for R in Items)
            Total = sum(R["amount"] for R in Items)
            Tw = 960
            Rows = len(Items) + 2
            RowH = min(104, (BODY_BOTTOM - BODY_TOP - 10) / Rows)
            Size = 32 if RowH >= 72 else (30 if RowH >= 62 else 27)
            if Detail:  # a detail that wraps needs two lines in its row
                Size = min(Size, B.FitAll([str(R.get("detail", "")) for R in Items], Size, Tw - 270 - 190 - 32,
                                          RowH - 8))
            Cols = 3 if Detail else 2
            Ty = BODY_TOP + 10
            Gf = S.shapes.add_table(Rows, Cols, B.X(M), B.Y(Ty), B.X(Tw), B.Y(RowH * Rows))
            Gf.name = "CostTable"
            Tbl = Gf.table
            Widths = [270, Tw - 270 - 190, 190] if Detail else [Tw - 240, 240]
            for C, Wd in enumerate(Widths):
                Tbl.columns[C].width = B.X(Wd)
            for R in range(Rows):
                Tbl.rows[R].height = B.Y(RowH)
            Head = ["Item"] + (["What it pays for"] if Detail else []) + [Unit or "Cost"]
            for C, Text in enumerate(Head):
                self.Cell(Tbl, 0, C, Text, Size, TEXT, BG, Bold=True, Right=C == Cols - 1)
            Hi = Sl.get("highlight")
            for Ri, Row in enumerate(Items, 1):
                On = Hi == Ri - 1
                Fill, Fg = (ACCENT, BG) if On else ((QUIET if Ri % 2 == 0 else BG), TEXT)
                Vals = [Row["item"]] + ([Row.get("detail", "")] if Detail else []) + [self.Amount(Row["amount"],
                                                                                                    Decimals)]
                for C, Val in enumerate(Vals):
                    self.Cell(Tbl, Ri, C, Val, Size, Fill, Fg, Bold=On, Right=C == Cols - 1)
            Last = Rows - 1
            Vals = [Sl.get("total_label", "Total")] + ([""] if Detail else []) + [self.Amount(Total, Decimals)]
            for C, Val in enumerate(Vals):
                self.Cell(Tbl, Last, C, Val, Size, TEXT, BG, Bold=True, Right=C == Cols - 1)
            kSlideText.Alt(Gf, Sl.get("alt") or "Cost table: " + "; ".join(
                f"{R['item']} {self.Amount(R['amount'], Decimals)}" for R in Items)
                + f"; {Sl.get('total_label', 'Total')} {self.Amount(Total, Decimals)} {Unit}.")
            # the total, large, beside the table - the number the decision is about
            Px = M + Tw + 56
            Pw = W - M - Px
            Big = self.Amount(Total, Decimals)
            Inline = bool(Unit) and len(str(Unit)) <= 4  # a short unit ('k', 'MSEK') reads with the number
            if Inline:
                Big = f"{Big}{NBSP}{Unit}"  # one line: a lone 'k' under '410' reads as a stray letter
            Bs = B.FitLine([Big], 110, Pw)
            Bh = B.Need(Big, Bs, Pw, Heading=True)
            Note = Sl.get("note", "")
            Ns = B.Fit([Note], 32, Pw, 260) if Note else 0
            Nh = B.Need(Note, Ns, Pw) if Note else 0
            Uh = 52 if Unit and not Inline else 0
            Block = 40 + 8 + Bh + Uh + (24 + Nh if Note else 0)
            Top = Ty + max(0, (RowH * Rows - Block) / 2)
            B.Text(S, "TotalLabel", Sl.get("total_label", "Total"), Px, Top, Pw, 40, 28, MUTED, Bold=True)
            B.Text(S, "TotalValue", Big, Px, Top + 48, Pw, Bh, Bs, ACCENT, Bold=True, Heading=True)
            if Unit and not Inline:
                B.Text(S, "TotalUnit", Unit, Px, Top + 48 + Bh, Pw, Uh, 36, ACCENT, Bold=True, Heading=True)
            if Note:
                B.Text(S, "CostNote", Note, Px, Top + 48 + Bh + Uh + 24, Pw, Nh, Ns, TEXT)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.CostTable")
            return

    def Image(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            from PIL import Image, ImageOps
            from cover_crop import kCoverCrop
            B = self._b
            B.Title(S, Sl["title"])
            BoxW, BoxH = W - 2 * M, BODY_BOTTOM - BODY_TOP - (90 if Sl.get("caption") else 0)
            Img = ImageOps.exif_transpose(Image.open(Sl["image"]))  # phone photos: honour the rotation tag
            Img = Img.convert("RGBA" if "A" in Img.getbands() else "RGB")  # CMYK/P/L -> something PNG can hold
            Buf = io.BytesIO()
            Crop = kCoverCrop.Crop(Img, BoxW * B._kx, BoxH * B._ky, Sl.get("focus_x", 0.5), Sl.get("focus_y", 0.5))
            Tw = max(Crop.width, 960)  # resample to the exact box ratio: rounding on small images reads as stretch
            Crop.resize((Tw, round(Tw * BoxH * B._ky / (BoxW * B._kx))), Image.LANCZOS).save(Buf, "PNG")
            Buf.seek(0)
            Pic = S.shapes.add_picture(Buf, B.X(M), B.Y(BODY_TOP), B.X(BoxW), B.Y(BoxH))
            Pic.name = "Photo"
            kSlideText.Alt(Pic, Sl.get("alt") or Sl["title"])
            if Sl.get("caption"):
                B.Text(S, "Caption", Sl["caption"], M, BODY_BOTTOM - 74, BoxW, 74, SIZE["caption"], MUTED)
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kSlidePatterns.Image(image={Sl.get('image')})")
            return

    def Matrix(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Q = Sl["quadrants"]
            Hi = Sl.get("highlight")
            Left, Top = M + 70, BODY_TOP + 10
            Wd, Ht = (W - Left - M - GAP) / 2, (BODY_BOTTOM - Top - 50 - GAP) / 2
            Hs = B.FitAll([str(X["heading"]) for X in Q], SIZE["h2"] - 2, Wd - 56, 54, Heading=True)
            Ts = B.FitAll([str(X.get("text", "")) for X in Q if X.get("text")], GROW["detail"], Wd - 56, Ht - 104)
            for I, Quad in enumerate(Q):
                Lx = Left + (I % 2) * (Wd + GAP)
                Ty = Top + (I // 2) * (Ht + GAP)
                B.Card(S, f"Quadrant{I + 1}", Lx, Ty, Wd, Ht, I == Hi, "left")
                Fg = BG if I == Hi else TEXT
                B.Text(S, f"QHeading{I + 1}", Quad["heading"], Lx + 28, Ty + 24, Wd - 56, 54, Hs, Fg,
                       Bold=True, Heading=True)
                if Quad.get("text"):
                    B.Text(S, f"QText{I + 1}", Quad["text"], Lx + 28, Ty + 90, Wd - 56, Ht - 104, Ts, Fg)
            if Sl.get("x_axis"):
                B.Text(S, "XAxis", Sl["x_axis"] + " →", Left, BODY_BOTTOM - 40, W - Left - M, 40, SIZE["label"],
                       MUTED, Align=PP_ALIGN.CENTER)
            if Sl.get("y_axis"):
                Tb = B.Text(S, "YAxis", Sl["y_axis"] + " →", M - 200 + 30, Top + Ht, 400, 40, SIZE["label"],
                            MUTED, Align=PP_ALIGN.CENTER)
                Tb.rotation = -90
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Matrix")
            return

    def Email(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Calls = self.ReadingOrder(Sl.get("callouts", []))
            Ew = 900 if Calls else W - 2 * M
            Lx, Ty, Eh = M, BODY_TOP - 16, BODY_BOTTOM - BODY_TOP + 16
            Keys = [(C.get("target", "body"), C.get("line", 0) if C.get("target", "body") == "body" else 0)
                    for C in Calls]
            PerRow = max([Keys.count(K) for K in Keys] or [0])
            Pad, Dm = 28, 46
            Iw = Ew - 2 * Pad - (Dm + 52 * (PerRow - 1) + 16 if PerRow else 0)
            B.Rect(S, "Email", Lx, Ty, Ew, Eh, BG, Line=MUTED)
            B.Rect(S, "EmailBar", Lx, Ty, Ew, 44, QUIET)
            B.Text(S, "EmailBarText", "Inbox · Message", Lx + Pad, Ty + 8, 420, 36, SIZE["label"], MUTED,
                   Anchor=MSO_ANCHOR.MIDDLE)
            Y = Ty + 44 + 14
            RowY = {}
            for Key, Label in (("from", "From:"), ("to", "To:")):
                if Sl.get(Key):
                    Rh = max(42, B.Need(str(Sl[Key]), 27, Iw - 100))
                    B.Text(S, f"Email{Label[:-1]}Label", Label, Lx + Pad, Y, 96, 42, 27, MUTED, Bold=True)
                    B.Text(S, f"Email{Label[:-1]}", str(Sl[Key]), Lx + Pad + 100, Y, Iw - 100, Rh, 27, TEXT)
                    RowY[(Key, 0)] = Y + 21
                    Y += Rh + 4
            Sh = B.Need(str(Sl["subject"]), 28, Iw, Bold=True)
            B.Text(S, "EmailSubject", str(Sl["subject"]), Lx + Pad, Y + 4, Iw, Sh, 28, TEXT, Bold=True)
            RowY[("subject", 0)] = Y + 4 + 22
            Y += 4 + Sh + 8
            Att = Sl.get("attachment")
            if Att:  # where mail programs show it: under the subject
                Cw = B.LineWidth(f"Attachment: {Att}", 26, Heading=False, Bold=False) / B._kx + 56  # both faces
                Chip = B.Rect(S, "EmailAttachment", Lx + Pad, Y, min(Iw, Cw), 48, QUIET, Line=MUTED)
                B.ShapeText(Chip, f"Attachment: {Att}", 26, TEXT, Bold=False, Heading=False, Inset=14)
                RowY[("attachment", 0)] = Y + 24
                Y += 48 + 12
            B.Rect(S, "EmailRule", Lx + Pad, Y, Ew - 2 * Pad, 2, QUIET)
            Y += 16
            Bottom = Ty + Eh - 20
            Body = [str(X) for X in Sl["body"]]
            Size = 30
            # each paragraph measured by its real wrapped height, so cutting words shows up at once
            while Size > FLOOR and sum(B.Wrapped(X, Size, Iw) for X in Body) + 10 * (len(Body) - 1) > Bottom - Y:
                Size -= 1
            BodyH = sum(B.Wrapped(X, Size, Iw) for X in Body) + 10 * (len(Body) - 1)
            if BodyH > Bottom - Y + 2:
                B.Report("EmailBody", "fit", f"the email body does not fit even at {Size} pt "
                                  f"(needs ~{BodyH:.0f} pt of {Bottom - Y:.0f}); shorten its paragraphs or drop one")
            for I, Para in enumerate(Body):
                Ph = B.Wrapped(Para, Size, Iw)
                Tb = B.Text(S, f"EmailBody{I + 1}", Para, Lx + Pad, Y, Iw, Ph, Size, TEXT)
                if Para.lower().startswith(("http://", "https://", "www.")):  # a link looks like one
                    for Run in Tb.text_frame.paragraphs[0].runs:
                        Run.font.underline = True
                        Run.font.color.theme_color = ACCENT
                RowY[("body", I)] = Y + Size * 0.62
                Y += Ph + 10
            Used = {}
            for I, (Key, Line) in enumerate(Keys):
                Cy = RowY.get((Key, Line), RowY.get(("subject", 0)))
                K = Used.get((Key, Line), 0)
                Used[(Key, Line)] = K + 1
                self.Badge(S, f"Marker{I + 1}", str(I + 1), Lx + Ew - Pad - Dm - 52 * K, Cy - Dm / 2, Dm, ACCENT, BG, 24)
            if Calls:
                Px = Lx + Ew + 44
                Pw = W - M - Px
                Tw = Pw - 72
                Notes = [str(C.get("note", "")) for C in Calls]
                Ns = B.FitAll(Notes, 34, Tw, (Eh - 24 * (len(Notes) - 1)) / len(Notes))
                Hs = [max(B.Need(X, Ns, Tw), 54) for X in Notes]
                Gap = min(56, max(16, (Eh - sum(Hs)) / max(1, len(Hs) - 1)))
                Y = Ty + max(0, (Eh - sum(Hs) - Gap * (len(Hs) - 1)) / 2)
                for I, Note in enumerate(Notes):
                    self.Badge(S, f"CalloutNo{I + 1}", str(I + 1), Px, Y, 54, ACCENT, BG, 26)
                    B.Text(S, f"Callout{I + 1}", Note, Px + 72, Y + 2, Tw, Hs[I], Ns, TEXT)
                    Y += Hs[I] + Gap
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Email")
            return

    def Options(self, S, Sl, Answer):
        """The quiz options as lettered cards; on the answer slide the right ones in the accent."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            Opts = Sl["options"]
            N = len(Opts)
            Explain = Sl.get("explain") if Answer else ""
            Bottom = BODY_BOTTOM - (140 if Explain else 0)
            Grid = N == 4
            Cols = 2 if Grid else 1
            RowsN = 2 if Grid else N
            Wd = (W - 2 * M - GAP * (Cols - 1)) / Cols
            Ht = (Bottom - BODY_TOP - 10 - GAP * (RowsN - 1)) / RowsN
            Pad, D = 28, 64
            Tw = Wd - 2 * Pad - D - 24 - (150 if Answer else 0)
            Ts = B.FitAll([str(O["text"]) for O in Opts], 40, Tw, Ht - 2 * 16)
            for I, O in enumerate(Opts):
                Lx = M + (I % Cols) * (Wd + GAP)
                Ty = BODY_TOP + 10 + (I // Cols) * (Ht + GAP)
                Right = Answer and O.get("correct")
                B.Card(S, f"Option{I + 1}", Lx, Ty, Wd, Ht, bool(Right), None if Answer else "left")
                self.Badge(S, f"Letter{I + 1}", "ABCD"[I], Lx + Pad, Ty + (Ht - D) / 2, D, BG if Right else (
                    MUTED if Answer else ACCENT), ACCENT if Right else BG, 32)
                B.Text(S, f"OptionText{I + 1}", str(O["text"]), Lx + Pad + D + 24, Ty + 12, Tw, Ht - 24, Ts,
                       BG if Right else (MUTED if Answer else TEXT), Anchor=MSO_ANCHOR.MIDDLE)
                if Answer:
                    B.Text(S, f"Verdict{I + 1}", "Correct" if Right else "Not this one", Lx + Wd - Pad - 150, Ty + 12,
                           150, Ht - 24, 24, BG if Right else MUTED, Bold=bool(Right), Align=PP_ALIGN.RIGHT,
                           Anchor=MSO_ANCHOR.MIDDLE)
            if Explain:
                B.Text(S, "Explain", Explain, M, Bottom + 24, W - 2 * M, 116, 34, TEXT)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Options")
            return

    def Quiz(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            self._b.Title(S, Sl["title"])
            self.Options(S, Sl, False)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Quiz")
            return

    def QuizAnswer(self, S, Sl):
        """The slide after a quiz with reveal 'slide': the same options, the right ones marked, the why below."""
        if kS.ErrorMode:
            return
        try:
            Letters = " and ".join("ABCD"[I] for I, O in enumerate(Sl["options"]) if O.get("correct"))
            Default = f"Answer: {Letters}"
            self._b.Title(S, Sl.get("answer_title") or Default)
            self.Options(S, Sl, True)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.QuizAnswer")
            return

    def Risks(self, S, Sl):
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Rk = Sl["risks"]
            Hi = Sl.get("highlight")
            N = len(Rk)
            Wd = (W - 2 * M - GAP * (N - 1)) / N
            Ty, Ht = BODY_TOP + 10, BODY_BOTTOM - BODY_TOP - 10
            Pad = 28
            Iw = Wd - 2 * Pad
            Hs = B.FitAll([str(R["risk"]) for R in Rk], 40, Iw, 150, Heading=True)
            Hh = max(B.Need(str(R["risk"]), Hs, Iw, Heading=True) for R in Rk)
            Levels = [LEVELS.get(str(R.get(Key, "medium")).lower(), "MEDIUM") for R in Rk
                      for Key in ("likelihood", "impact")]
            ChipS, ChipH, RowH = 26, 44, 52
            # a chip is as wide as its longest word plus padding, so 'MEDIUM' never breaks
            ChipW = max(kMeasure.SafeWidth(L, B.Major, ChipS * B.TypeScale(), True) for L in Levels) / B._kx + 36
            ChipsTop = Ty + EDGE + Pad + Hh + 18
            MitTop = ChipsTop + 2 * RowH + 20
            Room = Ty + Ht - Pad - MitTop - 44
            Ms = B.FitAll([str(R.get("mitigation", "")) for R in Rk], GROW["detail"], Iw, Room)
            Fills = {"HIGH": (ACCENT, BG), "MEDIUM": (TEXT, BG), "LOW": (BG, TEXT)}
            for I, R in enumerate(Rk):
                Lx = M + I * (Wd + GAP)
                B.Card(S, f"Risk{I + 1}", Lx, Ty, Wd, Ht, False, "top", EdgeOn=Hi is None or I == Hi)
                B.Text(S, f"RiskName{I + 1}", str(R["risk"]), Lx + Pad, Ty + EDGE + Pad, Iw, Hh, Hs, TEXT, Bold=True,
                       Heading=True)
                for J, (Key, Label) in enumerate((("likelihood", "Likelihood"), ("impact", "Impact"))):
                    Level = LEVELS.get(str(R.get(Key, "medium")).lower(), "MEDIUM")
                    Cy = ChipsTop + J * RowH
                    B.Text(S, f"{Label}{I + 1}", Label, Lx + Pad, Cy, Iw - ChipW - 8, ChipH, 28, MUTED,
                           Anchor=MSO_ANCHOR.MIDDLE)
                    Chip = B.Rect(S, f"{Label}Chip{I + 1}", Lx + Pad + Iw - ChipW, Cy, ChipW, ChipH, Fills[Level][0],
                                  Line=TEXT if Level == "LOW" else None)
                    B.ShapeText(Chip, Level, ChipS, Fills[Level][1])
                B.Rect(S, f"MitigationRule{I + 1}", Lx + Pad, MitTop - 10, Iw, 2, SOFT)
                B.Text(S, f"MitigationLabel{I + 1}", "Mitigation", Lx + Pad, MitTop, Iw, 40, 26, ACCENT if B.Themed
                       else MUTED, Bold=True)
                if R.get("mitigation"):
                    B.Text(S, f"Mitigation{I + 1}", str(R["mitigation"]), Lx + Pad, MitTop + 44, Iw, Room, Ms, TEXT)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Risks")
            return

    def DecisionSize(self, Sl, Wd, Size=40, Who=None):
        """(text size, text width, text height, box height) of the decision box Wd wide - measured once, so a
        layout can reserve exactly the room the box will take."""
        if kS.ErrorMode:
            return Size, Wd, 0, 0
        try:
            B = self._b
            Pad = 36
            Iw = Wd - 2 * Pad - EDGE
            Text = str(Sl["decision"])
            Ds = B.Fit([Text], Size, Iw, 170, Heading=True, Bold=True, What="the decision")
            Dh = B.Need(Text, Ds, Iw, Heading=True)
            for Try in (Ds - 2, Ds - 4):  # a line saved by a slightly smaller size beats a mostly empty second line
                if Try >= 30 and B.Need(Text, Try, Iw, Heading=True) < Dh - Try * 0.6:
                    Ds = Try
                    break
            Bw = B.BalancedWidth(Text, Ds, Iw)  # no lone last word in the ask
            Dh = B.Need(Text, Ds, Bw, Heading=True)
            return Ds, Bw, Dh, Pad + KICKER_H + 10 + Dh + (52 if Who else 0) + Pad - 10
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.DecisionSize")
            return Size, Wd, 0, 0

    def DecisionBox(self, S, Sl, Lx, Ty, Wd, Size=40, Name="Decision", MinHt=0, Who=None):
        """The ask in a box of its own: the text colour as fill, an accent edge, a tracked label and the decision
        itself. Returns the box height (1440 grid); Ty is its top. MinHt stretches the box (a figure beside it),
        its text block centred. Who ([(label, value)]) adds the owner / date line inside the box, under the ask."""
        if kS.ErrorMode:
            return 0
        try:
            B = self._b
            Pad = 36
            Iw = Wd - 2 * Pad - EDGE
            Text = str(Sl["decision"])
            Ds, Bw, Dh, Natural = self.DecisionSize(Sl, Wd, Size, Who)
            Ht = max(Natural, MinHt)
            Off = (Ht - Natural) / 2
            B.Rect(S, f"{Name}Box", Lx, Ty, Wd, Ht, TEXT)
            B.Rect(S, f"{Name}Edge", Lx, Ty, EDGE, Ht, ACCENT)
            Tb = B.Text(S, f"{Name}Label", str(Sl.get("decision_label") or "Decision requested").upper(),
                        Lx + EDGE + Pad, Ty + Off + Pad - 6, Iw, KICKER_H, KICKER_SIZE, BG, Bold=True)
            for R in Tb.text_frame.paragraphs[0].runs:
                R.font._rPr.set("spc", "200")
            B.Text(S, Name, Text, Lx + EDGE + Pad, Ty + Off + Pad + KICKER_H + 4, Bw, Dh, Ds, BG, Bold=True,
                   Heading=True)
            X, Y = Lx + EDGE + Pad, Ty + Off + Pad + KICKER_H + 4 + Dh + 12
            for Label, Val in Who or []:  # 'OWNER Executive team   BY Today' on one line inside the box
                Lw = kMeasure.SafeWidth(Label.upper(), B.Minor, KICKER_SIZE * B.TypeScale(), True) * 1.25 / B._kx + 12
                Tb = B.Text(S, f"{Label}Label", Label.upper(), X, Y + 4, Lw, 36, KICKER_SIZE, BG, Bold=True,
                            Anchor=MSO_ANCHOR.MIDDLE)
                for R in Tb.text_frame.paragraphs[0].runs:
                    R.font._rPr.set("spc", "200")
                Vw = min(Lx + Wd - Pad - X - Lw, kMeasure.SafeWidth(Val, B.Major, 30 * B.TypeScale(), True)
                         * 1.1 / B._kx + 16)
                B.Text(S, f"{Label}Value", Val.replace(" ", NBSP) if len(Val) <= 16 else Val, X + Lw, Y, Vw, 44,
                       30, BG, Bold=True, Heading=True, Anchor=MSO_ANCHOR.MIDDLE)
                X += Lw + Vw + 48
            return Ht
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.DecisionBox")
            return 0

    def Decision(self, S, Sl):
        """A statement with a decision: the claim as the title, the ask in a box, then who and by when, then the
        reasons - the designed close of a deck that asks for something."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Top = BODY_TOP + 14
            Fig = Sl.get("figure") if isinstance(Sl.get("figure"), dict) else None
            Lw = W - 2 * M  # the who/by line and the reasons run the full width under the ask
            Dw = Lw - (FIGURE_W + GAP if Fig else 0)
            Who = [(Label, str(Sl[Key])) for Key, Label in (("owner", "Owner"), ("date", "By")) if Sl.get(Key)]
            Bh = self.DecisionBox(S, Sl, M, Top, Dw, 44 if Fig else 46,
                                  MinHt=(FIGURE_TREND_H if Fig and Fig.get("trend") else FIGURE_H) if Fig else 0,
                                  Who=Who if Fig else None)  # beside a figure the owner line goes in the box
            if Fig:  # the number the decision moves, beside the ask: a tile (with its trend as labelled bars)
                self.Tiles(S, [Fig], None, W - M - FIGURE_W, Top, FIGURE_W, Bh, True, 96, 32)
            Y = Top + Bh + (24 if Fig else 36)
            if Who and not Fig:
                X = M
                for I, (Label, Val) in enumerate(Who):
                    B.Text(S, f"{Label}Label", Label.upper(), X, Y, 300, KICKER_H, KICKER_SIZE, MUTED, Bold=True)
                    Vw = min(M + Lw - X, kMeasure.SafeWidth(Val, B.Major, 36 * B.TypeScale(), True) / B._kx + 40)
                    B.Text(S, f"{Label}Value", Val.replace(" ", NBSP) if len(Val) <= 16 else Val, X, Y + KICKER_H + 4,
                           Vw, 54, 36, TEXT, Bold=True, Heading=True)
                    X += max(Vw, 300) + 80
                Y += KICKER_H + 4 + 54 + 30
            Points = [str(X) for X in Sl.get("points") or []]
            if Sl.get("support"):
                Support = str(Sl["support"])
                Sw = Lw - 200
                Room = BODY_BOTTOM - Y - 24
                if Points:  # the reasons follow: the support line takes only the height it needs
                    Room = min(Room, B.Need(Support, B.Fit([Support], 36, Sw, Room, What="the support line",
                                                           Floor=FLOOR), Sw))
                Ss = B.Fit([Support], 36 if Points else 40, Sw, Room, What="the support line", Floor=FLOOR)
                B.Rect(S, "SupportRule", M, Y, Lw, 2, QUIET)
                B.Text(S, "Support", Support, M, Y + 24, Sw, Room, Ss, MUTED, Floor=FLOOR)
                Y += 24 + Room + 12
            if Points:  # then the reasons: numbered cards in a row under the ask
                self.ReasonCards(S, Points, M, Y + 12, Lw, BODY_BOTTOM - Y - 12)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Decision")
            return

    def Value(self, Val):
        """A metric value that cannot wrap: every space no-break, so '112 %' and '5.9 M' stay whole."""
        if kS.ErrorMode:
            return ""
        try:
            return str(Val).strip().replace(" ", NBSP)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Value")
            return ""

    @staticmethod
    def TrendText(Mt):
        """A metric's trend as words, for a tile with no room for its bars: 'Q1 4.1 → Q4 5.9'."""
        if kS.ErrorMode:
            return ""
        try:
            Vals = [V for V in Mt.get("trend") or [] if isinstance(V, (int, float))]
            if not Vals:
                return ""
            Names = [str(X) for X in (Mt.get("trend_labels") or [])]
            First = kSlidePatterns.TrendNumber(Vals[0], Vals)
            Last = kSlidePatterns.TrendNumber(Vals[-1], Vals)
            if len(Names) >= 2:
                return f"{Names[0]} {First} \u2192 {Names[-1]} {Last}"
            return f"{First} \u2192 {Last}"
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.TrendText")
            return ""

    @staticmethod
    def TrendNumber(V, Values):
        """A trend value as text, with as many decimals as the most precise value in its series."""
        if kS.ErrorMode:
            return ""
        try:
            Places = max((len(f"{X:g}".split(".")[1]) if "." in f"{X:g}" else 0) for X in Values
                         if isinstance(X, (int, float)))
            return f"{V:.{min(Places, 2)}f}"
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.TrendNumber")
            return ""

    def Trend(self, S, I, Values, Periods, Lx, Ty, Wd, Ht, On):
        """A metric's trend as labelled bars (zero-based, so the change is honest): the first and the last value
        written above their bars, the first and last period (trend_labels) below them, the latest bar in the
        accent. Two values read as a before/after; every figure is at the label floor or above."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            Vals = [V for V in Values if isinstance(V, (int, float))]
            Top = max(Vals) if Vals else 0
            if not Vals or Top <= 0:
                return
            N = len(Values)
            Fh = TREND_LABEL * 1.3
            Ph = Fh if Periods else 0
            BarsTop, BarsH = Ty + Fh, Ht - Fh - Ph
            Slot = Wd / N
            Bw = min(Slot * (0.5 if N == 2 else 0.66), 120)
            Edge = [J for J, V in enumerate(Values) if isinstance(V, (int, float))]
            First, Last = Edge[0], Edge[-1]
            Names = [str(X) for X in (Periods or [])]
            Ends = {First: Names[0] if Names else "", Last: Names[-1] if Names else ""}
            for J, V in enumerate(Values):
                if not isinstance(V, (int, float)):
                    continue
                Bh = max(4, BarsH * V / Top)
                # two bars centred in their halves; more spread so the first starts and the last ends the row
                X = Lx + J * Slot + (Slot - Bw) / 2 if N == 2 else Lx + J * Slot + (Slot - Bw) * J / (N - 1)
                Lead = J == Last
                B.Rect(S, f"Trend{I + 1}Bar{J + 1}", X, BarsTop + BarsH - Bh, Bw, Bh,
                       BG if On else (ACCENT if Lead else SOFT))
                if J in (First, Last):  # the figure above its bar, the period below it
                    Lw = Wd * 0.48
                    Fx = min(max(Lx, X + Bw / 2 - Lw / 2), Lx + Wd - Lw) if N == 2 else (Lx if J == First
                                                                                        else Lx + Wd - Lw)
                    Align = PP_ALIGN.CENTER if N == 2 else (PP_ALIGN.LEFT if J == First else PP_ALIGN.RIGHT)
                    B.Text(S, f"Trend{I + 1}{'Last' if Lead else 'First'}", self.TrendNumber(V, Vals), Fx,
                           BarsTop + BarsH - Bh - Fh, Lw, Fh, TREND_LABEL, BG if On else (ACCENT if Lead else MUTED),
                           Bold=True, Align=Align, Anchor=MSO_ANCHOR.BOTTOM)
                    if Ends[J]:
                        B.Text(S, f"Trend{I + 1}{'To' if Lead else 'From'}", Ends[J], Fx, BarsTop + BarsH + 2, Lw,
                               Fh, TREND_LABEL, BG if On else MUTED, Align=Align)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Trend")
            return

    def Metrics(self, S, Sl):
        """Success metrics as a table: metric, baseline, target (in the accent) and, when given, owner and date;
        an optional rule (the stop or go condition) below it."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Items = Sl["rows"]
            Owner = any(R.get("owner") for R in Items)
            Date = any(R.get("date") for R in Items)
            Header = ["Metric", "Baseline", "Target"] + (["Owner"] if Owner else []) + (["By"] if Date else [])
            Keys = ["metric", "baseline", "target"] + (["owner"] if Owner else []) + (["date"] if Date else [])
            Rule = str(Sl.get("rule") or "")
            Rs = B.Fit([Rule], 34, W - 2 * M - EDGE - 64, 120) if Rule else 0
            Rh = B.Need(Rule, Rs, W - 2 * M - EDGE - 64) + 36 if Rule else 0
            Avail = BODY_BOTTOM - BODY_TOP - 10 - (Rh + 32 if Rule else 0)
            Rows = len(Items) + 1
            RowH = min(104, Avail / Rows)
            Size = 34 if RowH >= 90 else (32 if RowH >= 72 else (30 if RowH >= 62 else 27))
            Share = {"metric": 3.0, "baseline": 1.5, "target": 1.5, "owner": 2.0, "date": 1.1}
            Total = sum(Share[K] for K in Keys)
            Widths = [(W - 2 * M) * Share[K] / Total for K in Keys]
            Cols = [[str(R.get(K, "")) for R in Items] for K in Keys]
            Size = min([Size] + [B.FitAll(Col, Size, Widths[C] - 36, RowH - 10) for C, Col in enumerate(Cols)])
            Gf = S.shapes.add_table(Rows, len(Keys), B.X(M), B.Y(BODY_TOP + 10), B.X(W - 2 * M), B.Y(RowH * Rows))
            Gf.name = "MetricsTable"
            Tbl = Gf.table
            for C, Wd in enumerate(Widths):
                Tbl.columns[C].width = B.X(Wd)
            for R in range(Rows):
                Tbl.rows[R].height = B.Y(RowH)
            for C, Head in enumerate(Header):
                self.Cell(Tbl, 0, C, Head, Size - 2, TEXT, BG, Bold=True)
            Hi = Sl.get("highlight")
            for Ri, Row in enumerate(Items, 1):
                On = Hi == Ri - 1
                Fill = ACCENT if On else (QUIET if Ri % 2 == 0 else BG)
                for C, K in enumerate(Keys):
                    Target = K == "target"
                    Fg = BG if On else (ACCENT if Target and B.Themed else TEXT)
                    self.Cell(Tbl, Ri, C, str(Row.get(K, "")), Size, Fill, Fg, Bold=On or Target or K == "metric")
            kSlideText.Alt(Gf, Sl.get("alt") or "Success metrics: " + "; ".join(
                f"{R.get('metric')} from {R.get('baseline')} to {R.get('target')}" for R in Items) + ".")
            if Rule:
                Ry = BODY_TOP + 10 + RowH * Rows + 32
                B.Card(S, "MetricsRule", M, Ry, W - 2 * M, Rh, False, "left")
                B.Text(S, "MetricsRuleText", Rule, M + EDGE + 32, Ry + 18, W - 2 * M - EDGE - 64, Rh - 36, Rs, TEXT,
                       Bold=True, Anchor=MSO_ANCHOR.MIDDLE)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.Metrics")
            return

    def NextSteps(self, S, Sl):
        """Actions with an owner and a date, one numbered row each; an optional decision box above them."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            B.Title(S, Sl["title"])
            Steps = Sl["steps"]
            Top = BODY_TOP + 10
            if Sl.get("decision"):
                Top += self.DecisionBox(S, Sl, M, Top, W - 2 * M, 36) + 28
            N = len(Steps)
            RowH = min(124, (BODY_BOTTOM - Top) / N)
            D = 52
            Dx, Dw = M + D + 28, 250                     # date column
            Ow = 300 if any(X.get("owner") for X in Steps) else 0
            Ax = Dx + Dw + 24
            Aw = W - M - Ax - (Ow + 24 if Ow else 0)
            Actions = [str(X["action"]) for X in Steps]
            As = B.FitAll(Actions, 36, Aw, RowH - 20)
            Dates = [str(X.get("date", "")).replace(" ", NBSP) for X in Steps]
            Ds = B.FitLine([X for X in Dates if X] or ["-"], 32, Dw, Smallest=24)
            Owners = [str(X.get("owner", "")) for X in Steps]
            Os = B.FitAll([X for X in Owners if X], 30, Ow, RowH - 20) if Ow else 0
            for I in range(N):
                Ty = Top + I * RowH
                if I:
                    B.Rect(S, f"StepRule{I + 1}", Dx, Ty, W - M - Dx, 2, QUIET)
                self.Badge(S, f"StepNo{I + 1}", str(I + 1), M, Ty + (RowH - D) / 2, D, ACCENT, BG, 26)
                if Dates[I]:
                    B.Text(S, f"When{I + 1}", Dates[I], Dx, Ty + 10, Dw, RowH - 20, Ds, ACCENT if B.Themed else TEXT,
                           Bold=True, Heading=True, Anchor=MSO_ANCHOR.MIDDLE)
                B.Text(S, f"Action{I + 1}", Actions[I], Ax, Ty + 10, Aw, RowH - 20, As, TEXT,
                       Anchor=MSO_ANCHOR.MIDDLE)
                if Ow and Owners[I]:
                    B.Text(S, f"Owner{I + 1}", Owners[I], W - M - Ow, Ty + 10, Ow, RowH - 20, Os, MUTED,
                           Align=PP_ALIGN.RIGHT, Anchor=MSO_ANCHOR.MIDDLE)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.NextSteps")
            return

    def ProcessRows(self, S, Sl):
        """Five steps: one numbered row each (label, then its detail beside it), so the text stays large instead
        of squeezing into five narrow columns."""
        if kS.ErrorMode:
            return
        try:
            B = self._b
            Steps = Sl["steps"]
            Hi = Sl.get("highlight")
            N = len(Steps)
            Gap = 14
            Ht = (BODY_BOTTOM - BODY_TOP - 10 - Gap * (N - 1)) / N
            D = min(64, Ht - 24)
            Lx = M + EDGE + 28 + D + 28
            Lw = 380
            Dx = Lx + Lw + 32
            Dw = W - M - 32 - Dx
            Ls = B.FitAll([str(X["label"]) for X in Steps], 40, Lw, Ht - 16, Heading=True, Bold=True)
            Short = [str(X["label"]) for X in Steps if len(str(X["label"]).split()) <= 3]
            if Short:  # a short label stays on one line
                Ls = min(Ls, B.FitLine(Short, Ls, Lw, Smallest=LABEL_MIN + 4))
            Details = [str(X.get("detail", "")) for X in Steps]
            Ds = B.FitAll([X for X in Details if X], 36, Dw, Ht - 16) if any(Details) else 0
            for I, St in enumerate(Steps):
                Ty = BODY_TOP + 10 + I * (Ht + Gap)
                On = I == Hi
                B.Card(S, f"StepRow{I + 1}", M, Ty, W - 2 * M, Ht, On, "left")
                self.Badge(S, f"StepNo{I + 1}", str(I + 1), M + EDGE + 28, Ty + (Ht - D) / 2, D, BG if On else ACCENT,
                           ACCENT if On else BG, 28)
                Fg = BG if On else TEXT
                B.Text(S, f"StepLabel{I + 1}", str(St["label"]), Lx, Ty + 8, Lw, Ht - 16, Ls, Fg, Bold=True,
                       Heading=True, Anchor=MSO_ANCHOR.MIDDLE)
                if Details[I]:
                    B.Text(S, f"StepDetail{I + 1}", Details[I], Dx, Ty + 8, Dw, Ht - 16, Ds, Fg,
                           Anchor=MSO_ANCHOR.MIDDLE)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSlidePatterns.ProcessRows")
            return


class kPlanCheck:
    """The cheap checks --plan runs before anything is built: each slide's visible words against its pattern's
    budget, and a title measured in the deck's own heading font that needs three lines even at 40 pt."""

    @staticmethod
    def Words(Obj, Key=""):
        """Visible words under Obj (a slide), estimated from the spec: notes, ids, data and settings left out."""
        if kS.ErrorMode:
            return 0
        try:
            if Key in PLAN_SKIP:
                return 0
            if isinstance(Obj, dict):
                return sum(kPlanCheck.Words(V, K) for K, V in Obj.items())
            if isinstance(Obj, list):
                return sum(kPlanCheck.Words(V, Key) for V in Obj)
            return kRules.WordCount(Obj) if isinstance(Obj, str) else 0  # the count lint_deck.py uses
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPlanCheck.Words")
            return 0

    @staticmethod
    def Budget(Sl):
        """The lint word budget of the slide's pattern (the variants with points or a decision included)."""
        if kS.ErrorMode:
            return 999
        try:
            from lint_deck import PATTERN_BUDGET  # the same table lint_deck.py judges the built slide by
            Pat = Sl.get("pattern", "")
            if Pat == "statement" and Sl.get("decision"):
                Pat = "decision"
            elif Pat in ("statement", "big_number") and Sl.get("points"):
                Pat = f"{Pat}_points"
            return PATTERN_BUDGET.get(Pat, 999)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPlanCheck.Budget")
            return 999

    @staticmethod
    def Run(Spec):
        """Every pre-build finding as messages 'slide N (id): ...' (warnings, not errors)."""
        if kS.ErrorMode:
            return []
        try:
            Out = []
            Builder = kDeckBuilder(Spec)
            Builder.OpenDeck()
            Builder.Major = kTheme(Builder.Prs.slide_master).Font("+mj-lt")
            for N, Sl in enumerate(Spec.get("slides", []), 1):
                Tag = f"slide {N} ({Sl.get('id', Sl.get('pattern'))})"
                Words, Budget = kPlanCheck.Words(Sl), kPlanCheck.Budget(Sl)
                if Words > Budget * 1.3 and Sl.get("pattern") not in NO_CHROME:
                    Out.append(f"{Tag}: about {Words} visible words, budget {Budget} for this pattern - cut, or move "
                               "detail to the notes (advice: fine on a chart, quote or reference slide)")
                Title = Sl.get("title")
                if Title and Sl.get("pattern") not in NO_CHROME:
                    Full = Builder.TitleLines(str(Title), SIZE["title"], W - 2 * M - 15, Loose=True)
                    Least = Builder.TitleLines(str(Title), 40, W - 2 * M - 15, Loose=True)
                    if Least > 2:
                        Out.append(f"{Tag}: the title takes {Least} lines even at 40 pt ({len(str(Title))} "
                                   "characters) - shorten it to one claim")
                    elif Full > 2:
                        Out.append(f"{Tag}: the title wraps to {Full} lines at {SIZE['title']} pt and will be set "
                                   f"smaller ({len(str(Title))} characters) - shorten it to one claim")
            return Out
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kPlanCheck.Run")
            return []


class kCheckReport:
    """--check: every problem of every class (spec, fit, lint, notes) in one list sorted by slide, each with the
    spec edit that fixes it, then the contact sheet - one round of edits fixes everything."""

    def __init__(self):
        try:
            self.Items = []
            self.Sheet = ""
            self.Auto = []
            self.Passes = []
        except Exception as e:
            kS.GlobalErrorHandler(e, "kCheckReport.__init__")

    @staticmethod
    def SlideOf(Message):
        """The slide number a 'slide N ...' message starts with; 0 for a deck-level one."""
        if kS.ErrorMode:
            return 0
        try:
            Hit = re.match(r"slide (\d+)", str(Message))
            return int(Hit.group(1)) if Hit else 0
        except Exception as e:
            kS.GlobalErrorHandler(e, "kCheckReport.SlideOf")
            return 0

    @staticmethod
    def Suggest(Code, Shape, Spec, Message):
        """The spec edit that fixes one problem."""
        if kS.ErrorMode:
            return ""
        try:
            from _autofix import kAutoFix
            Where = f"slides[{Spec - 1}]" if Spec else "the spec"
            Field = kAutoFix.FieldOf(Shape)
            Target = f"{Where}.{Field}" if Field else (f"{Where} ({Shape})" if Shape else Where)
            if Code in ("fit", "text_overflow", "body_below_floor", "tile_text_below_floor", "footer_band"):
                return f"cut words in {Target}, move detail to the notes, or split the slide"
            if Code in ("fit_line", "unwanted_wrap"):
                return f"shorten the value in {Where} (e.g. '19.8 M' with the unit in the label)"
            if Code == "shape_overlap":
                return "usually follows from text that does not fit on this slide: fix that first, then rebuild"
            if Code == "kicker_title_overlap":
                return f"shorten {Where}.title or its kicker"
            if Code == "title_widow":
                return f"reword {Where}.title (a word more or less) so its last line is not one word"
            if Code in ("word_budget", "plan_words"):
                return f"cut words in {Where}, or move detail to {Where}.notes"
            if Code == "plan_title":
                return f"shorten {Where}.title to one claim"
            if Code == "missing_notes":
                return f"write {Where}.notes (what the speaker says)"
            if Code in ("spec", "spec_warning"):
                return f"edit {Where} as the message says"
            return "see the message"
        except Exception as e:
            kS.GlobalErrorHandler(e, "kCheckReport.Suggest")
            return ""

    def Add(self, Slide, Spec, Severity, Code, Message, Shape=""):
        """One problem."""
        if kS.ErrorMode:
            return
        try:
            self.Items.append({"slide": Slide, "spec_slide": Spec, "severity": Severity, "code": Code,
                               "shape": Shape or "", "message": Message,
                               "edit": kCheckReport.Suggest(Code, Shape, Spec, Message)})
        except Exception as e:
            kS.GlobalErrorHandler(e, "kCheckReport.Add")
            return

    def AddMessages(self, Messages, Severity, Code):
        """Spec errors, spec warnings or plan findings ('slide N ...', spec numbering)."""
        if kS.ErrorMode:
            return
        try:
            for Msg in Messages:
                No = kCheckReport.SlideOf(Msg)
                self.Add(No, No, Severity, Code, Msg)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kCheckReport.AddMessages")
            return

    def AddBuild(self, Builder):
        """The builder's own fit problems and every lint finding on the built deck (missing notes among them)."""
        if kS.ErrorMode:
            return
        try:
            for I in Builder.Issues:
                if I["code"] in ("fit", "fit_line", "footer_band"):
                    self.Add(I["slide"], I["spec"], "error", I["code"], f"slide {I['slide']} ({I['id']}): "
                             f"{I['message']}", I["shape"])
            for I in Builder.LintItems:
                Sev = "error" if I["code"] in LAYOUT_CODES else I["severity"]
                self.Add(I["slide"], I["spec"], Sev, I["code"], I["message"], I.get("shape") or "")
        except Exception as e:
            kS.GlobalErrorHandler(e, "kCheckReport.AddBuild")
            return

    @staticmethod
    def Key(Item):
        """Sort key: slide, then severity, then code."""
        if kS.ErrorMode:
            return (0, 0, "")
        try:
            return (Item["slide"], SEVERITY_ORDER.get(Item["severity"], 3), Item["code"])
        except Exception as e:
            kS.GlobalErrorHandler(e, "kCheckReport.Key")
            return (0, 0, "")

    def Print(self, Out, Slides, Seconds):
        """The summary: counts, then every problem sorted by slide with its edit, the sheet, and one JSON line."""
        if kS.ErrorMode:
            return
        try:
            Items = sorted(self.Items, key=kCheckReport.Key)
            Count = {"error": 0, "warn": 0, "info": 0}
            for I in Items:
                Count[I["severity"]] = Count.get(I["severity"], 0) + 1
            print(f"\n==== check: {Out}  ({Slides} slides, {Seconds:.1f} s) ====")
            print(f"auto-fixes: {len(self.Auto)}" + ("  (the auto: lines above; --no-auto to switch off)"
                                                     if self.Auto else ""))
            print(f"problems: {Count['error']} error(s), {Count['warn']} warning(s), {Count['info']} info")
            for I in Items:
                Where = f"slide {I['slide']}" if I["slide"] else "deck"
                Shape = f" [{I['shape']}]" if I["shape"] else ""
                print(f"  {Where:<9} {I['severity']:<5} {I['code']:<22}{Shape} {I['message']}")
                print(f"  {'':<9} edit: {I['edit']}")
            if self.Sheet:
                print(f"sheet: {self.Sheet}  <- look at it once; fix everything listed above in one edit")
            else:
                print("sheet: (not rendered - LibreOffice or pdftoppm not found; render on Windows instead)")
            print("CHECK-JSON " + json.dumps({"deck": Out, "slides": Slides, "seconds": round(Seconds, 1),
                                              "counts": Count, "auto_fixes": len(self.Auto), "passes": self.Passes,
                                              "sheet": self.Sheet, "problems": Items}, ensure_ascii=False))
        except Exception as e:
            kS.GlobalErrorHandler(e, "kCheckReport.Print")
            return


class kBuildDeckApp:
    """Command line: --print-schema, --list-directions, --plan, or build (exit 2 spec errors, 3 unfit text)."""

    @staticmethod
    def LimitOnly(Errors):
        """True when every spec error is a length or count limit, so --check can still build and report the rest."""
        if kS.ErrorMode:
            return False
        try:
            return all(re.search(r"characters, max|characters; with|items, needs", E) for E in Errors)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kBuildDeckApp.LimitOnly")
            return False

    @staticmethod
    def BuildWithFixes(Spec, Out, Only, Auto, Report):
        """Build; while text does not fit and Auto is on, apply kAutoFix to the slides with problems and build
        again (at most AUTO_PASSES times). Prints each change as an `auto:` line. Returns the last builder."""
        if kS.ErrorMode:
            return None
        try:
            from _autofix import kAutoFix
            Builder = None
            for Pass in range(1, AUTO_PASSES + 2):
                Builder = kDeckBuilder(Spec)
                Builder.Build(Out, Only)
                if kS.ErrorMode:
                    return Builder
                Report.Passes.append({"pass": Pass, "problems": len(Builder.Issues)})
                if not Auto or not Builder.Issues or Pass > AUTO_PASSES:
                    break
                Changes = kAutoFix.Apply(Spec, Builder.Issues, AllowSplit=not Only)
                if not Changes:
                    break
                for No, Id, Path, What in Changes:
                    print(f"auto: slide {No} ({Id}) {Path}: {What}")
                    Report.Auto.append({"spec_slide": No, "id": Id, "field": Path, "change": What, "pass": Pass})
            return Builder
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kBuildDeckApp.BuildWithFixes")
            return None

    @staticmethod
    def Render(Out, Report):
        """--check: render the deck with LibreOffice and make the contact sheet, when LibreOffice is installed."""
        if kS.ErrorMode:
            return
        try:
            from contact_sheet import kContactSheet
            from render_lo import kLoRenderer
            if not kLoRenderer.Available():
                return
            Dir = os.path.splitext(Out)[0] + "-render"
            Files = kLoRenderer.Render(Out, Dir)
            if Files:
                Report.Sheet = kContactSheet.Build(Files, 4, 480, os.path.join(Dir, "contact.png")) or ""
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kBuildDeckApp.Render")
            return

    def Run(self):
        if kS.ErrorMode:
            return 1
        try:
            Ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
            Ap.add_argument("spec", nargs="?")
            Ap.add_argument("--out")
            Ap.add_argument("--lint", action="store_true", help="run lint_deck.py on the result")
            Ap.add_argument("--check", action="store_true",
                            help="build + lint + render_lo --sheet, then one summary of every problem (sorted by "
                                 "slide, each with its spec edit) and a CHECK-JSON line")
            Ap.add_argument("--no-auto", action="store_true",
                            help="report unfit text only; do not drop filler words, move detail to the notes or "
                                 "split slides")
            Ap.add_argument("--force", action="store_true", help="build even if the spec breaks a pattern limit")
            Ap.add_argument("--list-directions", action="store_true")
            Ap.add_argument("--print-schema", action="store_true", help="print the spec's JSON Schema")
            Ap.add_argument("--plan", action="store_true", help="print the story and the pre-build checks, and stop")
            Ap.add_argument("--slides", help="build only these spec slides, e.g. 1,3 or 2-4 (the showcase-first "
                                             "step): same theme, page numbers and kickers as in the whole deck")
            Args = Ap.parse_args()
            if Args.print_schema:
                print(json.dumps(kSpecSchema.SpecSchema(), indent=2, ensure_ascii=False))
                return 0
            if Args.list_directions:
                for D in kSpecSchema.Directions():
                    print(f"{D['id']:<18} {D['tone']:<10} {D['heading']} / {D['body']}  #{D['accent']} on "
                          f"#{D['background']}  — {D['mood']}")
                return 0
            if not Args.spec or not (Args.out or Args.plan):
                Ap.error("spec and --out are required (or --plan)")
            Started = time.time()
            Spec = kSpecLoader.Load(Args.spec)
            if Spec is None:
                return 1
            Only = kSpecCheck.SlideSet(Args.slides, len(Spec.get("slides", []))) if Args.slides else None
            Schema = kSpecCheck.SchemaErrors(Spec)  # on the spec as written, before figure tokens are filled in
            Tokens = kFigures.Resolve(Spec, kSlidePatterns.Share)
            Report = kCheckReport()
            if not Args.no_auto and not Only:  # "allow_split": true and more bullets than the limit: split first
                from _autofix import kAutoFix
                for No, Id, Path, What in kAutoFix.PreSplit(Spec, LIMITS["bullets"]["items"][1]):
                    print(f"auto: slide {No} ({Id}) {Path}: {What}")
                    Report.Auto.append({"spec_slide": No, "id": Id, "field": Path, "change": What, "pass": 0})
                if Report.Auto:
                    Schema = kSpecCheck.SchemaErrors(Spec)
            Errors = kSpecCheck.ForSlides(kSpecCheck.CheckSpec(Spec) + Schema + Tokens, Only)
            Warnings = kSpecCheck.ForSlides(kSpecCheck.Warnings(Spec) + kFigures.SpecWarnings(Spec), Only)
            if Args.plan:
                kSpecLoader.PrintPlan(Spec)
                Pre = kPlanCheck.Run(Spec) if not Errors or self.LimitOnly(Errors) else []
                for E in Errors:
                    print(f"spec: {E}")
                for Wn in Warnings:
                    print(f"spec warning: {Wn}")
                for Pc in Pre:
                    print(f"plan: {Pc}")
                print(f"\nplan: {len(Errors)} spec error(s), {len(Warnings)} warning(s), {len(Pre)} pre-build "
                      "finding(s)" + (" - fix the spec lines, weigh the plan lines, then build with --check"
                                     if Errors or Warnings or Pre else " - build with --check"))
                return 2 if Errors else 0
            for E in Errors:
                print(f"spec: {E}", file=sys.stderr)
            for Wn in Warnings:
                print(f"spec warning: {Wn}", file=sys.stderr)
            Report.AddMessages(Errors, "error", "spec")
            Report.AddMessages(Warnings, "warn", "spec_warning")
            if Errors and not Args.force and not (Args.check and self.LimitOnly(Errors)):
                if Args.check:
                    Report.Print(Args.out, 0, time.time() - Started)
                return 2  # spec mistakes (expected state): listed above, nothing written
            Builder = self.BuildWithFixes(Spec, Args.out, Only, not Args.no_auto, Report)
            if kS.ErrorMode or Builder is None:
                return 1
            N = len(Builder.Numbers)
            print(f"{N} slides -> {Args.out}" + (f" (spec slides {Args.slides} only)" if Only else ""))
            for Pr in Builder.Problems:
                print(f"fit: {Pr}", file=sys.stderr)
            for Nn in Builder.NoNotes:
                print(f"notes: {Nn} has no speaker notes - write what the speaker says (reference/CONTENT.md)",
                      file=sys.stderr)
            Code = 0
            if Args.lint and not Args.check:
                Code = subprocess.run([sys.executable, os.path.join(HERE, "lint_deck.py"), Args.out]).returncode
            if Args.check:
                sys.stderr.flush()
                Report.AddBuild(Builder)
                self.Render(Args.out, Report)
                if kS.ErrorMode:
                    return 1
                Report.Print(Args.out, N, time.time() - Started)
            if Errors and not Args.force:
                return 2
            return 3 if Builder.Problems else Code  # 3: text that does not fit (expected state), listed above
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kBuildDeckApp.Run")
            return 1


if __name__ == "__main__":
    kRun.Main(kBuildDeckApp)
