"""Read a deck's or template's theme: colours, fonts, layouts and their placeholders. Any OS.

    uvx --with python-pptx python scripts/extract_theme.py template.potx            # JSON
    uvx --with python-pptx python scripts/extract_theme.py template.pptx --markdown > brand-spec.md

Use it before building from a company template: take colours from the theme slots (accent1-6,
dk1/lt1...) and fonts from the major (headings) / minor (body) pair, and pick slide layouts by
their placeholders instead of drawing boxes on blank slides. A .potx is read like a .pptx.
"""
import argparse
import io
import json
import os
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu

from kShared import ToolReportableException, kRun, kS, kToolException

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
SLOTS = ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
         "hlink", "folHlink"]
SLOT_USE = {"dk1": "main text", "lt1": "main background", "dk2": "secondary dark", "lt2": "secondary light",
            "accent1": "primary accent — the one highlight per slide", "hlink": "links", "folHlink": "visited links"}


class kThemeReader:
    """Reads theme colours, fonts and layouts from a .pptx/.potx."""

    @staticmethod
    def OpenAny(Path):
        """python-pptx refuses the .potx content type; copy to a .pptx in memory and patch the type."""
        if kS.ErrorMode:
            return None
        try:
            if not Path.lower().endswith((".potx", ".potm")):
                return Presentation(Path)
            Buffer = io.BytesIO()  # in memory: no temp file to clean up
            with zipfile.ZipFile(Path) as Source, zipfile.ZipFile(Buffer, "w", zipfile.ZIP_DEFLATED) as Target:
                for Item in Source.infolist():
                    Data = Source.read(Item.filename)
                    if Item.filename == "[Content_Types].xml":
                        Data = Data.replace(b"presentationml.template.main+xml",
                                            b"presentationml.presentation.main+xml")
                    Target.writestr(Item, Data)
            Buffer.seek(0)
            return Presentation(Buffer)
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kThemeReader.OpenAny(file={Path})")
            return None

    @staticmethod
    def Colour(Node):
        """The hex of an a:srgbClr or a:sysClr child, or None."""
        if kS.ErrorMode:
            return None
        try:
            if Node is None:
                return None
            Srgb = Node.find("a:srgbClr", NS)
            if Srgb is not None:
                return Srgb.get("val")
            System = Node.find("a:sysClr", NS)
            return System.get("lastClr") if System is not None else None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeReader.Colour")
            return None

    @staticmethod
    def ThemeOf(Master):
        """The parsed theme XML of a slide master."""
        if kS.ErrorMode:
            return None
        try:
            return etree.fromstring(Master.part.part_related_by(RT.THEME).blob)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeReader.ThemeOf")
            return None

    @staticmethod
    def Colours(Scheme):
        """Every slot's colour of a colour scheme ({} when there is none)."""
        if kS.ErrorMode:
            return {}
        try:
            if Scheme is None:
                return {}
            return {Slot: kThemeReader.Colour(Scheme.find(f"a:{Slot}", NS)) for Slot in SLOTS}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeReader.Colours")
            return {}

    @staticmethod
    def Fonts(FontScheme):
        """The major/minor latin typefaces of a font scheme."""
        if kS.ErrorMode:
            return {}
        try:
            if FontScheme is None:
                return {"major_latin": None, "minor_latin": None}
            return {"major_latin": FontScheme.find("a:majorFont/a:latin", NS).get("typeface"),
                    "minor_latin": FontScheme.find("a:minorFont/a:latin", NS).get("typeface")}
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeReader.Fonts")
            return {}

    @staticmethod
    def Extract(Path):
        """The theme as a JSON-ready dict."""
        if kS.ErrorMode:
            return None
        try:
            if not os.path.isfile(Path):
                raise ToolReportableException(f"not found: {Path}")
            Deck = kThemeReader.OpenAny(Path)
            Master = Deck.slide_master
            Theme = kThemeReader.ThemeOf(Master)
            Scheme = Theme.find(".//a:clrScheme", NS)
            Result = {
                "file": Path,
                "slide_size_pt": [round(Emu(Deck.slide_width).pt, 1), round(Emu(Deck.slide_height).pt, 1)],
                "theme_name": Theme.get("name"),
                "colour_scheme": Scheme.get("name") if Scheme is not None else None,
                "colours": kThemeReader.Colours(Scheme),
                "fonts": kThemeReader.Fonts(Theme.find(".//a:fontScheme", NS)),
                "layouts": [],
            }
            for Extra in list(Deck.slide_masters)[1:]:  # templates can carry several masters, each with its own theme
                ExtraTheme = kThemeReader.ThemeOf(Extra)
                Result.setdefault("additional_themes", []).append({
                    "theme_name": ExtraTheme.get("name"),
                    "colours": kThemeReader.Colours(ExtraTheme.find(".//a:clrScheme", NS)),
                    "fonts": kThemeReader.Fonts(ExtraTheme.find(".//a:fontScheme", NS)),
                    "layouts": [Layout.name for Layout in Extra.slide_layouts]})
            for Index, Layout in enumerate(Master.slide_layouts):
                Placeholders = []
                for Placeholder in Layout.placeholders:
                    Format = Placeholder.placeholder_format
                    Placeholders.append({
                        "idx": Format.idx, "type": str(Format.type).split(".")[-1].split(" ")[0],
                        "name": Placeholder.name,
                        "box_pt": [round(Emu(Value or 0).pt) for Value in
                                   (Placeholder.left, Placeholder.top, Placeholder.width, Placeholder.height)]})
                Result["layouts"].append({"index": Index, "name": Layout.name, "placeholders": Placeholders})
            return Result
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kThemeReader.Extract(file={Path})")
            return None

    @staticmethod
    def Markdown(Theme):
        """A brand-spec.md skeleton from an Extract() result."""
        if kS.ErrorMode:
            return ""
        try:
            Colours = Theme["colours"]
            Lines = [f"# Brand spec — {Theme['theme_name'] or 'theme'}", "",
                     f"Extracted from `{Theme['file']}` by scripts/extract_theme.py. Slide size: "
                     f"{Theme['slide_size_pt'][0]:g} × {Theme['slide_size_pt'][1]:g} pt.", "",
                     "## Fonts", "", f"- Headings (major): **{Theme['fonts']['major_latin']}**",
                     f"- Body (minor): **{Theme['fonts']['minor_latin']}**", "",
                     "## Colours", "", "| Slot | Hex | Use |", "|---|---|---|"]
            for Slot in SLOTS:
                Lines.append(f"| {Slot} | `#{Colours.get(Slot)}` | {SLOT_USE.get(Slot, 'chart series / secondary accent')} |")
            Lines += ["", "## Layouts", "", "| # | Layout | Placeholders |", "|---|---|---|"]
            for Layout in Theme["layouts"]:
                Types = ", ".join(Placeholder["type"] for Placeholder in Layout["placeholders"]) or "—"
                Lines.append(f"| {Layout['index']} | {Layout['name']} | {Types} |")
            for Extra in Theme.get("additional_themes", []):
                Accents = ", ".join(f"`#{Extra['colours'].get(f'accent{Number}')}`" for Number in range(1, 7))
                Lines += ["", f"## Additional theme — {Extra['theme_name']}", "",
                          f"Fonts: {Extra['fonts']['major_latin']} / {Extra['fonts']['minor_latin']}. "
                          "Accents: " + Accents + ".",
                          f"Layouts: {', '.join(Extra['layouts'])}."]
            Lines += ["", "## Fill in by hand", "", "- Logo file and where it goes:", "- Imagery style:",
                      "- Voice and tone:", "- Words to avoid:", ""]
            return "\n".join(Lines)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kThemeReader.Markdown")
            return ""


class kExtractThemeApp:
    """Command line."""

    def Run(self):
        if kS.ErrorMode:
            return 1
        try:
            Parser = argparse.ArgumentParser()
            Parser.add_argument("file")
            Parser.add_argument("--markdown", action="store_true", help="write a brand-spec.md skeleton instead of JSON")
            Args = Parser.parse_args()
            Theme = kThemeReader.Extract(Args.file)
            if Theme is None:
                return 1
            print(kThemeReader.Markdown(Theme) if Args.markdown else json.dumps(Theme, indent=2, ensure_ascii=False))
            return 0
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kExtractThemeApp.Run")
            return 1


if __name__ == "__main__":
    kRun.Main(kExtractThemeApp)
