"""Lint a .pptx for the defect codes in reference/AUDIT.md - from the file alone, on any OS.

    uvx --with python-pptx --with pillow python scripts/lint_deck.py deck.pptx
    uvx --with python-pptx --with pillow python scripts/lint_deck.py deck.pptx --json
    uvx --with python-pptx --with pillow python scripts/lint_deck.py deck.pptx --room-depth 45
    uvx --with python-pptx --with pillow python scripts/lint_deck.py deck.pptx --fix --out fixed.pptx

No PowerPoint and no rendering: geometry, text, fonts, colours, pictures, charts and alt text are
read from the OOXML. What it cannot see - real line breaks, text that overflows its box, how the
slide looks - still needs a render (render_slides.py on Windows, render_lo.py elsewhere).

Exit code: 1 if any finding has severity "error" (or "warn" with --fail-on warn), else 0.
--fix applies only safe fixes (deletes empty placeholders on slides that have other content)
and writes to --out; it never overwrites the input.
"""
import argparse
import io
import json
import os
import re
import sys
from collections import Counter

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

from kShared import ToolReportableException, kRun, kS, kToolException
from _rules import (CARTOON_HOSTS, DEFAULT_FACES, INSIGHT_WORDS, OFFICE_DEFAULT_SERIES, ORDINAL_RE,
                    STOCK_HOSTS, STOP_WORDS, kRules)
from _theme import kTheme
from _measure import kMeasure
from _figures import kFigures

PT = 12700  # EMU per point
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a16": "http://schemas.microsoft.com/office/drawing/2014/main"}
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
TEXT_PLACEHOLDERS = TITLE_TYPES | {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.OBJECT}
LABEL_MAX_PT = 12.0        # at or below: caption tier, never a body-floor finding (scaled per deck: kLintDeck.LabelMaxPt)
BODY_MIN_WORDS = 4         # a paragraph with fewer words is a label
TITLE_MAX_CHARS = 55
MAX_BULLETS = 7
MAX_COLOURS = 5
OVERLAP_WARN_PT2 = 4.0     # thresholds of the author's PowerPoint add-in: >= 4 pt2 warn, >= 200 pt2 error
OVERLAP_ERROR_PT2 = 200.0
EDGE_TOLERANCE_PT = 2.0
STRETCH_TOLERANCE = 0.03
SCALE = 1.0                # slide width / 960 pt (set per deck: kLintDeck.Scale)
IMAGE_EXT = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".emf", ".wmf", ".svg", ".webp")
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
SEVERITY_ORDER = {"error": 0, "warn": 1, "info": 2}
IDENTITY_TF = (1.0, 0.0, 1.0, 0.0)

# build_deck.py patterns, recognised by the shape names the builder gives them (first match wins), and the
# visible-word budget each one reads well with on a 960-pt slide's default of 12 (reference/LAYOUT.md#anchor-types-and-word-budgets).
PATTERN_MARKERS = [("DecisionBox+Support", "decision"), ("DecisionBox+OwnerValue", "decision"), ("MetricsTable", "metrics"), ("Action1", "next_steps"), ("EmailBar", "email"), ("CostTable", "cost_table"), ("Option1", "quiz"), ("Risk1", "risks"),
                   ("ChartCard", "kpi_chart"), ("Card1+Chart", "kpi_chart"), ("Chart", "chart"), ("Value1", "kpi"), ("HeroNumber+Points", "big_number_points"), ("HeroNumber", "big_number"),
                   ("Points", "bullets"), ("StatementPoint1", "statement_points"), ("Point1", "bullets"), ("Heading1", "compare"), ("StepLabel1", "process"), ("Rail", "timeline"),
                   ("QuoteMark", "quote"), ("Table", "table"), ("Photo", "image"), ("Quadrant1", "matrix"),
                   ("Subtitle", "title"), ("Eyebrow", "section"), ("Support", "statement"),
                   ("DecisionBox", "decision"), ("SectionPanel", "section"), ("CoverPanel", "title"),
                   ("AccentRule", "statement"), ("AnchorBar", "statement")]
CHROME_NAMES = {"Footer", "PageNumber", "Kicker", "CoverFooter"}  # the builder's deck furniture: not content
SINGLE_LINE_WORDS = 3      # a box this short (words) whose height holds one line is a single-line role
WRAP_MARGIN = 0.96         # measure single-line roles against 96 % of the width: other renderers set wider
PATTERN_BUDGET = {"title": 15, "section": 10, "statement": 20, "big_number": 12, "kpi": 32, "bullets": 30,
                  "compare": 32, "process": 36, "timeline": 30, "quote": 60, "chart": 16, "table": 999, "image": 16,
                  "matrix": 999, "email": 130, "kpi_chart": 26, "cost_table": 40, "quiz": 40, "risks": 64,
                  "metrics": 999, "next_steps": 60, "decision": 55,
                  "big_number_points": 30, "statement_points": 40}
LOOSE_LINE = 1.28          # line height of the loosest renderer (LibreOffice, serif faces): where a title's ink starts
LABEL_FLOOR_PT = 16.0      # the label floor on a 960-pt slide (24 pt on Full HD); text in metric tiles stays above it
DATA_LABELS = re.compile(r"Trend\d+(?:First|Last|From|To)|ShareLabel")
TILE_TEXT = re.compile(r"(?:Label|Note)\d+|Trend\d+(?:First|Last|From|To)|Lead(?:Label|Note)|ShareLabel")
MEASURABLE = re.compile(r"\d|[<>\u2264\u2265=\u00b1]")  # a target needs a number, date or comparison
COVER_PATTERNS = {"title", "section"}  # no figure_without_source: a cover's numbers are the talk's own headline

GROUP_TF = {}  # shape_id -> (sx, tx, sy, ty): child coordinates -> slide coordinates (pt)
GROWN = {}     # shape_id -> estimated height (pt) of a "resize to fit text" box once its text is laid out


class kFindings:
    """The findings of one lint run, plus the body floor it used."""

    def __init__(self):
        try:
            self.Items = []
            self.Floor = None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kFindings.__init__")

    def Add(self, Slide, Severity, Code, Message, Shape=None, ShapeId=None):
        """Record one finding."""
        if kS.ErrorMode:
            return None
        try:
            self.Items.append({"slide": Slide, "severity": Severity, "code": Code,
                               "shape": Shape, "shape_id": ShapeId, "message": Message})
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kFindings.Add")
            return None


class kLintDeck:
    """Stateless lint checks over a python-pptx Presentation. The per-slide caches are shared module dicts."""

    GroupTf = GROUP_TF
    Grown = GROWN
    LabelMaxPt = LABEL_MAX_PT
    Scale = SCALE

    # ------------------------------------------------------------ helpers

    @staticmethod
    def Rect(Sh):
        """(left, top, width, height) in pt as drawn on the slide. Shapes inside groups are mapped out of
        the group's child coordinate space; a box turned 90/270 degrees swaps width and height."""
        if kS.ErrorMode:
            return None
        try:
            L, T, W, H = (Emu(Sh.left or 0).pt, Emu(Sh.top or 0).pt, Emu(Sh.width or 0).pt, Emu(Sh.height or 0).pt)
            Sx, Tx, Sy, Ty = GROUP_TF.get(Sh.shape_id, IDENTITY_TF)
            L, T, W, H = L * Sx + Tx, T * Sy + Ty, W * Sx, H * Sy
            Rot = round(getattr(Sh, "rotation", 0) or 0) % 180
            if Sh.shape_id in GROWN and GROWN[Sh.shape_id] > H:
                H = GROWN[Sh.shape_id]
            if Rot == 90:
                Cx, Cy = L + W / 2, T + H / 2
                L, T, W, H = Cx - H / 2, Cy - W / 2, H, W
            return (L, T, W, H)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.Rect")
            return None

    @staticmethod
    def AttrPt(El, Key):
        """An EMU attribute of an XML element, in pt."""
        if kS.ErrorMode:
            return 0.0
        try:
            return int(El.get(Key)) / PT
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.AttrPt")
            return 0.0

    @staticmethod
    def Inset(Bp, Key, Default):
        """A bodyPr inset (lIns, tIns ...) in pt, or Default when it is not set."""
        if kS.ErrorMode:
            return 0.0
        try:
            return int(Bp.get(Key)) / PT if Bp is not None and Bp.get(Key) else Default
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.Inset")
            return 0.0

    @staticmethod
    def Walk(Shapes, Tf=IDENTITY_TF):
        """Every shape (a list), descending into groups. python-pptx reports group children in the group's
        own child coordinate space (a:chOff/a:chExt), so record how to map them back onto the slide."""
        if kS.ErrorMode:
            return []
        try:
            Out = []
            for Sh in Shapes:
                if Tf != IDENTITY_TF:
                    GROUP_TF[Sh.shape_id] = Tf
                Out.append(Sh)
                if Sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    Xfrm = Sh._element.find("p:grpSpPr/a:xfrm", NS)
                    Inner = Tf
                    if Xfrm is not None:
                        Off, Ext = Xfrm.find("a:off", NS), Xfrm.find("a:ext", NS)
                        ChOff, ChExt = Xfrm.find("a:chOff", NS), Xfrm.find("a:chExt", NS)
                        if None not in (Off, Ext, ChOff, ChExt):
                            V = kLintDeck.AttrPt
                            Kx = V(Ext, "cx") / V(ChExt, "cx") if V(ChExt, "cx") else 1.0
                            Ky = V(Ext, "cy") / V(ChExt, "cy") if V(ChExt, "cy") else 1.0
                            # local: x_parent = off + (x_child - chOff) * k ; then apply the outer transform
                            Sx, Tx = Kx, V(Off, "x") - V(ChOff, "x") * Kx
                            Sy, Ty = Ky, V(Off, "y") - V(ChOff, "y") * Ky
                            Inner = (Tf[0] * Sx, Tf[0] * Tx + Tf[1], Tf[2] * Sy, Tf[2] * Ty + Tf[3])
                    Out += kLintDeck.Walk(Sh.shapes, Inner)
            return Out
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.Walk")
            return []

    @staticmethod
    def PhType(Sh):
        """Placeholder type, or None for a plain shape."""
        if kS.ErrorMode:
            return None
        try:
            try:
                return Sh.placeholder_format.type if Sh.is_placeholder else None
            except ValueError:  # ERROR-SUPPRESSED-JUSTIFIED: python-pptx raises for a placeholder type it doesn't know; treat it as untyped
                return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.PhType")
            return None

    @staticmethod
    def IsTitle(Sh):
        """True for a title placeholder."""
        if kS.ErrorMode:
            return False
        try:
            return kLintDeck.PhType(Sh) in TITLE_TYPES
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.IsTitle")
            return False

    @staticmethod
    def LevelSize(LstStyle, Level):
        """Size from an <a:lstStyle>/<p:txStyles> list for paragraph level (1-based), or None."""
        if kS.ErrorMode:
            return None
        try:
            if LstStyle is None:
                return None
            Node = LstStyle.find(f"a:lvl{Level}pPr/a:defRPr", NS)
            Sz = Node.get("sz") if Node is not None else None
            return int(Sz) / 100 if Sz else None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.LevelSize")
            return None

    @staticmethod
    def InheritedSize(Sh, Para):
        """Font size of a paragraph when its runs don't set one: shape lstStyle -> layout placeholder
        -> master text style. Returns None when nothing in the chain sets a size."""
        if kS.ErrorMode:
            return None
        try:
            Level = (Para.level or 0) + 1
            Body = Sh.text_frame._txBody
            Size = kLintDeck.LevelSize(Body.find("a:lstStyle", NS), Level)
            if Size:
                return Size
            Pt = kLintDeck.PhType(Sh)
            if Pt is not None:
                try:
                    LayoutPh = Sh.part.slide.slide_layout.placeholders.get(idx=Sh.placeholder_format.idx)
                except (AttributeError, KeyError):  # ERROR-SUPPRESSED-JUSTIFIED: no layout placeholder to inherit from; fall through to the master
                    LayoutPh = None
                if LayoutPh is not None and LayoutPh.has_text_frame:
                    Size = kLintDeck.LevelSize(LayoutPh.text_frame._txBody.find("a:lstStyle", NS), Level)
                    if Size:
                        return Size
            Master = Sh.part.slide.slide_layout.slide_master._element
            Style = "titleStyle" if Pt in TITLE_TYPES else ("bodyStyle" if Pt is not None else "otherStyle")
            Size = kLintDeck.LevelSize(Master.find(f"p:txStyles/p:{Style}", NS), Level)
            return Size or (18.0 if Pt is None else None)  # PowerPoint's default text box size is 18 pt
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.InheritedSize")
            return None

    @staticmethod
    def ParaSize(Sh, Para):
        """Smallest run size of a paragraph, or its inherited size."""
        if kS.ErrorMode:
            return None
        try:
            Sizes = [R.font.size.pt for R in Para.runs if R.font.size is not None]
            return min(Sizes) if Sizes else kLintDeck.InheritedSize(Sh, Para)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.ParaSize")
            return None

    @staticmethod
    def SolidRgb(Fill):
        """RRGGBB of a solid RGB fill, or None."""
        if kS.ErrorMode:
            return None
        try:
            try:
                if Fill.type == 1:  # MSO_FILL.SOLID
                    return str(Fill.fore_color.rgb)
            except (AttributeError, TypeError, ValueError):  # ERROR-SUPPRESSED-JUSTIFIED: a theme/none colour has no rgb; None means "not a plain RGB fill"
                pass
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.SolidRgb")
            return None

    @staticmethod
    def RunRgb(Run):
        """RRGGBB set directly on a run, or None."""
        if kS.ErrorMode:
            return None
        try:
            try:
                return str(Run.font.color.rgb) if Run.font.color and Run.font.color.type is not None else None
            except (AttributeError, TypeError, ValueError):  # ERROR-SUPPRESSED-JUSTIFIED: a theme colour has no rgb; None means "not a plain RGB colour"
                return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.RunRgb")
            return None

    @staticmethod
    def AltText(Sh):
        """Alt text, treating a bare file name (python-pptx and some tools write one) as missing."""
        if kS.ErrorMode:
            return ""
        try:
            Nv = Sh._element.find(".//p:cNvPr", NS)
            Text = (Nv.get("descr") or Nv.get("title") or "").strip() if Nv is not None else ""
            return "" if Text.lower().endswith(IMAGE_EXT) and " " not in Text else Text
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.AltText")
            return ""

    @staticmethod
    def IsDecorative(Sh):
        """True when the shape is marked decorative."""
        if kS.ErrorMode:
            return False
        try:
            Nv = Sh._element.find(".//p:cNvPr", NS)
            if Nv is None:
                return False
            Dec = Nv.find(".//a16:decorative", NS)
            return Dec is not None and Dec.get("val") in ("1", "true")
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.IsDecorative")
            return False

    @staticmethod
    def PictureStretch(Sh):
        """Relative difference between the picture's shown aspect and its cropped native aspect."""
        if kS.ErrorMode:
            return 0.0
        try:
            try:
                from PIL import Image
                WPx, HPx = Image.open(io.BytesIO(Sh.image.blob)).size
            except Exception:  # ERROR-SUPPRESSED-JUSTIFIED: unreadable or vector image (EMF/SVG): nothing to compare
                return 0.0
            Cw = WPx * (1 - Sh.crop_left - Sh.crop_right)
            Ch = HPx * (1 - Sh.crop_top - Sh.crop_bottom)
            if Cw <= 0 or Ch <= 0 or not Sh.width or not Sh.height:
                return 0.0
            return (Sh.width / Sh.height) / (Cw / Ch) - 1
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.PictureStretch")
            return 0.0

    @staticmethod
    def ShapeFill(S, Theme):
        """Solid fill colour of a shape (theme-resolved), or None."""
        if kS.ErrorMode:
            return None
        try:
            Sp = S._element.find(".//p:spPr", NS)
            if Sp is None:
                return None
            Solid = Sp.find("a:solidFill", NS)
            if Solid is not None:
                return Theme.ColourIn(Solid)
            Ref = S._element.find("p:style/a:fillRef", NS)  # shape style default fill (idx 0 = none)
            if Ref is not None and Ref.get("idx", "0") != "0" and Sp.find("a:noFill", NS) is None \
                    and Sp.find("a:gradFill", NS) is None and Sp.find("a:blipFill", NS) is None:
                return Theme.ColourIn(Ref)
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.ShapeFill")
            return None

    @staticmethod
    def HasOtherFill(S):
        """True for a gradient, picture or pattern fill."""
        if kS.ErrorMode:
            return False
        try:
            Sp = S._element.find(".//p:spPr", NS)
            return Sp is not None and (Sp.find("a:gradFill", NS) is not None or Sp.find("a:blipFill", NS) is not None
                                       or Sp.find("a:pattFill", NS) is not None)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.HasOtherFill")
            return False

    @staticmethod
    def RunColour(Run, Theme):
        """Theme-resolved colour set on a run, or None."""
        if kS.ErrorMode:
            return None
        try:
            RPr = Run._r.find("a:rPr", NS)
            Solid = RPr.find("a:solidFill", NS) if RPr is not None else None
            return Theme.ColourIn(Solid) if Solid is not None else None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.RunColour")
            return None

    @staticmethod
    def SortKey(Item):
        """Print order: by slide, then error, warn, info."""
        if kS.ErrorMode:
            return (0, 0)
        try:
            return (Item["slide"], SEVERITY_ORDER[Item["severity"]])
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.SortKey")
            return (0, 0)

    # ------------------------------------------------------------ checks

    @staticmethod
    def PatternOf(Shapes):
        """The build_deck.py pattern a slide was built with, from its shape names; "" for any other slide."""
        if kS.ErrorMode:
            return ""
        try:
            Names = {S.name for S in Shapes}
            for Marker, Pattern in PATTERN_MARKERS:
                if all(Part in Names for Part in Marker.split("+")):
                    return Pattern
            return ""
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.PatternOf")
            return ""

    @staticmethod
    def LintSlide(N, Slide, Sw, Sh_, Floor, Budget, F, Theme):
        """Every per-slide check. Returns the slide's title text."""
        if kS.ErrorMode:
            return None
        try:
            GROUP_TF.clear()  # shape ids repeat across slides
            GROWN.clear()
            Rect = kLintDeck.Rect
            Shapes = kLintDeck.Walk(Slide.shapes)
            Titles = [S for S in Shapes if kLintDeck.IsTitle(S)]
            TitleText = Titles[0].text_frame.text.strip() if Titles and Titles[0].has_text_frame else ""
            SlideBg = Theme.Background(Slide)

            # titles
            if not Titles:
                F.Add(N, "warn", "missing_title", "No title placeholder. Every slide needs a title, even a hidden one, "
                      "for navigation and screen readers.")
            elif not TitleText:
                F.Add(N, "warn", "empty_title", "Title placeholder is empty.", Titles[0].name)
            else:
                Lines = kLintDeck.TitleLines(Titles[0], Theme)
                Display = (kLintDeck.ParaSize(Titles[0], Titles[0].text_frame.paragraphs[0]) or 0) >= 40 * kLintDeck.Scale
                if Lines > (3 if Display else 2) or (Lines == 0 and len(TitleText) > TITLE_MAX_CHARS):
                    F.Add(N, "warn", "headline_too_long",
                          (f"Title takes {Lines} lines at its size in its box" if Lines else
                           f"Title is {len(TitleText)} characters (> {TITLE_MAX_CHARS})") + "; cut it to one claim "
                          "of one or two lines.", Titles[0].name)
                kLintDeck.KickerCheck(N, Shapes, Titles[0], Theme, F)
                if "\n" in TitleText or "\v" in TitleText:
                    F.Add(N, "warn", "headline_two_line", "Title has a hard line break; make it one line or split "
                          "headline and subhead.", Titles[0].name)
                if kRules.LooksLikeLabel(TitleText):
                    F.Add(N, "info", "title_is_label", f"'{TitleText}' reads as a topic label, not a claim "
                          "(fine for dividers, agenda and Q&A).", Titles[0].name)

            ContentAll = [S for S in Shapes if S.shape_type != MSO_SHAPE_TYPE.GROUP]
            for S in ContentAll:  # first pass: estimate text layout, so later checks see grown boxes
                if S.has_text_frame and S.text_frame.text.strip():
                    kLintDeck.FitCheck(N, S, Theme, F)
                    if not kLintDeck.IsTitle(S):
                        kLintDeck.SpillCheck(N, S, ContentAll, F)
                    kLintDeck.TileFloorCheck(N, S, F)
                if getattr(S, "has_table", False) and S.has_table and S.name == "MetricsTable":
                    kLintDeck.TargetCheck(N, S, F)
            Pattern = kLintDeck.PatternOf(Shapes)
            Words, Colours, RawFills, Shadows, Accents = 0, set(), 0, 0, set()
            Sizes, BigTokens, ContrastDone = [], Counter(), False
            Content = [S for S in Shapes if not kLintDeck.IsTitle(S) and S.shape_type != MSO_SHAPE_TYPE.GROUP]
            TextShapes = [S for S in Shapes if S.has_text_frame and S.text_frame.text.strip()]
            for S in Content:
                Fill = kLintDeck.ShapeFill(S, Theme)
                if Fill:
                    Colours.add(Fill)
                El = S._element
                Sp = El.find(".//p:spPr", NS)
                if Sp is not None:
                    Srgb = Sp.find("a:solidFill/a:srgbClr", NS)
                    if Srgb is not None and not Theme.IsThemeHex(Srgb.get("val", "")):
                        RawFills += 1
                    if Sp.find("a:effectLst/a:outerShdw", NS) is not None:
                        Shadows += 1
                    Stops = Sp.findall("a:gradFill/a:gsLst/a:gs/a:srgbClr", NS)
                    if Stops and max(kTheme.Saturation(C.get("val", "000000")) for C in Stops) > 0.45:
                        F.Add(N, "warn", "gradient_high_chroma", "Saturated gradient that isn't built from theme colours "
                              "(a common machine-made look); use a solid accent or a two-stop brand gradient.", S.name)
                Accents.update(C.get("val") for C in El.iter(f"{{{NS['a']}}}schemeClr")
                               if C.get("val", "").startswith("accent"))

                # bounds
                L, T, W, H = Rect(S)
                if L < -EDGE_TOLERANCE_PT or T < -EDGE_TOLERANCE_PT or L + W > Sw + EDGE_TOLERANCE_PT \
                        or T + H > Sh_ + EDGE_TOLERANCE_PT:
                    F.Add(N, "warn", "offslide_shape", "Shape extends outside the slide.", S.name)
                if El.find(".//a:hlinkClick", NS) is not None and El.find(".//p:cNvPr/a:hlinkClick", NS) is not None \
                        and (W < 32 or H < 32):
                    F.Add(N, "warn", "tiny_click_target", f"Clickable shape is {W:.0f} × {H:.0f} pt; make it at least "
                          "44 × 44 pt for touch.", S.name)

                # placeholders left empty
                Pt = kLintDeck.PhType(S)
                if Pt in TEXT_PLACEHOLDERS and S.has_text_frame and not S.text_frame.text.strip():
                    if any(O is not S and (O.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.CHART,
                                                            MSO_SHAPE_TYPE.TABLE)
                                           or (O.has_text_frame and O.text_frame.text.strip()
                                               and not kLintDeck.IsTitle(O)))
                           for O in Content):
                        F.Add(N, "error", "unused_placeholder", "Empty placeholder next to real content shows "
                              "'Click to add text' in edit view and confuses screen readers.", S.name, S.shape_id)

                # text
                if S.has_text_frame and S.text_frame.text.strip():
                    Paras = [P for P in S.text_frame.paragraphs if P.text.strip()]
                    Chrome = S.name in CHROME_NAMES
                    Data = bool(DATA_LABELS.fullmatch(S.name or ""))  # a trend's figures and periods: chart labels
                    Words += 0 if Chrome or Data else sum(len(P.text.split()) for P in Paras)
                    Text = S.text_frame.text
                    if len(Paras) > MAX_BULLETS:
                        F.Add(N, "warn", "too_many_bullets", f"{len(Paras)} paragraphs (> {MAX_BULLETS}); split the "
                              "slide or reveal one at a time.", S.name)
                    if re.search(r"\blorem ipsum\b|\bdolor sit amet\b", Text, re.I):
                        F.Add(N, "error", "lorem_ipsum", "Placeholder 'lorem ipsum' text left in.", S.name)
                    for P in Paras:
                        Size = kLintDeck.ParaSize(S, P)
                        if Size and kLintDeck.LabelMaxPt < Size < Floor and len(P.text.split()) >= BODY_MIN_WORDS:
                            F.Add(N, "warn", "body_below_floor", f"{Size:g} pt body text (floor {Floor:g} pt): "
                                  f"'{P.text.strip()[:40]}'", S.name)
                            break
                    for P in Paras:
                        PtText = P.text.strip()
                        Size = kLintDeck.ParaSize(S, P) or 18.0
                        if not Chrome:
                            Sizes.append(Size)
                        if len(PtText) <= 24 and kRules.HasEmoji(PtText):
                            F.Add(N, "warn", "emoji_as_icon", f"Emoji used as an icon ('{PtText}'); use a real icon "
                                  "with a text label.", S.name)
                        if len(PtText) >= 8 and PtText.endswith(("…", "...")):
                            F.Add(N, "warn", "truncated_text", f"Text ends in an ellipsis ('…{PtText[-24:]}'); "
                                  "is something cut off?", S.name)
                        if len(PtText) >= 80 and P.alignment == 2:  # PP_ALIGN.CENTER
                            F.Add(N, "warn", "centered_long_body", "Long centred text is hard to read; left-align body "
                                  "text and keep centring for one-line headlines and quotes.", S.name)
                        if len(PtText) >= 90 and Size <= 28 * kLintDeck.Scale:
                            Cpl = kRules.CharsPerLine(Emu(S.width or 0).pt, Size)
                            if Cpl > 75:
                                F.Add(N, "warn", "measure_too_wide", f"About {Cpl:.0f} characters per line "
                                      "(comfortable: 45–75); narrow the box or raise the size.", S.name)
                        if Size >= 30 * kLintDeck.Scale:  # display type; grown body text (up to ~27 pt on 960) is not
                            BigTokens.update(Wd.lower() for Wd in re.findall(r"[^\W\d_]{4,}", PtText)
                                             if Wd.lower() not in STOP_WORDS)
                    # contrast: run colour (or theme text colour) against shape fill (or slide background)
                    Backing = Fill
                    if Backing is None and not kLintDeck.HasOtherFill(S):  # a card or band drawn behind the text box
                        for O in Content:
                            if O is S:
                                break
                            if kRules.Contains(Rect(O), Rect(S), Tol=2.0):
                                if kLintDeck.HasOtherFill(O) or O.shape_type == MSO_SHAPE_TYPE.PICTURE:  # text on a photo
                                    Backing = "?"
                                else:
                                    Backing = kLintDeck.ShapeFill(O, Theme) or Backing
                    if not ContrastDone and not kLintDeck.HasOtherFill(S) and Backing != "?":
                        Bg = Backing or SlideBg
                        Fill = Backing
                        DefaultFg = Theme.Colours.get(Theme.Map.get("tx1", "dk1"))
                        for P in Paras:
                            Size = kLintDeck.ParaSize(S, P) or 18.0
                            for R in P.runs:
                                if not R.text.strip():
                                    continue
                                Fg = kLintDeck.RunColour(R, Theme) or (DefaultFg if Fill else None)
                                if not Fg or not Bg:
                                    continue
                                Need = 3.0 if kRules.IsLargeText(Size / kLintDeck.Scale, bool(R.font.bold)) else 4.5
                                Ratio = kRules.ContrastRatio(Fg, Bg)
                                if Ratio < Need:
                                    F.Add(N, "error" if Ratio < 3.0 else "warn", "a11y_low_text_contrast",
                                          f"Text #{Fg} on #{Bg} is {Ratio:.1f}:1 (needs {Need}:1).", S.name)
                                    ContrastDone = True
                                    break
                            if ContrastDone:
                                break

                # pictures and charts
                if S.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.CHART) or getattr(S, "has_chart", False):
                    if not kLintDeck.AltText(S) and not kLintDeck.IsDecorative(S):
                        F.Add(N, "warn", "a11y_missing_alt_text", "No alt text (and not marked decorative).", S.name)
                if getattr(S, "has_table", False) and S.has_table:
                    for Cell in S.table.iter_cells():
                        Tc = Cell._tc
                        Solid = Tc.find("a:tcPr/a:solidFill", NS)
                        Cbg = Theme.ColourIn(Solid) if Solid is not None else SlideBg
                        for P in Cell.text_frame.paragraphs:
                            for R in P.runs:
                                Fg = kLintDeck.RunColour(R, Theme)
                                if not R.text.strip() or not Fg or not Cbg:
                                    continue
                                Size = (R.font.size.pt if R.font.size else 18.0)
                                Need = 3.0 if kRules.IsLargeText(Size / kLintDeck.Scale, bool(R.font.bold)) else 4.5
                                Ratio = kRules.ContrastRatio(Fg, Cbg)
                                if Ratio < Need:
                                    F.Add(N, "error" if Ratio < 3.0 else "warn", "a11y_low_text_contrast",
                                          f"Table cell text #{Fg} on #{Cbg} is {Ratio:.1f}:1 (needs {Need}:1).", S.name)
                                    break
                            else:
                                continue
                            break
                        else:
                            continue
                        break
                if S.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    Stretch = kLintDeck.PictureStretch(S)
                    if abs(Stretch) > STRETCH_TOLERANCE:
                        F.Add(N, "error" if abs(Stretch) > 0.15 else "warn", "picture_stretched",
                              f"Picture is {Stretch:+.0%} {'wider' if Stretch > 0 else 'narrower'} than its source; "
                              "crop instead (scripts/cover_crop.py).", S.name)
                    Hay = (kLintDeck.AltText(S) + " " + S.name + " " + " ".join(
                        Rel.target_ref for Rel in S.part.rels.values() if Rel.is_external)).lower()
                    if any(Host in Hay for Host in STOCK_HOSTS + CARTOON_HOSTS):
                        F.Add(N, "info", "stock_or_cartoon_image", "Looks like a stock photo or generic illustration; "
                              "real product, team or customer images carry more weight.", S.name)
                if getattr(S, "has_chart", False):
                    kLintDeck.LintChart(N, S, F, Theme)

            Allowed = round(Budget * PATTERN_BUDGET[Pattern] / 12) if Pattern else Budget
            if Words > Allowed:
                F.Add(N, "info", "word_budget", f"{Words} visible words (budget {Allowed}"
                      + (f" for a {Pattern} slide" if Pattern else "") + "). Fine for gallery, matrix, "
                      "chart, quote and reference slides; otherwise cut or move to the notes.")
            if len(Colours) > MAX_COLOURS:
                F.Add(N, "warn", "palette_too_many_colours", f"{len(Colours)} distinct fill colours (> {MAX_COLOURS}).")
            if RawFills >= 3:
                F.Add(N, "warn", "off_palette_fill", f"{RawFills} shapes use hard-coded colours outside the theme; "
                      "use theme colours so the deck re-themes cleanly.")
            if Shadows > 3:
                F.Add(N, "warn", "shadow_overuse", f"{Shadows} shapes have drop shadows; reserve elevation for one or two.")
            if len(Accents) >= 4:
                F.Add(N, "warn", "accent_overload", f"{len(Accents)} accent colours on one slide "
                      f"({', '.join(sorted(Accents))}); keep to 3 or fewer, ideally one highlight.")
            for Tok, C in BigTokens.items():
                if C >= 3:
                    F.Add(N, "warn", "repeated_word", f"'{Tok}' appears {C} times in large type; demote the repeats.")
                    break
            Big = sorted(set(Sizes), reverse=True)
            # build_deck.py layouts set their own hierarchy (a heading over its points, a value over its label,
            # cards that share one size); the check is for hand-made slides
            if len(Big) >= 2 and 1.08 < Big[0] / Big[1] < 1.6 and not Pattern:
                F.Add(N, "info", "weak_focal_hierarchy", f"Largest text sizes {Big[0]:g} and {Big[1]:g} pt are too "
                      "close; make one element clearly dominant (≥ 1.6×) or equal.")

            # grid monotony: 4+ body boxes, 3+ of them the same width and top as the first
            Area = Sw * Sh_
            Body = [Rect(S) for S in TextShapes if not kLintDeck.IsTitle(S) and Rect(S)[2] > 1 and Rect(S)[3] > 1
                    and Rect(S)[2] * Rect(S)[3] < 0.92 * Area]
            if len(Body) >= 4:
                Same = sum(1 for Rc in Body[1:] if abs(Rc[2] - Body[0][2]) <= 4 and abs(Rc[1] - Body[0][1]) <= 4)
                if Same >= 3:
                    F.Add(N, "info", "grid_monotony", f"{Same + 1} identical boxes in a row; vary size or emphasis so "
                          "one item leads.")

            # overlaps between text-bearing shapes and pictures (containment = a card/backing, not a defect)
            Boxes = [(S.name, Rect(S)) for S in Content
                     if (S.has_text_frame and S.text_frame.text.strip())
                     or S.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.CHART)]
            for I, (Na, Ra) in enumerate(Boxes):
                for Nb, Rb in Boxes[I + 1:]:
                    Ov = kRules.OverlapArea(Ra, Rb)
                    if Ov >= OVERLAP_WARN_PT2 and not kRules.Contains(Ra, Rb) and not kRules.Contains(Rb, Ra):
                        F.Add(N, "error" if Ov >= OVERLAP_ERROR_PT2 else "warn", "shape_overlap",
                              f"'{Na}' and '{Nb}' overlap by {Ov:.0f} pt².", Na)

            Notes = Slide.notes_slide.notes_text_frame.text if Slide.has_notes_slide else ""
            if not Notes.strip():
                F.Add(N, "info", "missing_notes", "No speaker notes.")
            Shown = " ".join(X.text_frame.text for X in Shapes if X.has_text_frame) + " " + " ".join(
                C.text for X in Shapes if getattr(X, "has_table", False) and X.has_table for C in X.table.iter_cells())
            HasChart = any(getattr(X, "has_chart", False) and X.has_chart for X in Shapes)
            HasFigure = HasChart or re.search(r"\d+(?:[.,]\d+)?\s*(?:%|percent|pt\b|[kKmMbB]n?\b|x\b)|[$€£¥]\s?\d",
                                              Shown)
            Cover = N == 1 or Pattern in COVER_PATTERNS or any(
                kLintDeck.PhType(X) == PP_PLACEHOLDER.CENTER_TITLE for X in Shapes)
            if HasFigure and not Cover and Notes.strip() and not re.search(
                    r"source|doi|https?://|www\.|according to|\(\d{4}\)|©|källa|quelle", Notes, re.I):
                F.Add(N, "info", "figure_without_source", "The slide shows figures but the notes name no source; add "
                      "where the numbers come from (SOURCES: …).")
            return TitleText
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kLintDeck.LintSlide(slide={N})")
            return None

    @staticmethod
    def TitleLines(T, Theme, Loose=False):
        """How many lines the title takes at its size in its box, from font metrics; 0 when its size is unknown.
        Loose: in the widest renderer (the face LibreOffice substitutes when the theme font is not installed)."""
        if kS.ErrorMode:
            return 0
        try:
            Paras = [P for P in T.text_frame.paragraphs if P.text.strip()]
            Size = kLintDeck.ParaSize(T, Paras[0]) if Paras else None
            if not Size:
                return 0
            Bp = T.text_frame._txBody.find("a:bodyPr", NS)
            Width = kLintDeck.Rect(T)[2] - kLintDeck.Inset(Bp, "lIns", 7.2) - kLintDeck.Inset(Bp, "rIns", 7.2)
            if Loose:
                return sum(kMeasure.LooseLines(P.text, Theme.Major, Size, Width, True) for P in Paras)
            return sum(len(kMeasure.LineWords(P.text, Theme.Major, Size, Width, True)) for P in Paras)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.TitleLines")
            return 0

    @staticmethod
    def KickerCheck(N, Shapes, T, Theme, F):
        """A kicker (the small label above a title) must end above the title's ink. The title's ink is measured
        with the line height and the line count of the loosest renderer (LibreOffice draws a missing theme font
        with a wider substitute), from its anchor: a bottom-anchored two-line title grows up."""
        if kS.ErrorMode:
            return
        try:
            Kick = next((S for S in Shapes if S.name == "Kicker" and S.has_text_frame), None)
            Paras = [P for P in T.text_frame.paragraphs if P.text.strip()]
            if Kick is None or not Paras:
                return
            Bp = T.text_frame._txBody.find("a:bodyPr", NS)
            L, Tp, W, H = kLintDeck.Rect(T)
            Size = kLintDeck.ParaSize(T, Paras[0]) or 44.0
            Lines = max(1, kLintDeck.TitleLines(T, Theme, True))
            Ink = Lines * Size * LOOSE_LINE
            Anchor = Bp.get("anchor") if Bp is not None else None
            if Anchor == "b":
                InkTop = Tp + H - kLintDeck.Inset(Bp, "bIns", 3.6) - Ink
            elif Anchor == "ctr":
                InkTop = Tp + (H - Ink) / 2
            else:
                InkTop = Tp + kLintDeck.Inset(Bp, "tIns", 3.6)
            Kl, Kt, Kw, Kh = kLintDeck.Rect(Kick)
            if Kt + Kh > InkTop + 1 and Kl < L + W and L < Kl + Kw:
                F.Add(N, "warn", "kicker_title_overlap", f"The kicker ends at {Kt + Kh:.0f} pt but the title's "
                      f"{Lines} line(s) start at ~{InkTop:.0f} pt; shorten the title to one line, lower its size or "
                      "move the kicker up.", Kick.name)
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kLintDeck.KickerCheck(slide={N})")
            return

    @staticmethod
    def SlideFigures(N, Slide, Contexts, Series):
        """One slide's figures for FigureCheck: chart series and numeric table columns (a total row is a claim on
        its column), the text of each card as one context (its largest number is the figure, the rest says what it
        is), every other text box, and the speaker notes."""
        if kS.ErrorMode:
            return
        try:
            Shapes = list(kLintDeck.Walk(Slide.shapes))
            Cards = [S for S in Shapes if kLintDeck.IsCard(S) and not (S.has_text_frame and S.text_frame.text.strip())]
            Groups = {}
            for S in Shapes:
                if getattr(S, "has_chart", False):
                    for Plot in S.chart.plots:
                        for Ser in Plot.series:
                            Series.append({"slide": N, "name": Ser.name or "chart series", "values": list(Ser.values)})
                if getattr(S, "has_table", False):
                    Rows = [[C.text.strip() for C in R.cells] for R in S.table.rows]
                    Cols, Totals = kFigures.TableColumns({"header": Rows[0] if Rows else [], "rows": Rows[1:]}, N)
                    Series += Cols
                    Contexts += [dict(Ctx, shape=S.name) for Ctx in Totals]
                    continue
                if not S.has_text_frame or not S.text_frame.text.strip():
                    continue
                L, T, W, H = kLintDeck.Rect(S)
                Home = None
                for C in Cards:
                    Cl, Ct, Cw, Ch = kLintDeck.Rect(C)
                    if Ch >= 40 and Cl - 2 <= L and L + W <= Cl + Cw + 2 and Ct - 2 <= T < Ct + Ch - 2 and \
                            (Home is None or Cw * Ch < Home[1]):
                        Home = (C.shape_id, Cw * Ch)
                Groups.setdefault(Home[0] if Home else f"own{S.shape_id}", []).append(S)
            for Members in Groups.values():
                Texts = [M.text_frame.text.strip() for M in Members]
                if len(Members) == 1:
                    Contexts.append({"slide": N, "where": Members[0].name, "shape": Members[0].name, "text": Texts[0],
                                     "primary": None})
                    continue
                Sized = [(kLintDeck.ParaSize(M, M.text_frame.paragraphs[0]) or 0, J) for J, M in enumerate(Members)
                         if kFigures.Numbers(Texts[J])]
                Lead = max(Sized)[1] if Sized else None
                for J, M in enumerate(Members):
                    if J != Lead:
                        Contexts.append({"slide": N, "where": M.name, "shape": M.name, "text": Texts[J],
                                         "primary": None})
                if Lead is not None:
                    Contexts.append({"slide": N, "where": Members[Lead].name, "shape": Members[Lead].name,
                                     "text": " \u00b7 ".join(Texts), "primary": Texts[Lead]})
            if Slide.has_notes_slide:
                Notes = Slide.notes_slide.notes_text_frame.text
                Contexts.append({"slide": N, "where": "notes", "shape": None, "text": Notes, "primary": None})
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kLintDeck.SlideFigures(slide={N})")
            return

    @staticmethod
    def FigureCheck(Prs, F):
        """figure_mismatch: a total, average, change or share written on a slide (or in its notes) that the
        deck's own data contradicts - '21.8 MUSD, sum of four quarters' beside a chart of 4.1, 4.6, 5.2, 5.9."""
        if kS.ErrorMode:
            return
        try:
            Contexts, Series = [], []
            for N, Slide in enumerate(Prs.slides, 1):
                if Slide._element.get("show") != "0":
                    kLintDeck.SlideFigures(N, Slide, Contexts, Series)
            for Found in kFigures.Check(Contexts, Series):
                F.Add(Found["slide"], "warn", "figure_mismatch", Found["message"] + ".", Found.get("shape"))
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.FigureCheck")
            return

    @staticmethod
    def SpillCheck(N, S, Shapes, F):
        """Text that starts on a card (or a frame such as a mocked email) but runs past its bottom: the box
        itself, or the box once grown to its text, ends below the card. The tightest card that holds the box's
        top is the one it belongs to."""
        if kS.ErrorMode:
            return
        try:
            L, T, W, H = kLintDeck.Rect(S)
            Best = None
            for O in Shapes:
                if O is S or (O.has_text_frame and O.text_frame.text.strip()) or not kLintDeck.IsCard(O):
                    continue
                Ol, Ot, Ow, Oh = kLintDeck.Rect(O)
                if Oh < 40 or not (Ol - 2 <= L and L + W <= Ol + Ow + 2 and Ot - 2 <= T < Ot + Oh - 2):
                    continue
                if Best is None or Ow * Oh < Best[2] * Best[3]:
                    Best = (Ol, Ot, Ow, Oh, O.name)
            if Best and T + H > Best[1] + Best[3] + 2:
                F.Add(N, "warn", "text_overflow", f"Text runs to {T + H:.0f} pt, past the bottom of '{Best[4]}' "
                      f"({Best[1] + Best[3]:.0f} pt); cut words, drop a paragraph or enlarge the card.", S.name)
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kLintDeck.SpillCheck(slide={N})")
            return

    @staticmethod
    def IsCard(Sh):
        """A filled (or outlined) autoshape that text can sit on: a card, a tile, a frame - not a picture, chart,
        table, line or placeholder."""
        if kS.ErrorMode:
            return False
        try:
            if Sh.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or Sh.is_placeholder:
                return False
            Sp = Sh._element.find(".//p:spPr", NS)
            return Sp is not None and (Sp.find("a:solidFill", NS) is not None or Sp.find("a:gradFill", NS) is not None
                                       or Sp.find("a:ln/a:solidFill", NS) is not None)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.IsCard")
            return False

    @staticmethod
    def TileFloorCheck(N, S, F):
        """Labels, notes and trend figures inside the builder's metric tiles never go below the label floor
        (16 pt on a 960-pt slide, 24 pt on Full HD): small text in a tile is unreadable from the back."""
        if kS.ErrorMode:
            return
        try:
            if not TILE_TEXT.fullmatch(S.name or ""):
                return
            Floor = LABEL_FLOOR_PT * kLintDeck.Scale
            for P in S.text_frame.paragraphs:
                Size = kLintDeck.ParaSize(S, P)
                if P.text.strip() and Size and Size < Floor - 0.05:
                    F.Add(N, "warn", "tile_text_below_floor", f"{Size:g} pt in a metric tile (label floor "
                          f"{Floor:g} pt): '{P.text.strip()[:30]}'; shorten the label or note, or drop a metric.",
                          S.name)
                    return
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kLintDeck.TileFloorCheck(slide={N})")
            return

    @staticmethod
    def TargetCheck(N, S, F):
        """A success-metrics table: every target needs a number, percent, date or comparison to be judged by."""
        if kS.ErrorMode:
            return
        try:
            Rows = list(S.table.rows)
            Head = [C.text.strip().lower() for C in Rows[0].cells] if Rows else []
            if "target" not in Head:
                return
            Col = Head.index("target")
            for R in Rows[1:]:
                Text = R.cells[Col].text.strip()
                if Text and not MEASURABLE.search(Text):
                    F.Add(N, "warn", "target_not_measurable", f"Target '{Text}' ({R.cells[0].text.strip()[:30]}) has "
                          "no number, percent, date or comparison; write it so the result can be judged (e.g. "
                          "'below 30 %', '>= 95 % of Q4').", S.name)
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kLintDeck.TargetCheck(slide={N})")
            return

    @staticmethod
    def FitCheck(N, S, Theme, F):
        """Estimate wrapped text height from font metrics and flag text that won't fit its box."""
        if kS.ErrorMode:
            return None
        try:
            Tf = S.text_frame
            Body = Tf._txBody
            Bp = Body.find("a:bodyPr", NS)
            if Bp is not None and Bp.get("wrap") == "none":
                return None
            Grows = Bp is not None and Bp.find("a:spAutoFit", NS) is not None
            if round(getattr(S, "rotation", 0) or 0) % 180 == 90:
                return None
            Ins = kLintDeck.Inset
            L, T, W, H = kLintDeck.Rect(S)
            Width = W - Ins(Bp, "lIns", 7.2) - Ins(Bp, "rIns", 7.2)
            Height = H - Ins(Bp, "tIns", 3.6) - Ins(Bp, "bIns", 3.6)
            if Width <= 0 or Height <= 0:
                return None
            Heading = kLintDeck.IsTitle(S)
            Paras = []
            for P in Tf.paragraphs:
                Size = kLintDeck.ParaSize(S, P) or 18.0
                Run = next((R for R in P.runs if R.text.strip()), None)
                Name = Run.font.name if Run is not None and Run.font.name else None
                Family = Theme.Font(Name) if Name else (Theme.Major if Heading else Theme.Minor)
                Bold = bool(Run.font.bold) if Run is not None and Run.font.bold is not None else Heading
                Before = P.space_before.pt if P.space_before is not None else 0
                Paras.append((P.text, Family, Size, Bold, Before))
            Need, Widest, Lines = kMeasure.TextHeight(Paras, Width)
            kLintDeck.WrapCheck(N, S, Paras, Width, Height, Heading, F)
            if Grows:  # "resize shape to fit text": the box will be as tall as its text; overlap/off-slide use that
                Grown = Need + Ins(Bp, "tIns", 3.6) + Ins(Bp, "bIns", 3.6)
                if Grown > H:
                    GROWN[S.shape_id] = Grown
                if Widest > Width + 1:
                    F.Add(N, "warn", "word_breaks", f"A word is wider than its box ({Widest:.0f} pt in {Width:.0f} pt) "
                          "and will break mid-word.", S.name)
                return None
            if Widest > Width + 1:
                F.Add(N, "warn", "word_breaks", f"A word is wider than its box ({Widest:.0f} pt in {Width:.0f} pt) and "
                      "will break mid-word; shorten it, widen the box or lower the size.", S.name)
            Shrink = Bp is not None and Bp.find("a:normAutofit", NS) is not None
            if Need > Height * 1.08:
                if Shrink:
                    F.Add(N, "info", "text_shrinks", f"Text needs ~{Need:.0f} pt in a {Height:.0f} pt box; PowerPoint "
                          "will shrink it to fit — check it stays above the floor.", S.name)
                else:
                    F.Add(N, "warn" if Need < Height * 1.5 else "error", "text_overflow",
                          f"Text needs ~{Need:.0f} pt ({Lines} lines) but the box is {Height:.0f} pt tall; it will "
                          "spill out. Cut words, widen or heighten the box, or split the slide.", S.name)
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kLintDeck.FitCheck(slide={N})")
            return None

    @staticmethod
    def WrapCheck(N, S, Paras, Width, Height, Heading, F):
        """Unwanted line breaks: a short single-line role (a KPI value, a chip, a label, a number) that wraps
        inside its box ('112 %' leaving the '%' on the label below), and a title whose last line is one word."""
        if kS.ErrorMode:
            return
        try:
            if len(Paras) == 1 and not Heading:
                Text, Family, Size, Bold, _ = Paras[0]
                Words = Text.split()
                OneLine = Height < Size * 1.2 * 1.7  # the box only has room for one line
                if Words and len(Words) <= SINGLE_LINE_WORDS and OneLine and len(Text.strip()) > 1:
                    Need = kMeasure.SafeWidth(Text.strip(), Family, Size, Bold)
                    if Need > Width * WRAP_MARGIN + 1:
                        F.Add(N, "warn", "unwanted_wrap", f"'{Text.strip()[:30]}' needs ~{Need:.0f} pt on one line but "
                              f"its box is {Width:.0f} pt wide; it will wrap (e.g. a unit under its number). Lower "
                              "the size, widen the box, or join number and unit with a no-break space.", S.name)
            if Heading and len(Paras) == 1:
                Text, Family, Size, Bold, _ = Paras[0]
                Sub = kMeasure.Substitute(Family)  # LibreOffice draws a missing theme font in another face
                for Face in (Family, Sub) if Sub else (Family,):
                    Lines = kMeasure.LineWords(Text, Face, Size, Width, Bold)
                    if len(Lines) >= 2 and len(Lines[-1]) == 1 and len(Lines[-1][0]) <= 14:
                        Where = "" if Face == Family else f" in {Face} (the face LibreOffice substitutes)"
                        F.Add(N, "warn", "title_widow", f"The title wraps leaving '{Lines[-1][0]}' alone on its last "
                              f"line{Where}; shorten it by a word, narrow the box to balance the lines, or lower the "
                              "size.", S.name)
                        break
            return
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kLintDeck.WrapCheck(slide={N})")
            return

    @staticmethod
    def LintChart(N, S, F, Theme):
        """Chart checks: palette, title, labels, legend, order, number format."""
        if kS.ErrorMode:
            return None
        try:
            Chart = S.chart
            Cx = Chart._chartSpace
            C = {"c": C_NS, "a": NS["a"]}
            Series = Cx.findall(".//c:ser", C)
            Colours = []
            for Ser in Series[:6]:
                Colours.append(Theme.ColourIn(Ser.find("c:spPr/a:solidFill", C)))
            if sum(1 for Col in Colours if Col in OFFICE_DEFAULT_SERIES) >= 2:
                F.Add(N, "warn", "chart_default_palette", "Series use Office default colours; bind them to the "
                      "theme, or highlight only the finding.", S.name)

            if Chart.has_title and Chart.chart_title.has_text_frame:
                T = Chart.chart_title.text_frame.text.strip()
                if T and not any(W in f" {T.lower()} " for W in INSIGHT_WORDS):
                    F.Add(N, "info", "chart_descriptive_title", f"Chart title '{T}' names the data; headline the "
                          "finding instead (e.g. 'East leads Q1, up 8 %').", S.name)

            LabelsOn = any(D.find("c:showVal", C) is not None and D.find("c:showVal", C).get("val") in ("1", "true")
                           for D in Cx.findall(".//c:dLbls", C))
            Gridlines = Cx.find(".//c:valAx/c:majorGridlines", C) is not None
            ValDeleted = Cx.find(".//c:valAx/c:delete", C)
            if LabelsOn and Gridlines and not (ValDeleted is not None and ValDeleted.get("val") in ("1", "true")):
                F.Add(N, "info", "chart_redundant_labels", "Data labels and a gridlined value axis say the same "
                      "thing; keep one.", S.name)

            Legend = Cx.find(".//c:legend", C)
            if Legend is not None:
                Pos = Legend.find("c:legendPos", C)
                Overlay = Legend.find("c:overlay", C)
                if Pos is not None and Pos.get("val") in ("t", "b") \
                        and (Overlay is None or Overlay.get("val") in ("0", "false")):
                    F.Add(N, "warn", "chart_legend_steals_plot", "Legend at the top/bottom takes height from the "
                          "plot; move it to the right, or drop it and colour-code the title.", S.name)

            Names = [(Ser.findtext(".//c:tx//c:v", namespaces=C) or "").strip().lower() for Ser in Series]
            Real = [Col for Col in Colours if Col]
            if len(Names) >= 3 and all(re.match(ORDINAL_RE, Nm) for Nm in Names if Nm) and all(Names) \
                    and len(Real) >= 2 and len({round(kTheme.Hue(Col), 2) for Col in Real}) >= 2:
                F.Add(N, "info", "chart_ordinal_categorical_color", "Ordered series (months, quarters, years) in "
                      "unrelated hues; use one hue from light to dark so the order reads at a glance.", S.name)

            for Fmt in Cx.findall(".//c:valAx/c:numFmt", C):
                Code = Fmt.get("formatCode", "")
                if "_(" in Code or ('"-"' in Code and "*" in Code):
                    F.Add(N, "warn", "chart_accounting_zero_dash", "Accounting number format shows zero as '$-' on "
                          "the axis; use a currency or number format.", S.name)
                    break

            Points = max((len(Ser.findall(".//c:val//c:pt", C)) for Ser in Series), default=0)
            if LabelsOn and (Points > 12 or len(Series) > 1):
                F.Add(N, "info", "chart_label_collision", f"Data labels on {len(Series)} series × {Points} points "
                      "will likely collide; label only the key points or rely on the axis. Check the render.", S.name)
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kLintDeck.LintChart(slide={N})")
            return None

    @staticmethod
    def Lint(Path, Floor, Budget):
        """Lint a deck file. Returns (Presentation, kFindings).
        Sizes are judged relative to a standard 960-pt-wide slide: on a 1440-pt (Full HD) slide an
        18 pt floor becomes 27 pt, because the same text is two-thirds as big on screen."""
        if kS.ErrorMode:
            return None, None
        try:
            GROUP_TF.clear()
            Prs = Presentation(Path)
            F = kFindings()
            Theme = kTheme(Prs.slide_master)
            Sw, Sh_ = Emu(Prs.slide_width).pt, Emu(Prs.slide_height).pt
            Scale = kLintDeck.Scale = max(1.0, Sw / 960)
            Floor, kLintDeck.LabelMaxPt = round(Floor * Scale, 1), round(12.0 * Scale, 1)
            if (round(Sw), round(Sh_)) != (1440, 810):
                Ratio = Sw / Sh_
                F.Add(0, "warn" if abs(Ratio - 16 / 9) > 0.01 else "info", "slide_size",
                      f"Slide size {Sw:g} x {Sh_:g} pt; Full HD is 1440 x 810 pt.")
            Titles = Counter()
            Fonts = Counter()
            Resolved = Counter()
            for N, Slide in enumerate(Prs.slides, 1):
                if Slide._element.get("show") == "0":
                    continue
                Title = kLintDeck.LintSlide(N, Slide, Sw, Sh_, Floor, Budget, F, Theme)
                if Title:
                    Titles[Title.lower()] += 1
                for S in kLintDeck.Walk(Slide.shapes):
                    if S.has_text_frame:
                        for P in S.text_frame.paragraphs:
                            for R in P.runs:
                                if not R.text.strip():
                                    continue
                                Name = R.font.name
                                if Name and not Name.startswith("+"):
                                    Fonts[Name] += 1
                                Face = Theme.Font(Name) if Name else (Theme.Major if kLintDeck.IsTitle(S)
                                                                      else Theme.Minor)
                                if Face:
                                    Resolved[Face.lower()] += 1
            kLintDeck.FigureCheck(Prs, F)
            for T, C in Titles.items():
                if C > 1:
                    F.Add(0, "warn", "duplicate_titles", f"{C} slides share the title '{T}'; screen-reader and "
                          "outline navigation can't tell them apart.")
            if len(Fonts) > 3:
                F.Add(0, "warn", "mixed_font_families", f"{len(Fonts)} font families set directly on text: "
                      f"{', '.join(sorted(Fonts))}. Use the theme fonts (display + body).")
            if Resolved and len(Resolved) <= 1 and set(Resolved) <= DEFAULT_FACES:
                F.Add(0, "info", "default_font_only", f"Only {next(iter(Resolved)).title()} is used. Fine for a "
                      "quick internal deck; for anything branded, pair a distinctive heading face with the body font.")
            F.Floor = Floor
            return Prs, F
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kLintDeck.Lint(file={Path})")
            return None, None

    @staticmethod
    def Fix(Prs, Findings):
        """Safe fixes only: delete empty placeholders flagged as unused. Returns how many, or -1."""
        if kS.ErrorMode:
            return -1
        try:
            Done = 0
            Targets = {(I["slide"], I["shape_id"]) for I in Findings.Items if I["code"] == "unused_placeholder"}
            for N, Slide in enumerate(Prs.slides, 1):
                for S in list(Slide.shapes):
                    if (N, S.shape_id) in Targets and S.is_placeholder and not S.text_frame.text.strip():
                        S._element.getparent().remove(S._element)
                        Done += 1
            return Done
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeck.Fix")
            return -1


class kLintDeckApp:
    """Command line."""

    def Run(self):
        if kS.ErrorMode:
            return 1
        try:
            Ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
            Ap.add_argument("file")
            Ap.add_argument("--floor", type=float, help="body font floor in pt (default 18, or from --room-depth)")
            Ap.add_argument("--room-depth", type=float,
                            help="viewing distance in feet; sets the floor (20->14, 30->18, 50->24, more->28)")
            Ap.add_argument("--budget", type=int, default=12,
                            help="visible words per slide before an info finding (default 12; slides built by "
                                 "build_deck.py get their pattern's budget, scaled by this)")
            Ap.add_argument("--json", action="store_true")
            Ap.add_argument("--fail-on", choices=["error", "warn"], default="error")
            Ap.add_argument("--fix", action="store_true", help="apply safe fixes; needs --out")
            Ap.add_argument("--out")
            A = Ap.parse_args()
            if A.fix and not A.out:
                Ap.error("--fix needs --out (the input is never overwritten)")
            if not os.path.isfile(A.file):
                raise ToolReportableException(f"file not found: {A.file}")
            Floor = A.floor or (kRules.FloorForRoom(A.room_depth) if A.room_depth else 18.0)

            Prs, F = kLintDeck.Lint(A.file, Floor, A.budget)
            if F is None:
                return 1
            Sev = Counter(I["severity"] for I in F.Items)
            if A.json:
                print(json.dumps({"file": A.file, "floor_pt": F.Floor, "counts": dict(Sev), "findings": F.Items},
                                 ensure_ascii=False, indent=2))
            else:
                for I in sorted(F.Items, key=kLintDeck.SortKey):
                    Where = f"slide {I['slide']}" if I["slide"] else "deck"
                    Shape = f" [{I['shape']}]" if I["shape"] else ""
                    print(f"{I['severity']:<5}  {Where:<9} {I['code']:<24}{Shape} {I['message']}")
                print(f"\n{Sev.get('error', 0)} error(s), {Sev.get('warn', 0)} warning(s), {Sev.get('info', 0)} info  "
                      f"(body floor {F.Floor:g} pt)")
            if A.fix:
                Done = kLintDeck.Fix(Prs, F)
                Prs.save(A.out)
                print(f"fixed {Done} issue(s) -> {A.out}", file=sys.stderr)
            Failing = {"error"} if A.fail_on == "error" else {"error", "warn"}
            return 1 if any(I["severity"] in Failing for I in F.Items) else 0
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kLintDeckApp.Run")
            return 1


if __name__ == "__main__":
    kRun.Main(kLintDeckApp)
