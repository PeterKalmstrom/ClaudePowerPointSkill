"""Write two small company-style templates for testing template builds (build_deck.py --template / --inspect).

    python scripts/make_test_templates.py examples/templates

  northwind-16x9.potx  16:9 .potx; deep teal accent, Cambria / Arial (installed, or a metric twin
                       on Linux CI, so it raises no font warning anywhere); layouts renamed (Cover, Chapter Break,
                       Headline Only, Headline and Text, Empty); a logo mark and footer text on the master; one
                       example slide listed in a section (both must be dropped by the build).
  harbor-4x3.pptx      4:3 .pptx; pale orange accent that fails 3:1 on white, a heading font that is not installed;
                       default layout names. Built to trigger the size, contrast and font warnings.

Fictional names; public-safe. Re-running overwrites both files.
"""
import argparse
import io
import os
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Pt

from kShared import kRun, kS, kToolException

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NS = {"a": A, "p": P}
NORTHWIND = {"name": "Northwind", "dk1": "1B2A2F", "lt1": "FFFFFF", "dk2": "3E5056", "lt2": "EEF4F4",
             "accent1": "00666E", "heading": "Cambria", "body": "Arial",
             "layouts": {"Title Slide": "Cover", "Section Header": "Chapter Break", "Title Only": "Headline Only",
                         "Title and Content": "Headline and Text", "Blank": "Empty"}}
HARBOR = {"name": "Harbor", "dk1": "222222", "lt1": "FFFFFF", "dk2": "666666", "lt2": "F2F2F2",
          "accent1": "FFB870", "heading": "Fictional Display Sans", "body": "Arial", "layouts": {}}
SECTIONS = ('<p:extLst xmlns:p="{P}"><p:ext uri="{{521415D9-36F7-43E2-AB2F-B90AF26B5E84}}">'
            '<p14:sectionLst xmlns:p14="{P14}"><p14:section name="Examples" '
            'id="{{6F1C2B9A-1D2E-4C3B-9A8F-0E1D2C3B4A59}}"><p14:sldIdLst><p14:sldId id="{Id}"/></p14:sldIdLst>'
            '</p14:section></p14:sectionLst></p:ext></p:extLst>')


class kTestTemplates:
    """Makes the two test templates."""

    @staticmethod
    def Theme(Prs, T):
        """Write T's colours and fonts into the master's theme."""
        if kS.ErrorMode:
            return
        try:
            Part = Prs.slide_master.part.part_related_by(RT.THEME)
            Root = etree.fromstring(Part.blob)
            Scheme = Root.find(".//a:clrScheme", NS)
            Scheme.set("name", T["name"])
            for Slot in ("dk1", "lt1", "dk2", "lt2", "accent1"):
                Node = Scheme.find(f"a:{Slot}", NS)
                for Child in list(Node):
                    Node.remove(Child)
                etree.SubElement(Node, f"{{{A}}}srgbClr", val=T[Slot])
            Fonts = Root.find(".//a:fontScheme", NS)
            Fonts.set("name", T["name"])
            Fonts.find("a:majorFont/a:latin", NS).set("typeface", T["heading"])
            Fonts.find("a:minorFont/a:latin", NS).set("typeface", T["body"])
            Part._blob = etree.tostring(Root, xml_declaration=True, encoding="UTF-8", standalone=True)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kTestTemplates.Theme")
            return

    @staticmethod
    def Brand(Prs, Text):
        """A logo mark (an accent square) in the master's top right corner and footer text on the master."""
        if kS.ErrorMode:
            return
        try:
            Scratch = Prs.slides.add_slide(Prs.slide_layouts[6])
            Logo = Scratch.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Prs.slide_width - Pt(70), Pt(18),
                                            Pt(44), Pt(44))
            Logo.name = "Logo"
            Logo.fill.solid()
            Logo.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
            Logo.line.fill.background()
            Logo._element.find(".//p:cNvPr", NS).set("id", "999")  # unique among the master's shapes
            Prs.slide_master.shapes._spTree.append(Logo._element)
            Ids = Prs.slides._sldIdLst
            Prs.part.drop_rel(Ids[-1].rId)
            Ids.remove(Ids[-1])
            for Ph in Prs.slide_master.placeholders:
                if Ph.placeholder_format.type == PP_PLACEHOLDER.FOOTER:
                    Ph.text_frame.text = Text
        except Exception as e:
            kS.GlobalErrorHandler(e, "kTestTemplates.Brand")
            return

    @staticmethod
    def Example(Prs):
        """One example slide (as company templates ship), listed in a section."""
        if kS.ErrorMode:
            return
        try:
            Slide = Prs.slides.add_slide(Prs.slide_layouts[0])
            Slide.shapes.title.text = "Example: replace this slide"
            Id = Prs.slides._sldIdLst[-1].get("id")
            Prs.part._element.append(etree.fromstring(SECTIONS.format(P=P, P14=P14, Id=Id)))
        except Exception as e:
            kS.GlobalErrorHandler(e, "kTestTemplates.Example")
            return

    @staticmethod
    def SaveAsPotx(Prs, Path):
        """Save with the template content type, as PowerPoint writes a .potx."""
        if kS.ErrorMode:
            return
        try:
            Buffer = io.BytesIO()
            Prs.save(Buffer)
            with zipfile.ZipFile(Buffer) as Source, zipfile.ZipFile(Path, "w", zipfile.ZIP_DEFLATED) as Target:
                for Item in Source.infolist():
                    Data = Source.read(Item.filename)
                    if Item.filename == "[Content_Types].xml":
                        Data = Data.replace(b"presentationml.presentation.main+xml",
                                            b"presentationml.template.main+xml")
                    Target.writestr(Item, Data)
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kTestTemplates.SaveAsPotx({Path})")
            return

    @staticmethod
    def Make(Dir):
        """Write both templates into Dir; returns their paths (16:9 .potx, 4:3 .pptx)."""
        if kS.ErrorMode:
            return []
        try:
            os.makedirs(Dir, exist_ok=True)
            Wide = Presentation()
            Wide.slide_width, Wide.slide_height = Pt(960), Pt(540)
            kTestTemplates.Theme(Wide, NORTHWIND)
            for Layout in Wide.slide_layouts:
                Layout.name = NORTHWIND["layouts"].get(Layout.name, Layout.name)
            kTestTemplates.Brand(Wide, "Northwind Example Co. - internal")
            kTestTemplates.Example(Wide)
            WidePath = os.path.join(Dir, "northwind-16x9.potx")
            kTestTemplates.SaveAsPotx(Wide, WidePath)
            Square = Presentation()  # python-pptx's default: 720 x 540 pt, 4:3
            kTestTemplates.Theme(Square, HARBOR)
            SquarePath = os.path.join(Dir, "harbor-4x3.pptx")
            Square.save(SquarePath)
            return [WidePath, SquarePath]
        except Exception as e:
            kS.GlobalErrorHandler(e, f"kTestTemplates.Make({Dir})")
            return []


class kMakeTestTemplatesApp:
    """Command line."""

    def Run(self):
        if kS.ErrorMode:
            return 1
        try:
            Parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
            Parser.add_argument("dir", nargs="?", default=os.path.join("examples", "templates"))
            Args = Parser.parse_args()
            for Path in kTestTemplates.Make(Args.dir):
                print(Path)
            return 0 if not kS.ErrorMode else 1
        except kToolException:
            raise
        except Exception as e:
            kS.GlobalErrorHandler(e, "kMakeTestTemplatesApp.Run")
            return 1


if __name__ == "__main__":
    kRun.Main(kMakeTestTemplatesApp)
