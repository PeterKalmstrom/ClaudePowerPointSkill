"""Self-test checks for improving an existing deck: extract_spec.py (restyle by rebuilding) and improve_deck.py
(keep the look). Run by selftest.py (kSelfTestEdit(Test).Run()); uses its Check / RunScript and temp folder.

The fixture is a deck no builder made - plain text boxes, no title placeholders, a slide number on every slide:
  1 cover              two lines, short notes                        -> title, notes verbatim
  2 four figure tiles  big figure over a label each                  -> kpi with four metrics
  3 column chart       a caption line                                -> chart with its data
  4 table              3 x 3                                         -> table
  5 three columns      bold heading over a detail                    -> compare
  6 dates              Jan / Feb / Mar / Apr over a label            -> timeline
  7 list               four lines, one over 100 characters           -> bullets, the overflow in the notes
"""
import json
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Pt

from kShared import kS

LONG_ITEM = ("Every slide keeps its words when the deck is rebuilt, because anything a pattern has no room for is "
             "moved to the speaker notes in full")


class kSelfTestEdit:
    """The checks; Test is the running kSelfTest."""

    def __init__(self, Test):
        try:
            self.Test = Test
            self.Dir = os.path.join(Test.Tmp, "edit")
            self.Source = os.path.join(self.Dir, "foreign.pptx")
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestEdit.__init__")

    @staticmethod
    def Box(Slide, Text, X, Y, W, H, Size, Bold=False):
        """A plain text box (no placeholder), one paragraph per line of Text."""
        if kS.ErrorMode:
            return None
        try:
            Shape = Slide.shapes.add_textbox(Pt(X), Pt(Y), Pt(W), Pt(H))
            Frame = Shape.text_frame
            Frame.word_wrap = True
            for I, Line in enumerate(Text.split("\n")):
                P = Frame.paragraphs[0] if I == 0 else Frame.add_paragraph()
                Run = P.add_run()
                Run.text = Line
                Run.font.size = Pt(Size)
                Run.font.bold = Bold
            return Shape
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestEdit.Box")
            return None

    def NewSlide(self, Prs, Title, Number):
        """A blank slide with a title box at the top and a small slide number bottom right."""
        if kS.ErrorMode:
            return None
        try:
            Slide = Prs.slides.add_slide(Prs.slide_layouts[6])
            if Title:
                self.Box(Slide, Title, 40, 30, 860, 60, 30, True)
            self.Box(Slide, str(Number), 880, 505, 60, 25, 11)
            return Slide
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestEdit.NewSlide")
            return None

    def FigureSlide(self, Prs):
        """Slide 2: four figure tiles."""
        if kS.ErrorMode:
            return None
        try:
            S = self.NewSlide(Prs, "Revenue rose in every region this year", 2)
            for I, (Fig, Label) in enumerate((("5.9 MUSD", "Q4 revenue"), ("112 %", "Net retention"),
                                              ("1.2 %", "Monthly churn"), ("191", "New customers"))):
                self.Box(S, Fig, 40 + I * 220, 150, 180, 60, 40, True)
                self.Box(S, Label, 40 + I * 220, 220, 180, 40, 16)
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestEdit.FigureSlide")
            return None

    def ColumnSlides(self, Prs):
        """Slides 5 (heading + detail columns) and 6 (dates)."""
        if kS.ErrorMode:
            return None
        try:
            S = self.NewSlide(Prs, "Three things change for the team", 5)
            for I, (Head, Detail) in enumerate((("Who", "All forty engineers"), ("Hours", "Thirty-two over four days"),
                                                ("Cover", "On-call rota unchanged"))):
                self.Box(S, Head, 40 + I * 300, 150, 260, 40, 22, True)
                self.Box(S, Detail, 40 + I * 300, 200, 260, 80, 16)
            S = self.NewSlide(Prs, "The pilot runs from January to April", 6)
            for I, (Date, Label) in enumerate((("Jan", "Kick-off"), ("Feb", "First check-in"), ("Mar", "Mid-point review"),
                                               ("Apr", "Final evaluation"))):
                self.Box(S, Date, 40 + I * 220, 160, 150, 40, 20, True)
                self.Box(S, Label, 40 + I * 220, 260, 150, 60, 14)
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestEdit.ColumnSlides")
            return None

    def BuildSource(self):
        """Write the fixture deck (960 x 540 pt, plain text boxes)."""
        if kS.ErrorMode:
            return None
        try:
            os.makedirs(self.Dir, exist_ok=True)
            Prs = Presentation()
            Prs.slide_width, Prs.slide_height = Pt(960), Pt(540)
            S = self.NewSlide(Prs, "", 1)
            self.Box(S, "Northwind review 2026", 60, 150, 820, 80, 44, True)
            self.Box(S, "Executive team | Q4", 60, 250, 820, 50, 20)
            S.notes_slide.notes_text_frame.text = "Welcome everyone. Today we review the year and ask for one decision."
            self.FigureSlide(Prs)
            S = self.NewSlide(Prs, "Revenue grew 44 % from Q1 to Q4", 3)
            Data = CategoryChartData()
            Data.categories = ["Q1", "Q2", "Q3", "Q4"]
            Data.add_series("Revenue (MUSD)", (4.1, 4.6, 5.2, 5.9))
            S.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(40), Pt(110), Pt(600), Pt(360), Data)
            self.Box(S, "Full year 19.8 MUSD", 680, 200, 240, 60, 18)
            S.notes_slide.notes_text_frame.text = "Q1 4.1, Q4 5.9."
            S = self.NewSlide(Prs, "Costs stay inside the budget", 4)
            Table = S.shapes.add_table(3, 3, Pt(40), Pt(120), Pt(860), Pt(200)).table
            for R, Row in enumerate((("Item", "kSEK", "Owner"), ("Buffer", "120", "Finance"), ("Tools", "30", "IT"))):
                for C, Text in enumerate(Row):
                    Table.cell(R, C).text = Text
            self.ColumnSlides(Prs)
            S =self.NewSlide(Prs, "Four rules keep the pilot safe", 7)
            self.Box(S, "\n".join(["Stop rule at the mid-point", "Weekly throughput tracking", "Overtime buffer for crunch",
                                   LONG_ITEM]), 40, 120, 860, 360, 20)
            Prs.save(self.Source)
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestEdit.BuildSource")
            return None

    def CheckExtract(self):
        """extract_spec: patterns, moved text, notes; then build, lint and the text comparison."""
        if kS.ErrorMode:
            return None
        try:
            T = self.Test
            Spec = os.path.join(self.Dir, "foreign.json")
            Code, Out = T.RunScript("extract_spec.py", self.Source, "--out", Spec)
            Data = json.loads(T.ReadText(Spec)) if Code == 0 and os.path.isfile(Spec) else {"slides": []}
            Slides = Data["slides"]
            Kinds = [S["pattern"] for S in Slides]
            T.Check("extract_spec: each slide mapped to the pattern that fits (title, kpi, chart, table, compare, "
                    "timeline, bullets)", Kinds == ["title", "kpi", "chart", "table", "compare", "timeline", "bullets"],
                    f"{Kinds} {Out[-400:]}")
            Ok = len(Slides) == 7
            T.Check("extract_spec: titles from plain text boxes, slide numbers dropped, chart data and metrics kept",
                    Ok and Slides[2]["title"] == "Revenue grew 44 % from Q1 to Q4"
                    and Slides[2]["series"][0]["values"] == [4.1, 4.6, 5.2, 5.9]
                    and [M["value"] for M in Slides[1]["metrics"]] == ["5.9 MUSD", "112 %", "1.2 %", "191"]
                    and Slides[1]["metrics"][0]["label"] == "Q4 revenue", json.dumps(Slides)[:600])
            Notes = json.dumps(Slides[6].get("notes", {})) if Ok else ""
            T.Check("extract_spec: an item over the limit is cut at a break, the full text MOVED FROM SLIDE to the "
                    "notes and flagged", Ok and "MOVED FROM SLIDE: " + LONG_ITEM in json.loads(Notes)["facts"]
                    and all(len(X) <= 100 for X in Slides[6]["items"]) and "flag: slide 7: moved" in Out, Notes[:400])
            T.Check("extract_spec: existing notes become the script verbatim; a slide without notes gets a marked "
                    "DRAFT from its own words", Ok and Slides[0]["notes"]["say"].startswith("Welcome everyone")
                    and "DRAFT" in Slides[3]["notes"]["pitfalls"][0] and "Costs stay inside the budget" in
                    " ".join(Slides[3]["notes"]["say"]), json.dumps([S.get("notes") for S in Slides[:4]])[:500])
            self.CheckRebuild(Spec)
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestEdit.CheckExtract")
            return None

    def CheckRebuild(self, Spec):
        """The extracted spec builds, lints with no errors and keeps every word and figure; a loss is caught."""
        if kS.ErrorMode:
            return None
        try:
            T = self.Test
            Rebuilt = os.path.join(self.Dir, "rebuilt.pptx")
            Code, Out = T.RunScript("build_deck.py", Spec, "--out", Rebuilt, "--lint")
            T.Check("extract_spec: the spec builds and lints with 0 errors", Code in (0, 3) and "\n0 error(s)" in "\n" + Out,
                    Out[-600:])
            Code, Out = T.RunScript("extract_spec.py", self.Source, "--compare", Rebuilt)
            T.Check("extract_spec --compare: every word and figure of the source is in the rebuilt deck",
                    Code == 0 and "all text and figures kept" in Out, Out[-600:])
            Code, Out = T.RunScript("extract_spec.py", self.Source, "--compare", T.Deck)
            T.Check("extract_spec --compare: a deck missing the source's text fails, naming the missing figures",
                    Code == 1 and "missing figures" in Out and "5.9" in Out, Out[-400:])
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestEdit.CheckRebuild")
            return None

    def CheckImprove(self):
        """improve_deck: safe fixes, drafted / repaired notes with the originals kept, the report, the modes."""
        if kS.ErrorMode:
            return None
        try:
            T = self.Test
            Out1 = os.path.join(self.Dir, "improved.pptx")
            Code, Out = T.RunScript("improve_deck.py", self.Source, "--out", Out1)
            Prs = Presentation(Out1) if os.path.isfile(Out1) else None
            Notes = [S.notes_slide.notes_text_frame.text if S.has_notes_slide else "" for S in Prs.slides] if Prs else []
            T.Check("improve_deck: a DRAFT script on every slide without notes, marked for review",
                    len(Notes) == 7 and all("DRAFT script" in N for N in Notes[1:]) and "Revenue rose" in Notes[1],
                    Out[-500:])
            T.Check("improve_deck: notes with no script get the draft in front, the original kept below the divider; "
                    "good notes are left alone", len(Notes) == 7 and Notes[2].rstrip().endswith("Q1 4.1, Q4 5.9.")
                    and "Original notes:" in Notes[2] and Notes[0].startswith("Welcome everyone"), Notes[:3])
            T.Check("improve_deck: lists its changes and ends with the lint report", "notes        slide 2" in Out
                    and "What's left for a person" in Out and "error(s)" in Out, Out[-400:])
            Before = os.path.getmtime(self.Source)
            Code, Out = T.RunScript("improve_deck.py", self.Source, "--dry-run")
            T.Check("improve_deck --dry-run: lists the changes and writes nothing",
                    Code == 0 and "dry run" in Out and os.path.getmtime(self.Source) == Before, Out[-300:])
            Code, Out = T.RunScript("improve_deck.py", self.Source)
            T.Check("improve_deck: refuses to run without --out, --in-place or --dry-run (exit 2)", Code == 2, Out)
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestEdit.CheckImprove")
            return None

    def Run(self):
        """Every check of this module."""
        if kS.ErrorMode:
            return None
        try:
            self.Test.Say("\n-- existing decks: extract_spec.py, improve_deck.py --")
            self.BuildSource()
            self.CheckExtract()
            self.CheckImprove()
            return None
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestEdit.Run")
            return None
