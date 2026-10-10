"""Self-test checks for round 10's post-mortem criticism: a light cover before dark body slides (kCoverTone) and a
weak plan-vs-actual timeline (kPlanActual). Run by selftest.py (kSelfTestPlanActual(Test).Run()); uses its Check /
RunScript / LintFindings and temp folder.
"""
import json
import os

from pptx import Presentation

from kShared import kS

SAY = {"say": "This slide is part of the plan-versus-actual self-test and says what the presenter would say aloud."}
EVENTS = [{"label": "Design and planning", "plan": [0, 1], "actual": [0, 1]},
          {"label": "Service migration", "plan": [1, 4], "actual": [1, 5]},
          {"label": "Data migration", "plan": [3, 5], "actual": [4, 8]},
          {"label": "Cutover and go-live", "plan": [6], "actual": [9]}]


class kSelfTestPlanActual:
    """The checks; Test is the running kSelfTest."""

    def __init__(self, Test):
        try:
            self.Test = Test
            self.Dir = os.path.join(Test.Tmp, "planactual")
            os.makedirs(self.Dir, exist_ok=True)
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestPlanActual.__init__")

    @staticmethod
    def Spec(Direction, Cover=None, Events=None):
        """A cover, a divider, a plan-vs-actual timeline and a kpi slide on Direction."""
        if kS.ErrorMode:
            return {}
        try:
            Spec = {"direction": Direction, "footer": "Plan vs actual test", "slides": [
                {"pattern": "title", "title": "Payments v2 shipped three months late", "subtitle": "Post-mortem"},
                {"pattern": "section", "title": "What happened", "eyebrow": "Part 1"},
                {"pattern": "timeline", "title": "Data migration slipped three months", "unit": "month",
                 "events": Events or EVENTS, "notes": SAY},
                {"pattern": "kpi", "title": "Delivery took 9 months against a plan of 6", "notes": SAY,
                 "metrics": [{"value": "9 months", "label": "Duration"}, {"value": "1.65 MSEK", "label": "Cost"}]}]}
            if Cover:
                Spec["cover"] = Cover
            return Spec
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestPlanActual.Spec")
            return {}

    def Build(self, Name, Spec):
        """Write Spec, build it; (exit code, output, deck path)."""
        if kS.ErrorMode:
            return 1, "", ""
        try:
            Path, Deck = os.path.join(self.Dir, f"{Name}.json"), os.path.join(self.Dir, f"{Name}.pptx")
            with open(Path, "w", encoding="utf-8") as Fh:
                json.dump(Spec, Fh)
            Code, Out = self.Test.RunScript("build_deck.py", Path, "--out", Deck)
            return Code, Out, Deck
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestPlanActual.Build")
            return 1, "", ""

    @staticmethod
    def Names(Deck):
        """Every slide's shapes by name."""
        if kS.ErrorMode:
            return []
        try:
            if not os.path.exists(Deck):
                return []
            return [{Shape.name: Shape for Shape in Slide.shapes} for Slide in Presentation(Deck).slides]
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestPlanActual.Names")
            return []

    def CheckCoverTone(self):
        """A dark direction keeps its cover and dividers dark; a light one keeps the deliberate dark panels."""
        if kS.ErrorMode:
            return
        try:
            T = self.Test
            Code, Out, Deck = self.Build("dark", self.Spec("night-terminal"))
            Dark = self.Names(Deck)
            T.Check("cover tone: a dark direction's cover and divider draw no light panel",
                    Code == 0 and len(Dark) == 4 and "CoverPanel" not in Dark[0] and "SectionPanel" not in Dark[1]
                    and "CoverBand" in Dark[0], Out[-300:])
            Code, Out, Deck = self.Build("light", self.Spec("clean-corporate"))
            Light = self.Names(Deck)
            T.Check("cover tone: a light direction keeps the cover and divider panels in the text colour",
                    Code == 0 and len(Light) == 4 and "CoverPanel" in Light[0] and "SectionPanel" in Light[1],
                    Out[-300:])
            Code, Out, Deck = self.Build("forced", self.Spec("blueprint", "panel"))
            Forced = self.Names(Deck)
            T.Check("cover tone: 'cover': 'panel' on a dark direction draws the panel and warns of the clash",
                    Code == 0 and Forced and "CoverPanel" in Forced[0]
                    and "read as a clash" in Out, Out[-400:])
            Code, Out, _ = self.Build("badcover", self.Spec("clean-corporate", "loud"))
            T.Check("cover tone: an unknown 'cover' is a spec error", Code != 0 and "'cover' is 'loud'" in Out,
                    Out[-300:])
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestPlanActual.CheckCoverTone")
            return

    def CheckDrawn(self):
        """Bar pairs, the overrun in the accent, the slip labelled in units, the lead row bold, notes in words."""
        if kS.ErrorMode:
            return
        try:
            T = self.Test
            Code, Out, Deck = self.Build("planactual", self.Spec("clean-corporate"))
            Bad = [f"{Sl}:{C}" for Sl, C, _, Sev in T.LintFindings(Deck) if Sev in ("error", "warn")]
            T.Check("plan vs actual: the timeline builds and lints clean", Code == 0 and not Bad,
                    ", ".join(Bad) + Out[-300:])
            Names = self.Names(Deck)
            if len(Names) < 3:
                return
            Sl = Names[2]
            T.Check("plan vs actual: a planned bar and an actual bar per milestone, the overrun separate",
                    {"PaPlan1", "PaActual1", "PaPlan3", "PaActual3", "PaOverrun3", "PaSlipLine4"} <= set(Sl)
                    and "PaOverrun1" not in Sl and Sl["PaPlan3"].top < Sl["PaActual3"].top, str(sorted(Sl)))
            Slips = [Sl[f"PaSlip{I}"].text_frame.text for I in range(1, 5) if f"PaSlip{I}" in Sl]
            T.Check("plan vs actual: the slip is labelled in units at the row's end",
                    Slips == ["On time", "+1 month", "+3 months", "+3 months"], repr(Slips))
            Runs = [R for P in Sl["PaLabel3"].text_frame.paragraphs for R in P.runs]
            T.Check("plan vs actual: the largest slip's label is bold, the axis is labelled M1..M9",
                    Runs and Runs[0].font.bold and "PaTick9" in Sl and Sl["PaTick1"].text_frame.text == "M1",
                    str(sorted(Sl)))
            Notes = Presentation(Deck).slides[2].notes_slide.notes_text_frame.text
            T.Check("plan vs actual: the notes give each milestone's plan, actual and slip in words",
                    "PLAN VS ACTUAL:" in Notes and "Data migration: planned M4-M5, actual M5-M8 (+3 months)"
                    in Notes and "Cutover and go-live: planned end of M6, actual end of M9" in Notes, Notes[-400:])
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestPlanActual.CheckDrawn")
            return

    def CheckErrors(self):
        """A missing actual, an end before its start, a bar against a milestone: spec errors."""
        if kS.ErrorMode:
            return
        try:
            T = self.Test
            Events = [{"label": "Design", "plan": [0, 1], "actual": [0, 1]},
                      {"label": "Build", "plan": [1, 4]},
                      {"label": "Test", "plan": [4, 3], "actual": [4, 6]},
                      {"label": "Launch", "plan": [6], "actual": [6, 9]}]
            Code, Out, _ = self.Build("planbad", self.Spec("clean-corporate", Events=Events))
            T.Check("plan vs actual: missing actual, end before start and a mixed bar/milestone are spec errors",
                    Code != 0 and "give every event a 'plan' and an 'actual'" in Out and "end >= start" in Out
                    and "mixes a bar and a milestone" in Out, Out[-600:])
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestPlanActual.CheckErrors")
            return

    def Run(self):
        if kS.ErrorMode:
            return
        try:
            self.Test.Say("plan vs actual and cover tone (round 10)")
            self.CheckCoverTone()
            self.CheckDrawn()
            self.CheckErrors()
        except Exception as e:
            kS.GlobalErrorHandler(e, "kSelfTestPlanActual.Run")
            return
